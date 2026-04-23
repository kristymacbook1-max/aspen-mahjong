"""Build the Unit Pricing workbook (unit_pricing.xlsx).

Tabs:
  1. README
  2. Deliverables
  3. Factors
  4. TechTable
  5. TokenTable
  6. Inputs
  7. Calc
  8. Calibration

Key tables (Excel structured tables, referenced by name in formulas):
  Deliverables, Factors, TechTable, TokenTable, CalcFactors, Calibration
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.utils import get_column_letter

WB = Workbook()
WB.remove(WB.active)

BLUE = Font(color="0047AB", bold=False)       # inputs
BLACK = Font(color="000000", bold=False)      # formulas
HEADER = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
NOTE_FILL = PatternFill("solid", fgColor="FFF2CC")
THIN = Side(border_style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

CURRENCY_FMT = '$#,##0;($#,##0);-'
PCT_FMT = "0.0%"

TABLE_STYLE = TableStyleInfo(
    name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
    showRowStripes=True, showColumnStripes=False,
)


def autosize(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def write_headers(ws, row, headers):
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=c, value=h)
        cell.font = HEADER
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="left", vertical="center")
        cell.border = BOX


def add_table(ws, name, ref):
    t = Table(displayName=name, ref=ref)
    t.tableStyleInfo = TABLE_STYLE
    ws.add_table(t)


# ---------------------------------------------------------------- README
def build_readme():
    ws = WB.create_sheet("README")
    ws["A1"] = "Unit Pricing Workbook"
    ws["A1"].font = Font(bold=True, size=16)

    notes = [
        "",
        "Purpose: price a deliverable by combining a base price, factor scores,",
        "a tech-stack tier, a token-usage tier, residual T&M, and an",
        "AI-automation discount.",
        "",
        "Convention:",
        "  • Blue font = input (editable)",
        "  • Black font = formula (do not overwrite)",
        "",
        "Workflow:",
        "  1. Fill Inputs: matter, deliverable, tiers, T&M hours/rate.",
        "  2. Assign 7 scores on Calc (1, 3, or 5) — AI is a discount only.",
        "  3. Read Final Unit Price (highlighted) and Low/High range.",
        "  4. If overriding, populate Override_Price, Override_Reason, Approver.",
        "  5. Copy the Calc snapshot row → paste values onto next Calibration row.",
        "  6. After the matter closes, fill Actual_Final_Price / Actual_Effort_Cost.",
        "",
        "Scoring rule (see Factors tab):",
        "  • 3 = baseline",
        "  • 1 and 5 require explanation",
        "  • AI automation is a downward adjustment only",
        "",
        "Model knobs (in order to calibrate):",
        "  1. Base deliverable prices (Deliverables tab)",
        "  2. Tech / Token tiers",
        "  3. Multiplier coefficient (Inputs!Mult_Coef, default 0.35)",
        "  4. Factor weights (Calc!CalcFactors) — must sum to 100%",
        "",
        "Pilot: run 10-20 recent matters through the file;",
        "compare to partner's proposal and actual cost; adjust in the order above.",
    ]
    for i, line in enumerate(notes, start=2):
        ws.cell(row=i, column=1, value=line)
    autosize(ws, [90])


# ---------------------------------------------------------------- Deliverables
def build_deliverables():
    ws = WB.create_sheet("Deliverables")
    headers = ["Deliverable_Name", "Base_Unit_Price", "Typical_Turnaround_Days", "Notes"]
    rows = [
        ["Quick advisory email",        1500,  2,  "Informal internal response"],
        ["Research memo",               6500,  7,  "Structured analytical memo"],
        ["Formal opinion",             22000, 21,  "Workpaper- or client-ready, defensible"],
        ["Transaction cost analysis",  18000, 14,  "Tax/accounting analysis"],
        ["Workpaper package",          12000, 10,  "Full supporting documentation"],
        ["Review / second-look",        4000,  5,  "Reviewer pass on prior deliverable"],
        ["Oral consultation",           2500,  1,  "Call-based, minimal writeup"],
    ]
    write_headers(ws, 1, headers)
    for r, row in enumerate(rows, start=2):
        for c, v in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = BOX
            if c == 2:
                cell.number_format = CURRENCY_FMT
    end = f"D{len(rows) + 1}"
    add_table(ws, "Deliverables", f"A1:{end}")
    autosize(ws, [32, 18, 24, 50])


# ---------------------------------------------------------------- Factors
def build_factors():
    ws = WB.create_sheet("Factors")

    # Note block above the table
    ws["A1"] = "Scoring notes"
    ws["A1"].font = Font(bold=True)
    notes = [
        "• 3 = baseline",
        "• 1 and 5 require explanation",
        "• AI automation is a downward adjustment only",
    ]
    for i, n in enumerate(notes, start=2):
        c = ws.cell(row=i, column=1, value=n)
        c.fill = NOTE_FILL

    header_row = 6
    headers = ["Factor_Name", "Score_1", "Score_3", "Score_5"]
    write_headers(ws, header_row, headers)

    rows = [
        ["Technical complexity / guidance profile",
         "Settled issue, routine authority, limited judgment",
         "Moderate analysis, standard research burden",
         "Novel, ambiguous, conflicting authority, specialist depth needed"],
        ["Economic / stakeholder risk",
         "Low impact, low visibility",
         "Moderate impact and visibility",
         "High materiality, executive visibility, significant exposure"],
        ["Controversy risk",
         "Unlikely to be challenged",
         "Plausible review/exam interest",
         "High likelihood of challenge or controversy sensitivity"],
        ["Deliverable rigor",
         "Informal internal response",
         "Structured memo or analytical package",
         "Formal, defensible, workpaper-ready or client-ready package"],
        ["Scope breadth",
         "Narrow issue, one entity/period",
         "Moderate spread",
         "Broad, multi-entity, multi-year, multi-issue"],
        ["Data quality / accessibility",
         "Clean, structured, available",
         "Some gaps/cleanup",
         "Fragmented, incomplete, conflicting, manual"],
        ["Data sensitivity / governance burden",
         "Ordinary internal-use data",
         "Moderate confidentiality constraints",
         "Highly sensitive, privileged, restricted, governance-heavy"],
        ["AI automation level",
         "Minimal reusable AI support",
         "Standard AI-assisted process",
         "Highly automated, reusable, efficient workflow"],
    ]
    for r, row in enumerate(rows, start=header_row + 1):
        for c, v in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = BOX
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    end = f"D{header_row + len(rows)}"
    add_table(ws, "Factors", f"A{header_row}:{end}")
    autosize(ws, [42, 42, 42, 55])
    for r in range(header_row + 1, header_row + 1 + len(rows)):
        ws.row_dimensions[r].height = 30


# ---------------------------------------------------------------- TechTable
def build_tech():
    ws = WB.create_sheet("TechTable")
    headers = ["Tech_Tier", "Description", "Multiplier"]
    rows = [
        ["T1", "No specialized tooling",                        1.00],
        ["T2", "Standard research tooling",                     1.10],
        ["T3", "Specialized platform / licensed data",          1.25],
        ["T4", "Heavy custom tooling / modeling",               1.45],
        ["T5", "Bespoke build, multi-system integration",       1.75],
    ]
    write_headers(ws, 1, headers)
    for r, row in enumerate(rows, start=2):
        for c, v in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = BOX
            if c == 3:
                cell.number_format = "0.00"
    end = f"C{len(rows) + 1}"
    add_table(ws, "TechTable", f"A1:{end}")
    autosize(ws, [12, 50, 14])


# ---------------------------------------------------------------- TokenTable
def build_tokens():
    ws = WB.create_sheet("TokenTable")
    headers = ["Token_Tier", "Est_Tokens", "Token_Cost"]
    rows = [
        ["K1 (<100k)",       50000,    25],
        ["K2 (100k-500k)",  300000,   150],
        ["K3 (500k-2M)",   1200000,   600],
        ["K4 (2M-10M)",    5000000,  2500],
        ["K5 (>10M)",     15000000,  7500],
    ]
    write_headers(ws, 1, headers)
    for r, row in enumerate(rows, start=2):
        for c, v in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = BOX
            if c == 2:
                cell.number_format = "#,##0"
            if c == 3:
                cell.number_format = CURRENCY_FMT
    end = f"C{len(rows) + 1}"
    add_table(ws, "TokenTable", f"A1:{end}")
    autosize(ws, [18, 16, 16])


# ---------------------------------------------------------------- Inputs
def build_inputs():
    ws = WB.create_sheet("Inputs")
    ws["A1"] = "Matter inputs"
    ws["A1"].font = Font(bold=True, size=14)

    ws["A3"] = "Blue = input. Black = formula."
    ws["A3"].font = Font(italic=True, color="595959")

    pairs = [
        ("Matter_ID",         "M-0001"),
        ("Matter_Name",       "Sample matter"),
        ("Deliverable_Type",  "Research memo"),
        ("Tech_Tier",         "T2"),
        ("Token_Tier",        "K2 (100k-500k)"),
        ("Override_Price",    None),
        ("Override_Reason",   None),
        ("Approver",          None),
        ("TM_Hours",          0),
        ("TM_Rate",           350),
        ("Mult_Coef",         0.35),
    ]
    for i, (label, val) in enumerate(pairs, start=5):
        lc = ws.cell(row=i, column=1, value=label)
        lc.font = Font(bold=True)
        lc.fill = PatternFill("solid", fgColor="F2F2F2")
        lc.border = BOX
        vc = ws.cell(row=i, column=2, value=val)
        vc.font = BLUE
        vc.border = BOX
        if label in ("Override_Price", "TM_Rate"):
            vc.number_format = CURRENCY_FMT
        if label == "Mult_Coef":
            vc.number_format = "0.000"
        if label == "TM_Hours":
            vc.number_format = "0.0"

    # Named references for easy formula use
    defined_names = {
        "in_MatterID":        "Inputs!$B$5",
        "in_MatterName":      "Inputs!$B$6",
        "in_Deliverable":     "Inputs!$B$7",
        "in_TechTier":        "Inputs!$B$8",
        "in_TokenTier":       "Inputs!$B$9",
        "in_Override":        "Inputs!$B$10",
        "in_OverrideReason":  "Inputs!$B$11",
        "in_Approver":        "Inputs!$B$12",
        "in_TMHours":         "Inputs!$B$13",
        "in_TMRate":          "Inputs!$B$14",
        "in_MultCoef":        "Inputs!$B$15",
    }
    for nm, ref in defined_names.items():
        from openpyxl.workbook.defined_name import DefinedName
        WB.defined_names[nm] = DefinedName(name=nm, attr_text=ref)

    # Data validation dropdowns
    dv_deliv = DataValidation(
        type="list",
        formula1="=INDIRECT(\"Deliverables[Deliverable_Name]\")",
        allow_blank=True,
    )
    dv_deliv.error = "Pick from Deliverables"
    dv_deliv.prompt = "Pick from Deliverables"
    ws.add_data_validation(dv_deliv)
    dv_deliv.add("B7")

    dv_tech = DataValidation(
        type="list",
        formula1="=INDIRECT(\"TechTable[Tech_Tier]\")",
        allow_blank=True,
    )
    ws.add_data_validation(dv_tech)
    dv_tech.add("B8")

    dv_token = DataValidation(
        type="list",
        formula1="=INDIRECT(\"TokenTable[Token_Tier]\")",
        allow_blank=True,
    )
    ws.add_data_validation(dv_token)
    dv_token.add("B9")

    # Override reason required if override price present
    ws["D10"] = "If Override_Price is set, Override_Reason and Approver are required."
    ws["D10"].font = Font(italic=True, color="C00000")
    ws["D11"] = ("=IF(AND(ISNUMBER(in_Override),"
                 "OR(LEN(in_OverrideReason)=0,LEN(in_Approver)=0)),"
                 "\"Missing override justification\",\"OK\")")
    ws["D11"].font = Font(bold=True)

    autosize(ws, [22, 32, 4, 55])


# ---------------------------------------------------------------- Calc
def build_calc():
    ws = WB.create_sheet("Calc")
    ws["A1"] = "Unit Price Calculation"
    ws["A1"].font = Font(bold=True, size=14)

    ws["A3"] = "Matter:"
    ws["B3"] = "=in_MatterName"
    ws["B3"].font = BLACK
    ws["A4"] = "Deliverable:"
    ws["B4"] = "=in_Deliverable"
    ws["B4"].font = BLACK

    # --- CalcFactors table: factor name, weight, score, contribution ---
    header_row = 6
    headers = ["Factor_Name", "Weight", "Score", "Contribution"]
    write_headers(ws, header_row, headers)

    # AI automation is excluded from the weighted multiplier (treated
    # purely as a downward discount below). Its weight is held at 0 and
    # the other seven renormalize to sum to 1.0.
    factors = [
        ("Technical complexity / guidance profile", 0.21),
        ("Economic / stakeholder risk",             0.16),
        ("Controversy risk",                        0.16),
        ("Deliverable rigor",                       0.16),
        ("Scope breadth",                           0.11),
        ("Data quality / accessibility",            0.10),
        ("Data sensitivity / governance burden",    0.10),
        ("AI automation level",                     0.00),
    ]
    first_row = header_row + 1
    for i, (name, weight) in enumerate(factors):
        r = first_row + i
        a = ws.cell(row=r, column=1, value=name)
        b = ws.cell(row=r, column=2, value=weight)
        c = ws.cell(row=r, column=3, value=3)     # baseline score (input)
        d = ws.cell(row=r, column=4,
                    value=f"=[@Weight]*[@Score]")
        a.border = BOX
        b.border = BOX; b.number_format = PCT_FMT; b.font = BLUE
        c.border = BOX; c.font = BLUE
        d.border = BOX; d.number_format = "0.000"; d.font = BLACK

    last_row = first_row + len(factors) - 1
    add_table(ws, "CalcFactors", f"A{header_row}:D{last_row}")

    # Heatmap on Score column C
    score_range = f"C{first_row}:C{last_row}"
    ws.conditional_formatting.add(
        score_range,
        CellIsRule(operator="lessThanOrEqual", formula=["2"],
                   fill=PatternFill("solid", fgColor="C6EFCE")),
    )
    ws.conditional_formatting.add(
        score_range,
        CellIsRule(operator="equal", formula=["3"],
                   fill=PatternFill("solid", fgColor="FFEB9C")),
    )
    ws.conditional_formatting.add(
        score_range,
        CellIsRule(operator="greaterThanOrEqual", formula=["4"],
                   fill=PatternFill("solid", fgColor="F4CCCC")),
    )

    # --- Totals & price build-up ---
    out_row = last_row + 3
    labels_and_formats = [
        ("Weighted factor score (w̄)",                          "0.000"),
        ("Multiplier coefficient",                              "0.000"),
        ("Factor multiplier = 1 + coef × (w̄ − 3)",             "0.000"),
        ("Base deliverable price",                              CURRENCY_FMT),
        ("Tech multiplier",                                     "0.00"),
        ("Token add-on",                                        CURRENCY_FMT),
        ("Residual T&M hours",                                  "0.0"),
        ("Residual T&M rate",                                   CURRENCY_FMT),
        ("Residual T&M cost",                                   CURRENCY_FMT),
        ("AI automation score (1-5)",                           "0"),
        ("AI automation discount (downward only)",              PCT_FMT),
        ("Subtotal (base × factor × tech + token + T&M)",       CURRENCY_FMT),
        ("Unit Price (calculated)",                             CURRENCY_FMT),
        ("Override in effect?",                                 "@"),
        ("Final Unit Price",                                    CURRENCY_FMT),
        ("Override_Check",                                      "@"),
        ("Low_End_Range",                                       CURRENCY_FMT),
        ("High_End_Range",                                      CURRENCY_FMT),
        ("Weight_Sum_Check (must = 100%)",                      PCT_FMT),
    ]
    rows_map = {i: out_row + i for i in range(len(labels_and_formats))}
    r_wbar = rows_map[0]
    r_coef = rows_map[1]
    r_fmul = rows_map[2]
    r_base = rows_map[3]
    r_tech = rows_map[4]
    r_tok  = rows_map[5]
    r_tmh  = rows_map[6]
    r_tmr  = rows_map[7]
    r_tm   = rows_map[8]
    r_ai   = rows_map[9]
    r_disc = rows_map[10]
    r_sub  = rows_map[11]
    r_calc = rows_map[12]
    r_ovr  = rows_map[13]
    r_fin  = rows_map[14]
    r_chk  = rows_map[15]
    r_lo   = rows_map[16]
    r_hi   = rows_map[17]
    r_wsum = rows_map[18]

    resolved = [
        "=SUMPRODUCT(CalcFactors[Weight],CalcFactors[Score])",
        "=in_MultCoef",
        f"=1+B{r_coef}*(B{r_wbar}-3)",
        "=VLOOKUP(in_Deliverable,Deliverables[[Deliverable_Name]:[Base_Unit_Price]],2,FALSE)",
        "=VLOOKUP(in_TechTier,TechTable[[Tech_Tier]:[Multiplier]],3,FALSE)",
        "=VLOOKUP(in_TokenTier,TokenTable[[Token_Tier]:[Token_Cost]],3,FALSE)",
        "=in_TMHours",
        "=in_TMRate",
        f"=B{r_tmh}*B{r_tmr}",
        "=INDEX(CalcFactors[Score],MATCH(\"AI automation level\",CalcFactors[Factor_Name],0))",
        f"=MAX(0,(B{r_ai}-3)/2)*0.15",
        f"=B{r_base}*B{r_fmul}*B{r_tech}+B{r_tok}+B{r_tm}",
        f"=B{r_sub}*(1-B{r_disc})",
        "=IF(ISNUMBER(in_Override),\"YES\",\"no\")",
        f"=IF(ISNUMBER(in_Override),in_Override,B{r_calc})",
        ("=IF(AND(ISNUMBER(in_Override),"
         "OR(LEN(in_OverrideReason)=0,LEN(in_Approver)=0)),"
         "\"Missing override reason\",\"OK\")"),
        f"=B{r_fin}*0.9",
        f"=B{r_fin}*1.1",
        "=SUM(CalcFactors[Weight])",
    ]

    for i, (label, fmt) in enumerate(labels_and_formats):
        r = rows_map[i]
        la = ws.cell(row=r, column=1, value=label)
        la.font = Font(bold=True)
        la.fill = PatternFill("solid", fgColor="F2F2F2")
        la.border = BOX
        vb = ws.cell(row=r, column=2, value=resolved[i])
        vb.font = BLACK
        vb.border = BOX
        if fmt != "@":
            vb.number_format = fmt

    # Highlight final price
    ws.cell(row=r_fin, column=1).fill = PatternFill("solid", fgColor="DDEBF7")
    ws.cell(row=r_fin, column=2).fill = PatternFill("solid", fgColor="DDEBF7")
    ws.cell(row=r_fin, column=2).font = Font(bold=True, color="000000", size=12)

    # Flag when weights don't sum to exactly 100%
    ws.conditional_formatting.add(
        f"B{r_wsum}",
        FormulaRule(formula=[f"ROUND(B{r_wsum},4)<>1"],
                    fill=PatternFill("solid", fgColor="F4CCCC"),
                    font=Font(bold=True, color="9C0006")),
    )

    # --- Snapshot block (copy → paste values onto next Calibration row) ---
    snap_header_row = r_wsum + 3
    ws.cell(row=snap_header_row - 1, column=1,
            value="Calibration snapshot — copy row below, paste values onto next Calibration row:"
            ).font = Font(italic=True, bold=True, color="595959")

    snap_cols = [
        ("Matter_ID",            "=in_MatterID",                        "@"),
        ("Matter_Name",          "=in_MatterName",                      "@"),
        ("Deliverable_Type",     "=in_Deliverable",                     "@"),
        ("Tech_Tier",            "=in_TechTier",                        "@"),
        ("Token_Tier",           "=in_TokenTier",                       "@"),
        ("Score_TechCplx",       "=INDEX(CalcFactors[Score],1)",        "0"),
        ("Score_EconRisk",       "=INDEX(CalcFactors[Score],2)",        "0"),
        ("Score_Controversy",    "=INDEX(CalcFactors[Score],3)",        "0"),
        ("Score_Rigor",          "=INDEX(CalcFactors[Score],4)",        "0"),
        ("Score_Scope",          "=INDEX(CalcFactors[Score],5)",        "0"),
        ("Score_DataQ",          "=INDEX(CalcFactors[Score],6)",        "0"),
        ("Score_DataSens",       "=INDEX(CalcFactors[Score],7)",        "0"),
        ("Score_AI",             "=INDEX(CalcFactors[Score],8)",        "0"),
        ("TM_Hours",             "=in_TMHours",                         "0.0"),
        ("TM_Rate",              "=in_TMRate",                          CURRENCY_FMT),
        ("Mult_Coef",            "=in_MultCoef",                        "0.000"),
        ("Proposed_Unit_Price",  f"=B{r_fin}",                          CURRENCY_FMT),
    ]
    for c, (lbl, _, _) in enumerate(snap_cols, start=1):
        cell = ws.cell(row=snap_header_row, column=c, value=lbl)
        cell.font = HEADER
        cell.fill = HEADER_FILL
        cell.border = BOX
    for c, (_, formula, fmt) in enumerate(snap_cols, start=1):
        cell = ws.cell(row=snap_header_row + 1, column=c, value=formula)
        cell.font = BLACK
        cell.border = BOX
        if fmt != "@":
            cell.number_format = fmt

    autosize(ws, [50, 22, 16, 14, 18, 18, 18, 18, 18, 18, 18, 18, 14, 14, 14, 22])


# ---------------------------------------------------------------- Calibration
def build_calibration():
    ws = WB.create_sheet("Calibration")
    headers = [
        # audit snapshot (populate via Calc log row → paste values)
        "Matter_ID", "Matter_Name", "Deliverable_Type",
        "Tech_Tier", "Token_Tier",
        "Score_TechCplx", "Score_EconRisk", "Score_Controversy",
        "Score_Rigor", "Score_Scope", "Score_DataQ",
        "Score_DataSens", "Score_AI",
        "TM_Hours", "TM_Rate", "Mult_Coef",
        "Proposed_Unit_Price",
        # outcomes (fill in after the matter closes)
        "Actual_Final_Price", "Actual_Effort_Cost", "Margin_%",
        "Partner_Feedback", "Client_Reaction",
        "Adjustment_Needed", "Notes",
    ]
    write_headers(ws, 1, headers)

    seed_row = 2
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=seed_row, column=c, value=None)
        cell.border = BOX

    currency_cols = {"TM_Rate", "Proposed_Unit_Price",
                     "Actual_Final_Price", "Actual_Effort_Cost"}
    for idx, h in enumerate(headers, start=1):
        if h in currency_cols:
            ws.cell(row=seed_row, column=idx).number_format = CURRENCY_FMT

    margin_col = headers.index("Margin_%") + 1
    ws.cell(row=seed_row, column=margin_col,
            value=("=IFERROR(([@Actual_Final_Price]-[@Actual_Effort_Cost])"
                   "/[@Actual_Final_Price],0)"))
    ws.cell(row=seed_row, column=margin_col).number_format = PCT_FMT
    ws.cell(row=seed_row, column=margin_col).font = BLACK

    last_col = get_column_letter(len(headers))
    add_table(ws, "Calibration", f"A1:{last_col}{seed_row}")
    widths = [
        12, 26, 22,
        12, 18,
        14, 14, 16, 12, 12, 12, 14, 12,
        12, 12, 12,
        20,
        20, 20, 10,
        26, 22, 22, 40,
    ]
    autosize(ws, widths)


# ---------------------------------------------------------------- build
build_readme()
build_deliverables()
build_factors()
build_tech()
build_tokens()
build_inputs()
build_calc()
build_calibration()

WB.active = WB.sheetnames.index("Calc")

out = "/home/user/aspen-mahjong/unit_pricing.xlsx"
WB.save(out)
print(f"wrote {out}")
