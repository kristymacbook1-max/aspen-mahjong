"""Phase 1 Fixed Asset / Depreciation Excel Report Generator.

Produces a 7-tab workbook:
  1. Executive Summary
  2. Depreciation Analysis
  3. Bonus Depreciation
  4. Section 179
  5. Like-Kind Exchanges
  6. Cost Segregation
  7. Phase 2 Roadmap
"""

import os, sys
from decimal import Decimal
from openpyxl import Workbook

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', '..'))
from revenue_recognition.utils.excel_styles import (
    FONT_TITLE, FONT_BODY, FONT_BODY_BOLD, THIN_BORDER, FMT_CURRENCY,
    apply_header_row, apply_section_header, apply_title,
    auto_fit_columns, write_data_row, format_currency_cell, get_risk_font,
)
from revenue_recognition.utils.currency_helpers import to_decimal, format_currency, format_millions


class Phase1FixedAssetReport:

    def generate(self, results: dict, output_path: str) -> str:
        self._r = results
        self._wb = Workbook()
        self._wb.remove(self._wb.active)

        self._create_executive_summary()
        self._create_depreciation()
        self._create_bonus()
        self._create_179()
        self._create_lke()
        self._create_cost_seg()
        self._create_roadmap()

        os.makedirs(os.path.dirname(output_path) or '.', exist_ok=True)
        self._wb.save(output_path)
        return output_path

    def _create_executive_summary(self):
        ws = self._wb.create_sheet("Executive Summary")
        apply_title(ws, 1, 1, 5, f"Fixed Asset Analysis — {self._r['company'].name}")
        row = 3

        dep = self._r["depreciation_analysis"]
        bonus = self._r["bonus_depreciation"]
        info = [
            ("Total Book Depreciation", format_currency(dep["total_book_depreciation"])),
            ("Total Tax Depreciation", format_currency(dep["total_tax_depreciation"])),
            ("Book-Tax Difference", format_currency(dep["total_difference"])),
            ("Bonus Depreciation", format_currency(bonus["bonus_amount"])),
            (f"Bonus Rate ({bonus['year']})", bonus["bonus_rate_pct"]),
            ("Cost Seg Opportunities", str(self._r["cost_segregation"]["count"])),
            ("§1031 Deferred Gain", format_currency(self._r["like_kind_exchanges"]["total_deferred_gain"])),
        ]
        for label, val in info:
            ws.cell(row=row, column=1, value=label).font = FONT_BODY_BOLD
            ws.cell(row=row, column=2, value=val).font = FONT_BODY
            row += 1

        row += 1
        apply_section_header(ws, row, 1, 5, "Key Findings")
        row += 1
        for f in self._r["findings"]:
            ws.cell(row=row, column=1, value=f["area"]).font = FONT_BODY_BOLD
            ws.cell(row=row, column=2, value=f["description"]).font = FONT_BODY
            ws.cell(row=row, column=3, value=f["risk_level"].upper()).font = get_risk_font(f["risk_level"])
            row += 1
        auto_fit_columns(ws)

    def _create_depreciation(self):
        ws = self._wb.create_sheet("Depreciation Analysis")
        apply_title(ws, 1, 1, 10, "Book vs. Tax Depreciation by Asset Category")
        row = 3

        headers = ["Category", "Gross Value", "Net Book Value", "Book Life", "Book Method",
                    "MACRS Class", "Book Depr", "Tax Depr", "Difference", "Additions"]
        for i, h in enumerate(headers, 1):
            ws.cell(row=row, column=i, value=h)
        apply_header_row(ws, row, 1, len(headers))
        row += 1

        for item in self._r["depreciation_analysis"]["items"]:
            write_data_row(ws, row, [
                item["name"], float(to_decimal(item["gross_value"])),
                float(to_decimal(item["net_book_value"])), item["book_life"],
                item["book_method"], item["macrs_class"],
                float(to_decimal(item["book_depreciation"])),
                float(to_decimal(item["tax_depreciation"])),
                float(to_decimal(item["difference"])),
                float(to_decimal(item["additions"])),
            ], alternate=(row % 2 == 0))
            for c in [2, 3, 7, 8, 9, 10]:
                format_currency_cell(ws.cell(row=row, column=c))
            row += 1

        row += 1
        dep = self._r["depreciation_analysis"]
        for label, col, val in [("Total Book", 7, dep["total_book_depreciation"]),
                                ("Total Tax", 8, dep["total_tax_depreciation"]),
                                ("Difference", 9, dep["total_difference"])]:
            ws.cell(row=row, column=1, value=label).font = FONT_BODY_BOLD
            c = ws.cell(row=row, column=col, value=float(to_decimal(val)))
            format_currency_cell(c)
            c.font = FONT_BODY_BOLD
            row += 1
        auto_fit_columns(ws)

    def _create_bonus(self):
        ws = self._wb.create_sheet("Bonus Depreciation")
        apply_title(ws, 1, 1, 3, "§168(k) Bonus Depreciation Analysis")
        row = 3

        bonus = self._r["bonus_depreciation"]
        lines = [
            ("Qualifying Additions", bonus["qualifying_additions"]),
            ("Bonus Rate", bonus["bonus_rate_pct"]),
            ("Bonus Amount", bonus["bonus_amount"]),
        ]
        for label, val in lines:
            ws.cell(row=row, column=1, value=label).font = FONT_BODY_BOLD
            if isinstance(val, str):
                ws.cell(row=row, column=2, value=val).font = FONT_BODY
            else:
                c = ws.cell(row=row, column=2, value=float(to_decimal(val)))
                format_currency_cell(c)
            ws.cell(row=row, column=1).border = THIN_BORDER
            row += 1

        row += 1
        ws.cell(row=row, column=1, value=bonus["phase_down_note"]).font = FONT_BODY

        row += 2
        apply_section_header(ws, row, 1, 3, "Bonus Depreciation Phase-Down Schedule")
        row += 1
        for yr, rate in [(2022, "100%"), (2023, "80%"), (2024, "60%"), (2025, "40%"), (2026, "20%"), (2027, "0%")]:
            ws.cell(row=row, column=1, value=str(yr)).font = FONT_BODY_BOLD if yr == bonus["year"] else FONT_BODY
            ws.cell(row=row, column=2, value=rate).font = FONT_BODY_BOLD if yr == bonus["year"] else FONT_BODY
            row += 1
        auto_fit_columns(ws)

    def _create_179(self):
        ws = self._wb.create_sheet("Section 179")
        apply_title(ws, 1, 1, 3, "§179 Expensing Analysis")
        row = 3

        s = self._r["section_179"]
        lines = [
            ("Qualifying Property", s["qualifying_property"]),
            ("§179 Limit (2024)", s["limit"]),
            ("Phase-Out Threshold", s["phase_out_threshold"]),
            ("Total Placed in Service", s["total_placed_in_service"]),
            ("Phase-Out Reduction", s["phase_out_reduction"]),
            ("Available §179", s["available"]),
            ("Elected Amount", s["elected"]),
        ]
        for label, val in lines:
            ws.cell(row=row, column=1, value=label).font = FONT_BODY
            c = ws.cell(row=row, column=2, value=float(to_decimal(val)))
            format_currency_cell(c)
            ws.cell(row=row, column=1).border = THIN_BORDER
            c.border = THIN_BORDER
            row += 1
        auto_fit_columns(ws)

    def _create_lke(self):
        ws = self._wb.create_sheet("Like-Kind Exchanges")
        apply_title(ws, 1, 1, 8, "§1031 Like-Kind Exchange Analysis")
        row = 3

        if not self._r["like_kind_exchanges"]["items"]:
            ws.cell(row=row, column=1, value="No §1031 exchanges identified.").font = FONT_BODY
        else:
            headers = ["Description", "Relinquished", "Replacement", "FMV", "Basis",
                        "Realized Gain", "Recognized", "Deferred"]
            for i, h in enumerate(headers, 1):
                ws.cell(row=row, column=i, value=h)
            apply_header_row(ws, row, 1, len(headers))
            row += 1

            for item in self._r["like_kind_exchanges"]["items"]:
                write_data_row(ws, row, [
                    item["description"], item["relinquished"], item["replacement"],
                    float(to_decimal(item["fmv"])), float(to_decimal(item["basis"])),
                    float(to_decimal(item["realized_gain"])),
                    float(to_decimal(item["recognized_gain"])),
                    float(to_decimal(item["deferred_gain"])),
                ], alternate=(row % 2 == 0))
                for c in [4, 5, 6, 7, 8]:
                    format_currency_cell(ws.cell(row=row, column=c))
                row += 1
        auto_fit_columns(ws)

    def _create_cost_seg(self):
        ws = self._wb.create_sheet("Cost Segregation")
        apply_title(ws, 1, 1, 5, "Cost Segregation Opportunities")
        row = 3

        cs = self._r["cost_segregation"]
        if not cs["candidates"]:
            ws.cell(row=row, column=1, value="No cost segregation candidates identified.").font = FONT_BODY
        else:
            headers = ["Property", "Gross Value", "Current Class", "Potential Reclassification", "Est. Acceleration"]
            for i, h in enumerate(headers, 1):
                ws.cell(row=row, column=i, value=h)
            apply_header_row(ws, row, 1, len(headers))
            row += 1

            for c in cs["candidates"]:
                write_data_row(ws, row, [
                    c["name"], float(to_decimal(c["gross_value"])),
                    c["current_class"], c["potential_reclassification"],
                    float(to_decimal(c["estimated_acceleration"])),
                ], alternate=(row % 2 == 0))
                format_currency_cell(ws.cell(row=row, column=2))
                format_currency_cell(ws.cell(row=row, column=5))
                row += 1

            row += 1
            ws.cell(row=row, column=1, value="Total Potential Acceleration").font = FONT_BODY_BOLD
            c = ws.cell(row=row, column=5, value=float(cs["total_potential_acceleration"]))
            format_currency_cell(c)
            c.font = FONT_BODY_BOLD
        auto_fit_columns(ws)

    def _create_roadmap(self):
        ws = self._wb.create_sheet("Phase 2 Roadmap")
        apply_title(ws, 1, 1, 4, "Phase 2 Recommendations")
        row = 3

        headers = ["Priority", "Area", "Description", "Data Needed"]
        for i, h in enumerate(headers, 1):
            ws.cell(row=row, column=i, value=h)
        apply_header_row(ws, row, 1, len(headers))
        row += 1

        for item in self._r["phase2_roadmap"]:
            write_data_row(ws, row, [
                item["priority"], item["area"], item["description"], item["data_needed"],
            ], alternate=(row % 2 == 0))
            row += 1
        auto_fit_columns(ws)
