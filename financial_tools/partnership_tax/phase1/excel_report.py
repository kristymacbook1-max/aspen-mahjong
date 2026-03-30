"""Phase 1 Partnership Tax Excel Report Generator.

Produces a 7-tab workbook:
  1. Executive Summary
  2. Allocation Analysis
  3. §704(c) Contributed Property
  4. Basis Adjustments
  5. Distributions
  6. Liability Allocations
  7. Phase 2 Roadmap
"""

import os, sys
from decimal import Decimal
from openpyxl import Workbook

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', '..'))
from revenue_recognition.utils.excel_styles import (
    FONT_TITLE, FONT_BODY, FONT_BODY_BOLD, THIN_BORDER, FMT_CURRENCY,
    apply_header_row, apply_section_header, apply_title,
    auto_fit_columns, write_data_row, format_currency_cell, get_risk_font,
)
from revenue_recognition.utils.currency_helpers import to_decimal, format_currency


class Phase1PartnershipReport:

    def generate(self, results: dict, output_path: str) -> str:
        self._r = results
        self._wb = Workbook()
        self._wb.remove(self._wb.active)

        self._create_executive_summary()
        self._create_allocations()
        self._create_704c()
        self._create_basis_adjustments()
        self._create_distributions()
        self._create_liabilities()
        self._create_roadmap()

        os.makedirs(os.path.dirname(output_path) or '.', exist_ok=True)
        self._wb.save(output_path)
        return output_path

    def _create_executive_summary(self):
        ws = self._wb.create_sheet("Executive Summary")
        apply_title(ws, 1, 1, 5, f"Partnership Tax Analysis — {self._r['partnership'].name}")
        row = 3

        p = self._r["partnership"]
        alloc = self._r["allocation_analysis"]
        sec704c = self._r["section_704c_analysis"]
        basis = self._r["basis_adjustment_analysis"]
        info = [
            ("Partnership", p.name),
            ("Tax Year", str(p.tax_year)),
            ("§754 Election", "Yes" if p.has_754_election else "No"),
            ("BBA Partnership", "Yes" if p.is_bba_partnership else "No"),
            ("Partners", str(p.number_of_partners)),
            ("Total Allocations", format_currency(alloc["total_allocations"])),
            ("Special Allocations", str(alloc["special_allocation_count"])),
            ("§704(c) Properties", str(sec704c["count"])),
            ("Net Built-In Gain/(Loss)", format_currency(sec704c["net_builtin"])),
            ("Book-Tax Asset Difference", format_currency(basis["book_tax_difference"])),
        ]
        for label, val in info:
            ws.cell(row=row, column=1, value=label).font = FONT_BODY_BOLD
            ws.cell(row=row, column=2, value=val).font = FONT_BODY
            row += 1

        row += 1
        apply_section_header(ws, row, 1, 5, "Key Findings")
        row += 1
        for f in self._r["findings"]:
            ws.cell(row=row, column=1, value=f["area"]).font = FONT_BODY_BOLD
            ws.cell(row=row, column=2, value=f["description"]).font = FONT_BODY
            ws.cell(row=row, column=3, value=f["risk_level"].upper()).font = get_risk_font(f["risk_level"])
            row += 1
        auto_fit_columns(ws)

    def _create_allocations(self):
        ws = self._wb.create_sheet("Allocation Analysis")
        apply_title(ws, 1, 1, 5, "§704(b) Partnership Allocations")
        row = 3

        headers = ["Item", "Amount", "Type", "Special Allocation", "SEE Test Note"]
        for i, h in enumerate(headers, 1):
            ws.cell(row=row, column=i, value=h)
        apply_header_row(ws, row, 1, len(headers))
        row += 1

        for item in self._r["allocation_analysis"]["items"]:
            write_data_row(ws, row, [
                item["description"],
                float(to_decimal(item["amount"])),
                item["type"],
                "Yes" if item["has_special_allocation"] else "No",
                item["see_test"],
            ], alternate=(row % 2 == 0))
            format_currency_cell(ws.cell(row=row, column=2))
            row += 1
        auto_fit_columns(ws)

    def _create_704c(self):
        ws = self._wb.create_sheet("§704(c) Property")
        apply_title(ws, 1, 1, 7, "§704(c) Contributed Property Analysis")
        row = 3

        items = self._r["section_704c_analysis"]["items"]
        if not items:
            ws.cell(row=row, column=1, value="No §704(c) contributed properties identified.").font = FONT_BODY
        else:
            headers = ["Property", "Contributing Partner", "FMV at Contribution",
                        "Tax Basis", "Built-In Gain/(Loss)", "Method", "Ceiling Rule Risk"]
            for i, h in enumerate(headers, 1):
                ws.cell(row=row, column=i, value=h)
            apply_header_row(ws, row, 1, len(headers))
            row += 1

            for item in items:
                write_data_row(ws, row, [
                    item["description"], item["contributing_partner"],
                    float(to_decimal(item["fmv_at_contribution"])),
                    float(to_decimal(item["tax_basis"])),
                    float(to_decimal(item["initial_builtin"])),
                    item["method"],
                    "Yes" if item["ceiling_rule_risk"] else "No",
                ], alternate=(row % 2 == 0))
                for c in [3, 4, 5]:
                    format_currency_cell(ws.cell(row=row, column=c))
                row += 1

            row += 1
            sec = self._r["section_704c_analysis"]
            ws.cell(row=row, column=1, value="Total Built-In Gain").font = FONT_BODY_BOLD
            c = ws.cell(row=row, column=3, value=float(to_decimal(sec["total_builtin_gain"])))
            format_currency_cell(c)
            c.font = FONT_BODY_BOLD
            row += 1
            ws.cell(row=row, column=1, value="Total Built-In Loss").font = FONT_BODY_BOLD
            c = ws.cell(row=row, column=3, value=float(to_decimal(sec["total_builtin_loss"])))
            format_currency_cell(c)
            c.font = FONT_BODY_BOLD
        auto_fit_columns(ws)

    def _create_basis_adjustments(self):
        ws = self._wb.create_sheet("Basis Adjustments")
        apply_title(ws, 1, 1, 3, "§754 / §743(b) / §734(b) Basis Adjustment Analysis")
        row = 3

        ba = self._r["basis_adjustment_analysis"]
        lines = [
            ("§754 Election in Effect", "Yes" if ba["has_754_election"] else "No"),
            ("Book Value of Assets", ba["book_assets"]),
            ("Tax Basis of Assets", ba["tax_assets"]),
            ("Book-Tax Difference", ba["book_tax_difference"]),
            ("Substantial Built-In Loss (>$250K)", "Yes" if ba["substantial_built_in_loss"] else "No"),
        ]
        for label, val in lines:
            ws.cell(row=row, column=1, value=label).font = FONT_BODY_BOLD
            if isinstance(val, str):
                ws.cell(row=row, column=2, value=val).font = FONT_BODY
            else:
                c = ws.cell(row=row, column=2, value=float(to_decimal(val)))
                format_currency_cell(c)
            row += 1

        row += 1
        ws.cell(row=row, column=1, value="Recommendation").font = FONT_BODY_BOLD
        ws.cell(row=row, column=2, value=ba["recommendation"]).font = FONT_BODY
        if ba["mandatory_743b_note"]:
            row += 1
            ws.cell(row=row, column=1, value="WARNING").font = get_risk_font("high")
            ws.cell(row=row, column=2, value=ba["mandatory_743b_note"]).font = FONT_BODY
        auto_fit_columns(ws)

    def _create_distributions(self):
        ws = self._wb.create_sheet("Distributions")
        apply_title(ws, 1, 1, 7, "Partnership Distributions Analysis")
        row = 3

        items = self._r["distribution_analysis"]["items"]
        if not items:
            ws.cell(row=row, column=1, value="No distributions identified.").font = FONT_BODY
        else:
            headers = ["Partner", "Type", "Cash", "Property FMV", "Property Basis",
                        "Outside Basis", "Gain Recognized"]
            for i, h in enumerate(headers, 1):
                ws.cell(row=row, column=i, value=h)
            apply_header_row(ws, row, 1, len(headers))
            row += 1

            for item in items:
                write_data_row(ws, row, [
                    item["partner"], item["type"],
                    float(to_decimal(item["cash"])),
                    float(to_decimal(item["property_fmv"])),
                    float(to_decimal(item["property_basis"])),
                    float(to_decimal(item["partner_outside_basis"])),
                    float(to_decimal(item["gain_recognized"])),
                ], alternate=(row % 2 == 0))
                for c in [3, 4, 5, 6, 7]:
                    format_currency_cell(ws.cell(row=row, column=c))
                row += 1
        auto_fit_columns(ws)

    def _create_liabilities(self):
        ws = self._wb.create_sheet("Liability Allocations")
        apply_title(ws, 1, 1, 5, "§752 Liability Allocations")
        row = 3

        liab = self._r["liability_analysis"]
        ws.cell(row=row, column=1, value="Total Liabilities").font = FONT_BODY_BOLD
        c = ws.cell(row=row, column=2, value=float(to_decimal(liab["total_liabilities"])))
        format_currency_cell(c)
        row += 1
        ws.cell(row=row, column=1, value="Recourse").font = FONT_BODY_BOLD
        c = ws.cell(row=row, column=2, value=float(to_decimal(liab["recourse"])))
        format_currency_cell(c)
        row += 1
        ws.cell(row=row, column=1, value="Nonrecourse").font = FONT_BODY_BOLD
        c = ws.cell(row=row, column=2, value=float(to_decimal(liab["nonrecourse"])))
        format_currency_cell(c)
        row += 2

        if liab["partner_allocations"]:
            headers = ["Partner", "Recourse", "Nonrecourse", "Qual. Nonrecourse", "Total"]
            for i, h in enumerate(headers, 1):
                ws.cell(row=row, column=i, value=h)
            apply_header_row(ws, row, 1, len(headers))
            row += 1

            for pa in liab["partner_allocations"]:
                write_data_row(ws, row, [
                    pa["partner"],
                    float(to_decimal(pa["recourse"])),
                    float(to_decimal(pa["nonrecourse"])),
                    float(to_decimal(pa["qualified_nonrecourse"])),
                    float(to_decimal(pa["total"])),
                ], alternate=(row % 2 == 0))
                for c in [2, 3, 4, 5]:
                    format_currency_cell(ws.cell(row=row, column=c))
                row += 1

        row += 1
        ws.cell(row=row, column=1, value=liab["note"]).font = FONT_BODY
        auto_fit_columns(ws)

    def _create_roadmap(self):
        ws = self._wb.create_sheet("Phase 2 Roadmap")
        apply_title(ws, 1, 1, 4, "Phase 2 Recommendations")
        row = 3

        headers = ["Priority", "Area", "Description", "Data Needed"]
        for i, h in enumerate(headers, 1):
            ws.cell(row=row, column=i, value=h)
        apply_header_row(ws, row, 1, len(headers))
        row += 1

        for item in self._r["phase2_roadmap"]:
            write_data_row(ws, row, [
                item["priority"], item["area"], item["description"], item["data_needed"],
            ], alternate=(row % 2 == 0))
            row += 1
        auto_fit_columns(ws)
