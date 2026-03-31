"""Phase 1 CAMT Excel Report Generator.

Produces a 6-tab workbook:
  1. Executive Summary
  2. Applicable Corporation Test
  3. AFSI Computation
  4. CAMT Computation
  5. AFSI Adjustments
  6. Phase 2 Roadmap
"""

import os
import sys
from decimal import Decimal
from openpyxl import Workbook

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', '..'))
from revenue_recognition.utils.excel_styles import (
    FONT_TITLE, FONT_BODY, FONT_BODY_BOLD, FILL_HIGHLIGHT_GREEN,
    FILL_HIGHLIGHT_RED, THIN_BORDER, FMT_CURRENCY,
    apply_header_row, apply_section_header, apply_title,
    auto_fit_columns, write_data_row, format_currency_cell,
    get_risk_font,
)
from revenue_recognition.utils.currency_helpers import to_decimal, format_currency, format_millions


class Phase1CAMTReport:

    def generate(self, results: dict, output_path: str) -> str:
        self._r = results
        self._wb = Workbook()
        self._wb.remove(self._wb.active)

        self._create_executive_summary()
        self._create_threshold_test()
        self._create_afsi_computation()
        self._create_camt_computation()
        self._create_adjustments()
        self._create_roadmap()

        os.makedirs(os.path.dirname(output_path) or '.', exist_ok=True)
        self._wb.save(output_path)
        return output_path

    def _create_executive_summary(self):
        ws = self._wb.create_sheet("Executive Summary")
        company = self._r["company"]
        apply_title(ws, 1, 1, 5, f"CAMT Analysis — {company.name}")
        row = 3

        thresh = self._r["threshold_test"]
        comp = self._r["camt_computation"]

        info = [
            ("Company", company.name),
            ("Tax Year", str(company.tax_year)),
            ("3-Year Average AFSI", format_millions(thresh.three_year_average)),
            ("Applicable Corporation?", "YES" if thresh.is_applicable_corporation else "NO"),
            ("Estimated CAMT Liability", format_currency(comp.camt_liability)),
        ]
        for label, val in info:
            ws.cell(row=row, column=1, value=label).font = FONT_BODY_BOLD
            ws.cell(row=row, column=2, value=val).font = FONT_BODY
            row += 1

        row += 1
        apply_section_header(ws, row, 1, 5, "Key Findings")
        row += 1
        for f in self._r["findings"]:
            ws.cell(row=row, column=1, value=f["area"]).font = FONT_BODY_BOLD
            ws.cell(row=row, column=2, value=f["description"]).font = FONT_BODY
            ws.cell(row=row, column=3, value=f["risk_level"].upper()).font = get_risk_font(f["risk_level"])
            row += 1

        auto_fit_columns(ws)

    def _create_threshold_test(self):
        ws = self._wb.create_sheet("Applicable Corp Test")
        apply_title(ws, 1, 1, 3, "§59(k) Applicable Corporation Threshold Test")
        row = 3

        thresh = self._r["threshold_test"]
        lines = [
            ("Year 1 AFSI", thresh.year1_afsi),
            ("Year 2 AFSI", thresh.year2_afsi),
            ("Year 3 AFSI", thresh.year3_afsi),
            ("3-Year Average", thresh.three_year_average),
            ("Threshold", thresh.threshold),
        ]
        for label, val in lines:
            ws.cell(row=row, column=1, value=label).font = FONT_BODY
            c = ws.cell(row=row, column=2, value=float(to_decimal(val)))
            format_currency_cell(c)
            ws.cell(row=row, column=1).border = THIN_BORDER
            c.border = THIN_BORDER
            row += 1

        row += 1
        result = "APPLICABLE CORPORATION — CAMT APPLIES" if thresh.is_applicable_corporation else "NOT an applicable corporation — CAMT does not apply"
        ws.cell(row=row, column=1, value=result).font = get_risk_font("high" if thresh.is_applicable_corporation else "low")
        auto_fit_columns(ws)

    def _create_afsi_computation(self):
        ws = self._wb.create_sheet("AFSI Computation")
        afsi = self._r["afsi_computation"]
        apply_title(ws, 1, 1, 3, f"Adjusted Financial Statement Income — {afsi['year']}")
        row = 3

        lines = [("Pretax Book Income", afsi["pretax_book_income"])]
        for k, v in afsi["adjustments"].items():
            lines.append((f"Adjustment: {k.replace('_', ' ').title()}", v))
        lines.append(("Total Adjustments", afsi["total_adjustments"]))
        lines.append(("AFSI", afsi["afsi"]))

        for label, val in lines:
            ws.cell(row=row, column=1, value=label).font = FONT_BODY if "AFSI" not in label else FONT_BODY_BOLD
            c = ws.cell(row=row, column=2, value=float(to_decimal(val)))
            format_currency_cell(c)
            ws.cell(row=row, column=1).border = THIN_BORDER
            c.border = THIN_BORDER
            row += 1

        auto_fit_columns(ws)

    def _create_camt_computation(self):
        ws = self._wb.create_sheet("CAMT Computation")
        apply_title(ws, 1, 1, 3, "Corporate Alternative Minimum Tax Computation")
        row = 3

        comp = self._r["camt_computation"]
        lines = [
            ("Adjusted Financial Statement Income (AFSI)", comp.afsi),
            ("× 15% CAMT Rate", ""),
            ("Tentative Minimum Tax (net of CAMT FTC)", comp.tentative_minimum_tax),
            ("Less: Regular Tax + BEAT", comp.regular_tax_plus_base_erosion),
            ("CAMT Liability", comp.camt_liability),
        ]
        for label, val in lines:
            ws.cell(row=row, column=1, value=label).font = FONT_BODY_BOLD if "CAMT Liability" in label else FONT_BODY
            if isinstance(val, str):
                ws.cell(row=row, column=2, value=val).font = FONT_BODY
            else:
                c = ws.cell(row=row, column=2, value=float(to_decimal(val)))
                format_currency_cell(c)
            ws.cell(row=row, column=1).border = THIN_BORDER
            row += 1

        auto_fit_columns(ws)

    def _create_adjustments(self):
        ws = self._wb.create_sheet("AFSI Adjustments")
        apply_title(ws, 1, 1, 6, "AFSI Adjustment Detail")
        row = 3

        headers = ["Description", "Book Amount", "Adjustment", "Adjusted Amount", "IRC Section", "Impact"]
        for i, h in enumerate(headers, 1):
            ws.cell(row=row, column=i, value=h)
        apply_header_row(ws, row, 1, len(headers))
        row += 1

        for adj in self._r["adjustments"]:
            write_data_row(ws, row, [
                adj["description"], float(to_decimal(adj["book_amount"])),
                float(to_decimal(adj["tax_adjustment"])), float(to_decimal(adj["adjusted_amount"])),
                adj["irc_section"], adj["impact"],
            ], alternate=(row % 2 == 0))
            format_currency_cell(ws.cell(row=row, column=2))
            format_currency_cell(ws.cell(row=row, column=3))
            format_currency_cell(ws.cell(row=row, column=4))
            row += 1

        auto_fit_columns(ws)

    def _create_roadmap(self):
        ws = self._wb.create_sheet("Phase 2 Roadmap")
        apply_title(ws, 1, 1, 4, "Phase 2 Recommendations")
        row = 3

        headers = ["Priority", "Area", "Description", "Data Needed"]
        for i, h in enumerate(headers, 1):
            ws.cell(row=row, column=i, value=h)
        apply_header_row(ws, row, 1, len(headers))
        row += 1

        for item in self._r["phase2_roadmap"]:
            write_data_row(ws, row, [
                item["priority"], item["area"], item["description"], item["data_needed"],
            ], alternate=(row % 2 == 0))
            row += 1

        auto_fit_columns(ws)
