"""Trial-balance reader — the single, robust ingestion path.

Alias-based header detection + section-header skipping + amount parsing
(promoted from the original run_263a_classifier.py, which was the most complete
of the three readers the tool shipped). Crucially, it carries dollar AMOUNTS
into TBLine — the gap that made the original a labeler rather than a calc.
"""

import re
import warnings
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook
from .model import TBLine

_ALIASES = {
    "acct_num": ["account number", "acct number", "acct #", "acct no", "account no",
                 "gl account", "account #", "gl #", "account"],
    "acct_desc": ["account description", "acct desc", "acct description", "account desc",
                  "gl description", "gl desc", "description"],
    "cc_num": ["cost center number", "cc number", "cc #", "cc", "dept number",
               "dept. number", "dept #", "dept no", "department number",
               "cost center", "cost center #"],
    "cc_desc": ["cost center description", "cc description", "cc desc",
                "department description", "dept description", "dept desc",
                "department", "dept name", "cost center name"],
    # A single net/amount column is preferred. "debit" is NOT treated as a
    # standalone amount — on a two-column (Debit/Credit) TB that would zero out
    # credit-only balances; debit and credit are captured separately and netted.
    "amount": ["amount", "net balance", "net", "total book amount", "balance",
               "net amount", "ending balance", "amount (debit / <credit>)"],
    "debit": ["debit", "debit amount", "dr"],
    "credit": ["credit", "credit amount", "cr"],
}
_SECTION_HEADERS = {"assets", "liabilities", "equity", "revenue", "expenses",
                    "income", "cost of goods sold", "cogs"}
_TB_SHEET_HINTS = ["raw tb", "tb", "trial balance", "cy_trial_balance",
                   "tb import & classification"]
_EXCLUDED_SHEET_TITLES = ("instructions", "classification results",
                          "classification summary", "cost code reference")
# Cached formula-error literals (openpyxl data_only=True returns these as plain
# strings when a formula is broken) — must never be treated as account text.
_FORMULA_ERRORS = {"#ref!", "#n/a", "#div/0!", "#name?", "#null!", "#num!", "#value!"}
_HEADER_SCAN_ROWS = 100   # real ERP exports (SAP/Oracle/NetSuite) commonly have
                          # 20-40 rows of preamble before the header row


def _to_decimal(v):
    if v is None:
        return Decimal("0")
    if isinstance(v, (int, float)):
        return Decimal(str(v))
    s = str(v).strip().replace(",", "").replace("$", "").strip()
    neg = s.startswith("(") and s.endswith(")")
    s = s.strip("()")
    if s in ("", "-"):
        return Decimal("0")
    try:
        d = Decimal(s)
    except InvalidOperation:
        return Decimal("0")
    return -d if neg else d


def _looks_like_tb(ws):
    """True if this sheet has the minimum columns a trial balance needs
    (an account-description column plus an amount or debit/credit column) —
    used to rule out title/cover/notes sheets before the ambiguity check."""
    header_row = _detect_header(ws)
    cols = _map_columns(ws, header_row)
    return "acct_desc" in cols and ("amount" in cols or "debit" in cols or "credit" in cols)


def _pick_sheet(wb):
    for hint in _TB_SHEET_HINTS:
        for ws in wb.worksheets:
            if ws.title.strip().lower() == hint:
                return ws
    candidates = [ws for ws in wb.worksheets
                  if ws.sheet_state == "visible" and not ws.title.startswith("_")
                  and ws.title.lower() not in _EXCLUDED_SHEET_TITLES]
    # A title/cover/notes sheet with no TB-shaped columns shouldn't count toward
    # ambiguity — only raise when 2+ candidates actually look like a trial balance.
    tb_shaped = [ws for ws in candidates if _looks_like_tb(ws)]
    if len(tb_shaped) > 1:
        titles = ", ".join(repr(ws.title) for ws in tb_shaped)
        raise ValueError(
            f"Multiple candidate sheets ({titles}) and none matches a known trial-balance "
            f"sheet name — pass sheet=<name> explicitly to avoid picking the wrong one.")
    if tb_shaped:
        return tb_shaped[0]
    if candidates:
        return candidates[0]
    return wb.worksheets[0]


def _detect_header(ws):
    best_row, best_score = 1, 0
    all_aliases = {a for v in _ALIASES.values() for a in v}
    for r in range(1, min(ws.max_row, _HEADER_SCAN_ROWS) + 1):
        score = 0
        for c in range(1, min(ws.max_column, 30) + 1):
            v = str(ws.cell(r, c).value or "").strip().lower()
            if v in all_aliases:
                score += 1
        if score > best_score:
            best_row, best_score = r, score
    return best_row


def _map_columns(ws, header_row):
    cols = {}
    for c in range(1, min(ws.max_column, 40) + 1):
        v = str(ws.cell(header_row, c).value or "").strip().lower()
        if not v:
            continue
        for field, aliases in _ALIASES.items():
            if field in cols:
                continue
            if v in aliases:
                cols[field] = c
                break
    return cols


def read_trial_balance(path, sheet=None):
    wb = load_workbook(path, data_only=True)
    ws = wb[sheet] if sheet else _pick_sheet(wb)
    header_row = _detect_header(ws)
    cols = _map_columns(ws, header_row)
    if "acct_desc" not in cols:
        raise ValueError(f"Could not find an account-description column on '{ws.title}' "
                         f"(header row {header_row}).")

    has_amount = "amount" in cols
    has_debit_credit = "debit" in cols or "credit" in cols
    if not has_amount and not has_debit_credit:
        raise ValueError(
            f"Could not find an amount, debit, or credit column on '{ws.title}' "
            f"(header row {header_row}) — every line would silently read as $0.")
    if "acct_num" not in cols:
        warnings.warn(
            f"No account-number column detected on '{ws.title}' (header row {header_row}); "
            f"acct_num will be blank for every line.", stacklevel=2)

    lines = []
    for r in range(header_row + 1, ws.max_row + 1):
        def cell(field):
            ci = cols.get(field)
            return ws.cell(r, ci).value if ci else None
        desc = str(cell("acct_desc") or "").strip()
        if (not desc or desc.lower() in _SECTION_HEADERS or desc == "0"
                or desc.lower() in _FORMULA_ERRORS):
            continue
        if has_amount:
            amount = _to_decimal(cell("amount"))
        else:
            amount = _to_decimal(cell("debit")) - _to_decimal(cell("credit"))
        lines.append(TBLine(
            acct_num=str(cell("acct_num") or "").strip(),
            acct_desc=desc,
            cc_num=str(cell("cc_num") or "").strip(),
            cc_desc=str(cell("cc_desc") or "").strip(),
            amount=amount,
            row_index=r,
        ))
    return lines
