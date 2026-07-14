"""Report generation: must not crash on edge-case inputs."""

import os
from decimal import Decimal

from openpyxl import load_workbook

from financial_tools.cap263a.model import TBLine
from financial_tools.cap263a.analysis import analyze, EntityProfile
from financial_tools.cap263a.report import CapitalizationReport


def test_empty_trial_balance_does_not_crash(tmp_path):
    """A zero-row TB used to invert the conditional-formatting range
    (O4:O3) and crash workbook generation entirely."""
    r = analyze([], EntityProfile())
    path = os.path.join(tmp_path, "empty.xlsx")
    out = CapitalizationReport().generate(r, path)
    wb = load_workbook(out)
    assert "Summary Dashboard" in wb.sheetnames
    assert "Classified TB" in wb.sheetnames


def test_normal_trial_balance_generates_five_tabs(tmp_path):
    lines = [
        TBLine("5000", "Direct labor", "100", "Production", amount=Decimal("1000000")),
        TBLine("7000", "Advertising", "400", "Marketing", amount=Decimal("200000")),
    ]
    r = analyze(lines, EntityProfile(avg_gross_receipts=Decimal("75000000")))
    path = os.path.join(tmp_path, "tb.xlsx")
    out = CapitalizationReport().generate(r, path)
    wb = load_workbook(out)
    # Full ordered list, not just a subset/index-0 check — a set-based
    # assertion here wouldn't catch tabs 2-5 silently reordering.
    assert wb.sheetnames == ["Summary Dashboard", "Classified TB", "Asset Basis Schedule",
                             "Adjusted IS", "Method Changes"]


def test_review_workflow_surfaces(tmp_path):
    """The Classified TB must be workable as a review queue: autofilter,
    frozen header, provenance columns, and a bucket-name data validation
    (a typo'd bucket silently drops the line from the Summary SUMIFS)."""
    lines = [
        TBLine("5000", "Direct labor", "100", "Production", amount=Decimal("1000000")),
        TBLine("7000", "Advertising", "400", "Marketing", amount=Decimal("200000")),
    ]
    r = analyze(lines, EntityProfile(avg_gross_receipts=Decimal("75000000")))
    path = os.path.join(tmp_path, "tb.xlsx")
    wb = load_workbook(CapitalizationReport().generate(r, path))
    tb = wb["Classified TB"]
    assert tb.auto_filter.ref and tb.auto_filter.ref.startswith("A3:")
    assert tb.freeze_panes == "A4"
    assert tb.cell(3, 17).value == "Method (why)"
    assert tb.cell(3, 18).value == "Zone"
    assert tb.cell(4, 17).value          # provenance actually populated
    assert len(tb.data_validations.dataValidation) == 1


def test_no_label_stored_as_broken_formula(tmp_path):
    """A label written as "= Adjusted currently-deductible" was stored by
    openpyxl as a formula and rendered #NAME? in Excel."""
    r = analyze([], EntityProfile())
    path = os.path.join(tmp_path, "empty.xlsx")
    wb = load_workbook(CapitalizationReport().generate(r, path))
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.data_type == "f":   # formula cells must look like formulas
                    assert str(c.value).startswith("="), (ws.title, c.coordinate)
                    assert not str(c.value).startswith("= "), (ws.title, c.coordinate)


def test_asset_basis_includes_interest_and_266_and_exempt_stub_is_zero(tmp_path):
    """The Asset Basis 'Total basis additions' hardcoded §263A(f) to 0 and had
    no §266 row at all — understating total additions while Method Changes
    showed the amounts. And the APE×rate interest stub ignored the §263A(i)
    exemption entirely."""
    lines = [TBLine("5000", "Direct labor", "100", "Production", amount=Decimal("100000"))]
    p = EntityProfile(avg_gross_receipts=Decimal("75000000"),
                      accumulated_production_expenditures=Decimal("1000000"),
                      avoided_cost_rate=Decimal("0.05"), has_designated_property=True)
    r = analyze(lines, p)
    wb = load_workbook(CapitalizationReport().generate(r, os.path.join(tmp_path, "a.xlsx")))
    ab = wb["Asset Basis Schedule"]
    labels = {str(ab.cell(row, 1).value or ""): ab.cell(row, 2).value
              for row in range(5, 15)}
    interest_rows = [v for k, v in labels.items() if "§263A(f)" in k]
    sec266_rows = [v for k, v in labels.items() if "§266" in k]
    assert interest_rows and interest_rows[0] == 50000.0     # APE 1M × 5%
    assert sec266_rows != []                                 # §266 row exists

    # exempt entity: the stub must be zero even with designated property
    p2 = EntityProfile(avg_gross_receipts=Decimal("1000000"),
                       accumulated_production_expenditures=Decimal("1000000"),
                       avoided_cost_rate=Decimal("0.05"), has_designated_property=True)
    r2 = analyze(lines, p2)
    wb2 = load_workbook(CapitalizationReport().generate(r2, os.path.join(tmp_path, "b.xlsx")))
    mc = wb2["Method Changes"]
    vals = {str(mc.cell(row, 1).value or ""): mc.cell(row, 3).value for row in range(5, 14)}
    stub = [v for k, v in vals.items() if k.startswith("Interest capitalized")]
    assert stub and stub[0] == 0.0


def test_warnings_render_on_summary(tmp_path):
    """UNICAP computation warnings (unimplemented method, stale threshold,
    negative pools) must be visible in the workpaper itself, not just stderr."""
    lines = [TBLine("5000", "Direct labor", "100", "Production", amount=Decimal("100000"))]
    r = analyze(lines, EntityProfile(avg_gross_receipts=Decimal("75000000"),
                                     method="FACTS", tax_year=2027))
    path = os.path.join(tmp_path, "warn.xlsx")
    wb = load_workbook(CapitalizationReport().generate(r, path))
    summ = wb["Summary Dashboard"]
    cautions = [c.value for row in summ.iter_rows() for c in row
                if isinstance(c.value, str) and c.value.startswith("⚠")]
    assert any("FACTS" in c for c in cautions)
    assert any("2027" in c for c in cautions)


def _label_values(ws, col=3):
    """{label text (col A, or col B if A is blank): value in `col`} for every
    labeled row, so tests survive row-number reshuffling."""
    out = {}
    for row in ws.iter_rows():
        label = row[0].value or row[1].value
        if isinstance(label, str) and label.strip():
            out[label.strip()] = row[col - 1].value
    return out


def test_summary_waterfall_ties_to_adjusted_is_with_mixed_service(tmp_path):
    """Full-pipeline regression: a synthetic engagement with a partial (< 1)
    SSCM mixed-service ratio used to make the Summary Dashboard and the
    Adjusted IS tab report TWO DIFFERENT "final deductible" numbers for the
    same workbook, off by exactly mixed_deductible, with no warning anywhere
    — "Total capitalized" and Asset Basis Schedule's "Total basis additions"
    were both missing mixed_capitalized entirely (confirmed: $300,000/
    $100,000 apart on these facts). Total capitalized + Adjusted
    currently-deductible must equal the starting IS total, and must match
    the independently-computed Adjusted IS total."""
    lines = [
        TBLine("5000", "Direct labor", "100", "Production", amount=Decimal("1000000")),
        TBLine("5100", "Direct materials", "100", "Production", amount=Decimal("500000")),
        TBLine("6100", "Purchasing dept salaries", "150", "Purchasing", amount=Decimal("300000")),
        TBLine("6200", "Human resources - plant", "160", "HR", amount=Decimal("400000")),
        TBLine("7000", "Advertising", "400", "Marketing", amount=Decimal("200000")),
        TBLine("7100", "Sales commissions", "400", "Sales", amount=Decimal("600000")),
    ]
    p = EntityProfile(avg_gross_receipts=Decimal("75000000"), method="SPM",
                      ending_inventory_471=Decimal("200000"))
    r = analyze(lines, p)
    u = r["unicap"]
    assert Decimal("0") < u["mixed_alloc_ratio"] < Decimal("1")   # a real split, not 0/100%
    assert u["mixed_capitalized"] > 0 and u["mixed_deductible"] > 0

    path = os.path.join(tmp_path, "mixed.xlsx")
    wb = load_workbook(CapitalizationReport().generate(r, path))
    summ = _label_values(wb["Summary Dashboard"])
    ab = _label_values(wb["Asset Basis Schedule"], col=2)
    adj_is = _label_values(wb["Adjusted IS"], col=4)

    cap_bucket_sum = sum(r["bucket_totals"][b] for b in
        ["Inventory §471", "§263A Additional", "§263(a) Mandatory",
         "§263(a) Elective", "§263A(f) Interest", "§266 Carrying"])
    ded_bucket_sum = r["bucket_totals"]["Deductible"] + r["bucket_totals"]["Non-Operating"]
    true_total_capitalized = cap_bucket_sum + u["mixed_capitalized"]
    true_adjusted_deductible = ded_bucket_sum + u["mixed_deductible"]
    assert true_total_capitalized + true_adjusted_deductible == r["is_total"]

    assert summ["Mixed service — SSCM capitalized share"] == float(u["mixed_capitalized"])
    assert summ["Mixed service — SSCM remaining deductible"] == float(u["mixed_deductible"])
    # "Total capitalized"/"Adjusted currently-deductible" are live SUMIFS
    # formulas (openpyxl can't evaluate them without Excel/LibreOffice); pin
    # the fix by checking the mixed-capitalized row's number is inside the
    # "Total capitalized" SUM(...) range instead.
    ws = wb["Summary Dashboard"]
    mixed_cap_row = next(row[0].row for row in ws.iter_rows()
                         if row[0].value == "  Mixed service — SSCM capitalized share")
    total_cap_formula = next(row[2].value for row in ws.iter_rows()
                             if row[0].value == "Total capitalized")
    assert total_cap_formula.startswith("=SUM(")
    lo, hi = total_cap_formula[5:-1].split(":")
    lo_row, hi_row = int(lo[1:]), int(hi[1:])
    assert lo_row <= mixed_cap_row <= hi_row

    assert ab["§263A additional cost — mixed-service capitalized (SSCM)"] == \
        float(u["mixed_capitalized"])
    assert adj_is["Mixed-service deductible remainder (post-SSCM)"] == \
        float(u["mixed_deductible"])
    # "TOTAL remaining deductible" is itself a =SUM(...) formula (openpyxl
    # can't evaluate it); sum the literal per-row Amount values it covers —
    # each is a plain float, not a formula — to get the value Excel would.
    is_ws = wb["Adjusted IS"]
    total_formula = next(row[3].value for row in is_ws.iter_rows()
                         if row[0].value == "TOTAL remaining deductible")
    lo, hi = total_formula[5:-1].split(":")
    lo_row, hi_row = int(lo[1:]), int(hi[1:])
    computed_total = sum(is_ws.cell(rr, 4).value for rr in range(lo_row, hi_row + 1))
    assert computed_total == float(true_adjusted_deductible)
