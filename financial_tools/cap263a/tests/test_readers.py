"""Multi-format engagement ingestion (readers.py)."""

import json
from datetime import date
from decimal import Decimal

from openpyxl import Workbook

from financial_tools.cap263a.readers import read_engagement


def _json_engagement(tmp_path, **overrides):
    doc = {
        "trial_balance": [
            {"acct_num": "5000", "acct_desc": "Direct labor", "cc_num": "100",
             "cc_desc": "Production", "amount": "1000000"},
            {"acct_num": "7000", "acct_desc": "Advertising", "cc_num": "400",
             "cc_desc": "Marketing", "amount": "200000"},
        ],
        "btd": [
            {"acct_num": "5000", "cc_num": "100", "description": "Comp accrual",
             "adjustment": "-15000", "category": "timing"},
        ],
        "fixed_assets": [
            {"asset_id": "FA1", "description": "Plant building", "asset_type": "real",
             "cost": "2500000", "placed_in_service": "2020-06-15",
             "class_life": "39", "book_capitalized_interest": "10000"},
        ],
        "cip": [
            {"project_id": "P1", "description": "New line", "linked_asset_id": "FA1",
             "is_real_property": "yes", "production_start": "2026-01-01"},
        ],
        "cip_snapshots": [
            {"project_id": "P1", "measurement_date": "2026-03-31", "cumulative_ape": "3500000"},
            {"project_id": "P1", "measurement_date": "2026-06-30", "cumulative_ape": "5000000"},
        ],
        "debt": [
            {"debt_id": "L1", "description": "Construction loan", "principal": "3000000",
             "rate": "0.06", "interest_incurred": "180000", "traced_to": "P1"},
            {"debt_id": "L2", "description": "Revolver", "principal": "2800000",
             "interest_incurred": "200000"},
        ],
        "re": [
            {"re_id": "R1", "description": "Foreign lab", "amount": "150000",
             "domestic": "no", "tax_year": "2026"},
        ],
    }
    doc.update(overrides)
    p = tmp_path / "engagement.json"
    p.write_text(json.dumps(doc))
    return p


def test_json_engagement_reads_all_schedules(tmp_path):
    data = read_engagement(_json_engagement(tmp_path))
    assert len(data.tb_lines) == 2 and data.tb_lines[0].amount == Decimal("1000000")
    assert data.btds[0].adjustment == Decimal("-15000")
    fa = data.fixed_assets[0]
    assert fa.asset_id == "FA1" and fa.cost == Decimal("2500000")
    assert fa.placed_in_service == date(2020, 6, 15)
    assert fa.book_capitalized_interest == Decimal("10000")
    cip = data.cip_projects[0]
    assert cip.linked_asset_id == "FA1" and cip.is_real_property
    assert [s.cumulative_ape for s in cip.snapshots] == [Decimal("3500000"), Decimal("5000000")]
    traced = [d for d in data.debts if d.traced_to == "P1"]
    assert traced and traced[0].interest_incurred == Decimal("180000")
    assert data.re_expenditures[0].domestic is False
    assert data.validation.blocking == 0


def test_fk_violations_are_blocking_errors(tmp_path):
    p = _json_engagement(
        tmp_path,
        cip=[{"project_id": "P1", "linked_asset_id": "NOPE"}],
        debt=[{"debt_id": "L1", "traced_to": "GHOST", "principal": "1",
               "interest_incurred": "0"}])
    data = read_engagement(p)
    assert data.validation.blocking == 2
    joined = " ".join(data.validation.errors)
    assert "NOPE" in joined and "GHOST" in joined


def test_duplicate_pk_is_error(tmp_path):
    p = _json_engagement(tmp_path, fixed_assets=[
        {"asset_id": "FA1", "cost": "1"}, {"asset_id": "FA1", "cost": "2"}])
    data = read_engagement(p)
    assert any("Duplicate FixedAsset" in e for e in data.validation.errors)


def test_csv_schedule_with_alias_headers(tmp_path):
    csv_path = tmp_path / "assets.csv"
    csv_path.write_text(
        "Asset #,Asset Description,Property Type,Original Cost,PIS Date\n"
        "FA9,Warehouse,real,750000,03/01/2024\n")
    data = read_engagement(assets_path=csv_path)
    fa = data.fixed_assets[0]
    assert fa.asset_id == "FA9" and fa.cost == Decimal("750000")
    assert fa.placed_in_service == date(2024, 3, 1)


def test_xlsx_workbook_sheet_routing(tmp_path):
    wb = Workbook()
    tb = wb.active
    tb.title = "Raw TB"
    tb.append(["Account Number", "Account Description", "Cost Center", "Amount"])
    tb.append(["5000", "Direct labor", "100", 1000000])
    fa = wb.create_sheet("Fixed Asset Register")
    fa.append(["asset id", "description", "asset type", "cost"])
    fa.append(["FA1", "Plant", "real", 2500000])
    debt = wb.create_sheet("Debt Schedule")
    debt.append(["loan id", "principal", "interest incurred", "traced to"])
    debt.append(["L1", 3000000, 180000, ""])
    path = tmp_path / "eng.xlsx"
    wb.save(path)
    data = read_engagement(path)
    assert data.tb_lines and data.tb_lines[0].acct_desc == "Direct labor"
    assert data.fixed_assets[0].cost == Decimal("2500000")
    assert data.debts[0].principal == Decimal("3000000")
    assert data.validation.blocking == 0


def test_directory_of_schedule_files(tmp_path):
    (tmp_path / "fixed_assets.csv").write_text(
        "asset id,description,cost\nFA1,Plant,100\n")
    (tmp_path / "debt.csv").write_text(
        "debt id,principal,interest incurred\nL1,500,25\n")
    data = read_engagement(tmp_path)
    assert data.fixed_assets[0].asset_id == "FA1"
    assert data.debts[0].interest_incurred == Decimal("25")
