"""Reader: debit/credit netting (a two-column TB must not zero out credits)."""

import os
from decimal import Decimal

import pytest
from openpyxl import Workbook
from financial_tools.cap263a.reader import read_trial_balance, _to_decimal


def test_debit_credit_netting(tmp_path):
    wb = Workbook(); ws = wb.active; ws.title = "TB"
    ws.append(["Account Number", "Account Description", "Cost Center", "Debit", "Credit"])
    ws.append(["4000", "Sales revenue", "Sales", "", 500000])     # credit-only
    ws.append(["5000", "Direct labor", "Plant", 300000, ""])      # debit-only
    ws.append(["6000", "Net line", "Plant", 100000, 40000])       # both
    p = os.path.join(tmp_path, "dc.xlsx"); wb.save(p)
    lines = read_trial_balance(p)
    amounts = {l.acct_desc: l.amount for l in lines}
    assert amounts["Sales revenue"] == Decimal("-500000")   # not zero
    assert amounts["Direct labor"] == Decimal("300000")
    assert amounts["Net line"] == Decimal("60000")          # debit - credit


def test_single_amount_column_still_works(tmp_path):
    wb = Workbook(); ws = wb.active; ws.title = "TB"
    ws.append(["Account Number", "Account Description", "Amount"])
    ws.append(["5000", "Direct labor", 300000])
    p = os.path.join(tmp_path, "amt.xlsx"); wb.save(p)
    lines = read_trial_balance(p)
    assert lines[0].amount == Decimal("300000")


def test_no_amount_or_debit_credit_column_raises(tmp_path):
    """A TB with only account/description columns must error, not silently
    produce $0 for every line."""
    wb = Workbook(); ws = wb.active; ws.title = "TB"
    ws.append(["Account Number", "Account Description"])
    ws.append(["5000", "Direct labor"])
    p = os.path.join(tmp_path, "noamt.xlsx"); wb.save(p)
    with pytest.raises(ValueError, match="amount, debit, or credit"):
        read_trial_balance(p)


def test_formula_error_strings_are_skipped(tmp_path):
    """A broken formula cached as '#REF!'/'#N/A' must not become a phantom
    line item polluting totals."""
    wb = Workbook(); ws = wb.active; ws.title = "TB"
    ws.append(["Account Number", "Account Description", "Amount"])
    ws.append(["5000", "Direct labor", 300000])
    ws.append(["6000", "#REF!", 999999])
    ws.append(["7000", "#N/A", 888888])
    p = os.path.join(tmp_path, "err.xlsx"); wb.save(p)
    lines = read_trial_balance(p)
    assert len(lines) == 1
    assert lines[0].acct_desc == "Direct labor"


def test_ambiguous_sheets_require_explicit_choice(tmp_path):
    """Two equally-plausible sheets with no hint match must not silently pick
    the first one in workbook order — that risks reading a stale/wrong TB."""
    wb = Workbook()
    ws1 = wb.active; ws1.title = "2023 Data"
    ws1.append(["Account Number", "Account Description", "Amount"])
    ws1.append(["5000", "Direct labor", 300000])
    ws2 = wb.create_sheet("2024 Data")
    ws2.append(["Account Number", "Account Description", "Amount"])
    ws2.append(["5000", "Direct labor", 400000])
    p = os.path.join(tmp_path, "ambiguous.xlsx"); wb.save(p)
    with pytest.raises(ValueError, match="Multiple candidate sheets"):
        read_trial_balance(p)
    # explicit sheet= still works
    lines = read_trial_balance(p, sheet="2024 Data")
    assert lines[0].amount == Decimal("400000")


def test_accounting_negative_with_dollar_sign_and_space():
    """"$ (1,234.00)" (dollar sign, space, then parens) must parse as a real
    negative, not silently zero out."""
    assert _to_decimal("$ (1,234.00)") == Decimal("-1234.00")
    assert _to_decimal("$(1,234,567.89)") == Decimal("-1234567.89")


def test_nan_and_infinity_amounts_are_neutralized():
    """NaN/±Infinity are parseable by Decimal but poison every total, blank out
    in the workbook, and silently defeat the tie-check. They must read as 0."""
    for v in ["NaN", "nan", "Infinity", "-Infinity", "inf", float("nan"), float("inf")]:
        assert _to_decimal(v) == Decimal("0"), v


def test_nan_amount_does_not_defeat_tie_check(tmp_path):
    """End-to-end: a NaN amount cell must not make is_total/tie_check NaN."""
    import os
    from financial_tools.cap263a.analysis import analyze, EntityProfile
    wb = Workbook(); ws = wb.active; ws.title = "TB"
    ws.append(["Account Number", "Account Description", "Amount"])
    ws.append(["5000", "Direct labor", None])
    cell = ws.cell(2, 3); cell.value = "NaN"; cell.data_type = "s"
    p = os.path.join(tmp_path, "nan.xlsx"); wb.save(p)
    r = analyze(read_trial_balance(p), EntityProfile())
    assert r["tie_check"] == Decimal("0")
    assert r["is_total"] == Decimal("0")


def test_stray_far_cell_does_not_blow_up_row_scan(tmp_path):
    """A single cell at a huge row inflates ws.max_row to ~1M; the reader must
    stop after a long blank run rather than iterate to max_row."""
    import os, time
    wb = Workbook(); ws = wb.active; ws.title = "TB"
    ws.append(["Account Number", "Account Description", "Amount"])
    ws.append(["5000", "Direct labor", 100000])
    ws.cell(1_000_000, 1, "stray")     # inflate max_row
    p = os.path.join(tmp_path, "sparse.xlsx"); wb.save(p)
    t0 = time.time()
    lines = read_trial_balance(p)
    assert time.time() - t0 < 3.0
    assert len(lines) == 1


def test_title_sheet_plus_one_tb_sheet_is_not_ambiguous(tmp_path):
    """A cover/title sheet with no TB-shaped columns must not count toward
    ambiguity — only raise when 2+ candidates actually look like a trial
    balance. A prior fix for the ambiguous-sheet guard regressed this exact
    common workflow (title page + one real TB sheet whose name doesn't match
    a known hint)."""
    wb = Workbook()
    ws1 = wb.active; ws1.title = "Cover Page"
    ws1.append(["Prepared for:", "Acme Inc"])
    ws2 = wb.create_sheet("TB Detail")
    ws2.append(["Account Number", "Account Description", "Amount"])
    ws2.append(["5000", "Direct labor", 300000])
    p = os.path.join(tmp_path, "cover.xlsx"); wb.save(p)
    lines = read_trial_balance(p)
    assert len(lines) == 1
    assert lines[0].acct_desc == "Direct labor"
