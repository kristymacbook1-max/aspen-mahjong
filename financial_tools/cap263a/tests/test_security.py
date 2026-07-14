"""Red-team regressions: hostile trial-balance input must not weaponize the
output workbook or escape the output directory."""

import os
from decimal import Decimal

from openpyxl import Workbook, load_workbook

from financial_tools.cap263a.model import TBLine
from financial_tools.cap263a.analysis import analyze, EntityProfile
from financial_tools.cap263a.report import CapitalizationReport, _defuse
from financial_tools.cap263a.pipeline import CapitalizationPipeline


def test_formula_injection_is_neutralized():
    """A =/+/-/@-leading account description is otherwise stored as an active
    Excel formula that executes when a reviewer opens the workbook."""
    for payload in ['=HYPERLINK("http://evil","x")', "=1+2", "+1+1", "-1+1", "@SUM(1)"]:
        assert _defuse(payload) == "'" + payload
    assert _defuse("Normal description") == "Normal description"
    assert _defuse(1234) == 1234


def test_injection_payload_not_written_as_formula(tmp_path):
    payload = '=HYPERLINK("http://evil.example","click")'
    lines = [TBLine("5000", payload, "100", "Production", amount=Decimal("100000"))]
    r = analyze(lines, EntityProfile(avg_gross_receipts=Decimal("75000000")))
    wb = load_workbook(CapitalizationReport().generate(r, os.path.join(tmp_path, "inj.xlsx")))
    tb = wb["Classified TB"]
    cell = tb.cell(4, 2)
    assert cell.data_type == "s"          # text, not 'f' (formula)
    assert not str(cell.value).startswith("=")


def test_long_entity_name_does_not_crash(tmp_path):
    """A ~300-char client entity name would exceed the OS 255-char filename
    limit and raise OSError; the tag must be length-capped."""
    wb = Workbook(); ws = wb.active; ws.title = "TB"
    ws.append(["Account Number", "Account Description", "Amount"])
    ws.append(["5000", "Direct labor", 100000])
    tb = os.path.join(tmp_path, "tb.xlsx"); wb.save(tb)
    out = CapitalizationPipeline(os.path.join(tmp_path, "out")).run(
        tb, company_tag="A" * 300)
    assert os.path.exists(out["_output_path"])
    assert len(os.path.basename(out["_output_path"])) < 200
