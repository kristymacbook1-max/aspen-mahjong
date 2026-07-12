"""Red-team round 3 regressions — fixes to round 2's own fixes plus the
workbook/interview layers (docs/TAX_DECISIONS.md §17). Every test pins a
CONFIRMED finding; docstrings name the original failure."""

import json
from datetime import date
from decimal import Decimal

import pytest

from financial_tools.cap263a.analysis import EntityProfile
from financial_tools.cap263a.engines.interest import compute_263af
from financial_tools.cap263a.engines.inventory import compute_srm
from financial_tools.cap263a.interview import run_interview
from financial_tools.cap263a.model import (CIPProject, CIPSnapshot,
                                           DebtInstrument, REExpenditure)
from financial_tools.cap263a.pipeline import CapitalizationPipeline
from financial_tools.cap263a.readers import _hint_matches

QDATES = [date(2026, 3, 31), date(2026, 6, 30),
          date(2026, 9, 30), date(2026, 12, 31)]
MDATES = [date(2026, m, 28) for m in range(1, 13)]


def _unit(pid, apes, dates):
    return CIPProject(project_id=pid, snapshots=[
        CIPSnapshot(measurement_date=d, cumulative_ape=a)
        for d, a in zip(dates, apes)])


def test_mixed_measurement_grids_do_not_cross_contaminate():
    """ROUND-2'S OWN FIX WAS WRONG on mixed frequencies: a full-year
    quarterly unit next to a monthly unit was zero-padded against the
    16-date union — a 4x understatement of its average excess. Mixed
    (non-nested) grids must warn MIXED-MEASUREMENT-GRID and average each
    unit over its own dates."""
    q = _unit("Q", [Decimal("1000000")] * 4, QDATES)
    m = _unit("M", [Decimal("500000")] * 12, MDATES)
    nontraced = DebtInstrument(debt_id="L", principal=Decimal("10000000"),
                               interest_incurred=Decimal("1000000"))
    out = compute_263af([q, m], [nontraced])
    assert any("MIXED-MEASUREMENT-GRID" in w for w in out["warnings"])
    assert out["per_unit"]["Q"]["average_excess"] == Decimal("1000000")
    assert out["per_unit"]["M"]["average_excess"] == Decimal("500000")
    assert not any("PARTIAL-PERIOD-UNIT" in w for w in out["warnings"])


def test_nested_grids_still_zero_pad_partial_units():
    """The legitimate partial-period case (same convention, unit ends
    mid-year) keeps round 2's zero-padding: [400k, 600k] over the 4-date
    grid averages 250,000."""
    full = _unit("FULL", [Decimal("1000000")] * 4, QDATES)
    part = _unit("PART", [Decimal("400000"), Decimal("600000")], QDATES[:2])
    nontraced = DebtInstrument(debt_id="L", principal=Decimal("10000000"),
                               interest_incurred=Decimal("1000000"))
    out = compute_263af([full, part], [nontraced])
    assert out["per_unit"]["PART"]["average_excess"] == Decimal("250000")
    assert any("PARTIAL-PERIOD-UNIT" in w for w in out["warnings"])


def test_dated_balance_gaps_warn_instead_of_silent_principal():
    """A debt WITH a dated schedule missing a grid date silently substituted
    principal, distorting WAIR 3x in the counterexample."""
    unit = _unit("U", [Decimal("5000000")] * 4, QDATES)
    paid_down = DebtInstrument(
        debt_id="L", principal=Decimal("4000000"),
        interest_incurred=Decimal("60000"),
        outstanding_by_date={QDATES[0].isoformat(): Decimal("1000000")})
    out = compute_263af([unit], [paid_down])
    assert any("DATED-BALANCE-MISSING" in w for w in out["warnings"])


def test_partial_re_options_preserves_interview_election(tmp_path):
    """ROUND-2'S WIRING FIX WAS HALF-CONNECTED: a caller supplementing one
    option (re_options={'catchup_method': ...}) discarded the interview's
    §174A(c) election entirely — $500k flipped from capitalized to deducted
    with zero warnings."""
    doc = {"trial_balance": [{"acct_num": "5000", "acct_desc": "Direct labor",
                              "cc_num": "100", "amount": "1000"}],
           "re": [{"re_id": "R1", "amount": "500000", "domestic": "yes",
                   "tax_year": "2026"}]}
    src = tmp_path / "e.json"
    src.write_text(json.dumps(doc))
    answers = {"Q0.1": "c_corp", "Q0.2": 2026, "Q0.3": False,
               "Q0.4": "75000000", "Q0.5": True, "Q1.1a": True, "Q1.1b": False,
               "Q2.1": "FIFO", "Q2.2": "SPM", "Q2.4": False, "Q2.5": False,
               "Q3.1": "0", "Q8.0": True,
               "Q8.2a": True, "Q8.2b": 60}
    r = CapitalizationPipeline(str(tmp_path / "o")).run_engagement(
        src, answers=answers, re_options={"catchup_method": "one_year"},
        generate_workbook=False)
    assert r["re_174"]["capitalized_total"] == Decimal("500000")


def test_research_building_sheet_no_longer_routes_to_re():
    """'Research Building Depreciation' (a fixed-asset sheet) routed to the
    R&E parser via the bare word 'research' — a phantom $850k §174 deduction
    with zero warnings. 'Loan Covenant Fees' likewise routed to debt."""
    assert "re" not in _hint_matches("Research Building Depreciation")
    assert "debt" not in _hint_matches("Loan Covenant Fees")
    assert _hint_matches("R&D Schedule") == ["re"] or \
        "re" in _hint_matches("R&D Schedule")
    assert "debt" in _hint_matches("Debt Schedule")


def test_ingestion_warnings_reach_all_warnings(tmp_path):
    """The entire round-2 reader warning hardening was DEAD OUTPUT — no
    consumer read validation.warnings."""
    doc = {"trial_balance": [{"acct_num": "5000", "acct_desc": "Direct labor",
                              "cc_num": "100", "amount": "1000"}],
           "debt_balances": [{"debt_id": "GHOST", "date": "2026-06-30",
                              "outstanding": "1"}]}
    src = tmp_path / "e.json"
    src.write_text(json.dumps(doc))
    r = CapitalizationPipeline(str(tmp_path / "o")).run_engagement(
        src, EntityProfile(), generate_workbook=False)
    assert any("debt_balances row skipped" in w for w in r["all_warnings"])


def test_interview_nan_and_infinity_rejected_with_qid():
    """'NaN' receipts crashed with a raw InvalidOperation outside any
    handler; '-Infinity' GRANTED the small-business exemption."""
    with pytest.raises(ValueError, match="Q0.4"):
        run_interview({"Q0.4": "NaN"})
    with pytest.raises(ValueError, match="Q0.4"):
        run_interview({"Q0.4": "-Infinity"})


def test_interview_type_errors_carry_question_id():
    """A list answered to an int question raised a raw TypeError; a truthy
    non-dict prior_elections crashed or silently compared nothing."""
    with pytest.raises(ValueError, match="Q0.2"):
        run_interview({"Q0.2": []})
    with pytest.raises(ValueError, match="prior_elections"):
        run_interview({"Q0.5": False, "prior_elections": 5})
    with pytest.raises(ValueError, match="Q0.2"):
        run_interview({"Q0.2": True})   # bool into int


def test_unanswered_receipts_warns_exemption_undetermined():
    """Q0.4 unanswered: the walk ran full UNICAP gates while the returned
    profile's default-0 receipts read as EXEMPT — self-contradictory, with
    zero warnings."""
    res = run_interview({"Q0.1": "c_corp", "Q0.2": 2026, "Q0.3": False,
                         "Q1.1a": True, "Q1.1b": False, "Q2.2": "SPM"})
    assert any("EXEMPTION-UNDETERMINED" in w for w in res.warnings)


def test_answers_for_unasked_questions_warn():
    """A recognized-but-unasked answer was silently swallowed — hiding the
    Q0.5(first year)+Q0.6(prior method) contradiction entirely."""
    res = run_interview({"Q0.1": "c_corp", "Q0.2": 2026, "Q0.3": False,
                         "Q0.4": "75000000", "Q0.5": True, "Q0.6": "SPM"})
    assert any("ANSWER-FOR-UNASKED-QUESTION" in w and "Q0.6" in w
               for w in res.warnings)


def test_srm_availability_unknown_not_asserted_unavailable():
    """A producer whose de minimis determination was never asked got a hard
    SRM-METHOD-CONFLICT asserting 'above de minimis' — a determination never
    made. Unknown must read as unknown."""
    res = run_interview({"Q0.1": "c_corp", "Q0.2": 2026, "Q0.3": False,
                         "Q0.4": "75000000", "Q0.5": True,
                         "Q1.1a": True, "Q1.1b": False, "Q2.1": "FIFO",
                         "Q2.2": "SRM", "Q2.4": False, "Q2.5": False})
    assert not any("SRM-METHOD-CONFLICT" in w for w in res.warnings)
    assert any("SRM-AVAILABILITY-UNKNOWN" in w for w in res.warnings)


def test_empty_per_item_answers_not_recorded_as_elections():
    res = run_interview({"Q0.1": "c_corp", "Q0.2": 2026, "Q0.3": False,
                         "Q0.4": "75000000", "Q0.5": True, "Q10.1": True,
                         "Q10.2": []})
    assert not any(e["id"] == "Q10.2" for e in res.elections)


def test_variation_a_plus_b_combination_warns():
    p = EntityProfile(avg_gross_receipts=Decimal("75000000"), method="SRM",
                      produces=False, acquires_for_resale=True,
                      purchasing_costs=Decimal("60000"),
                      current_year_471_costs=Decimal("2000000"),
                      storage_handling_costs=Decimal("105000"),
                      beginning_inventory_471=Decimal("400000"),
                      ending_inventory_471=Decimal("500000"),
                      srm_variation_a=True, srm_variation_b=True,
                      ending_inventory_471_total_lifo=Decimal("800000"))
    u = compute_srm({"rows": [], "bucket_totals": {},
                     "mixed_total": Decimal("0"),
                     "deductible_total": Decimal("0")}, p)
    assert any("SRM-VARIATION-A-PLUS-B" in w for w in u["warnings"])


def test_implausible_magnitude_rejected_at_model():
    """Decimal('1e400') passed is_finite() but became float inf at the
    report layer — openpyxl wrote an EMPTY numeric cell, silently blanking
    the basis figure AND its total; with a recovery period it crashed the
    whole workbook via quantize InvalidOperation."""
    with pytest.raises(ValueError, match="implausible"):
        REExpenditure(amount="1e400")


def test_defuse_strips_control_chars_and_truncates():
    """A \\x00 in any description killed the whole workbook with
    IllegalCharacterError; >32,767 chars truncated silently."""
    from financial_tools.cap263a.report import _defuse
    assert _defuse("Bad\x00desc\x07") == "Baddesc"
    long = _defuse("x" * 40000)
    assert len(long) < 32767 and long.endswith("…[TRUNCATED]")
    assert _defuse("=HYPERLINK(1)").startswith("'")


def test_absurd_ratios_warn_in_mspm_and_srm():
    """FUZZ: SPM has warned on >100% absorption since round 1; MSPM/SRM had
    no analog — a 50,000,000x ratio computed a $5 trillion capitalization
    silently."""
    from financial_tools.cap263a.engines.inventory import compute_mspm
    res = {"rows": [], "bucket_totals": {}, "mixed_total": Decimal("0"),
           "deductible_total": Decimal("0")}
    p = EntityProfile(avg_gross_receipts=Decimal("75000000"), method="MSPM",
                      pre_production_471=Decimal("0.01"),
                      pre_production_additional_263A=Decimal("500000"),
                      pre_production_471_on_hand=Decimal("0.01"))
    u = compute_mspm(res, p)
    assert any("ABSORPTION RATIO" in w and ">100%" in w for w in u["warnings"])
    p2 = EntityProfile(avg_gross_receipts=Decimal("75000000"), method="SRM",
                       produces=False, acquires_for_resale=True,
                       purchasing_costs=Decimal("500000"),
                       current_year_471_costs=Decimal("0.01"),
                       ending_inventory_471=Decimal("100000"))
    u2 = compute_srm(res, p2)
    assert any("ABSORPTION RATIO" in w for w in u2["warnings"])


def test_mspm_negative_dm_denominator_warns_not_negative_dollars():
    """FUZZ: ending DM > beginning + purchased made the production
    denominator negative -> negative ratio -> -$500 'capitalized', silent."""
    from financial_tools.cap263a.engines.inventory import compute_mspm
    res = {"rows": [], "bucket_totals": {}, "mixed_total": Decimal("0"),
           "deductible_total": Decimal("0")}
    p = EntityProfile(avg_gross_receipts=Decimal("75000000"), method="MSPM",
                      production_471=Decimal("1000"),
                      production_additional_263A=Decimal("500"),
                      production_471_on_hand=Decimal("2000"),
                      ending_DM_not_yet_in_production=Decimal("3000"))
    u = compute_mspm(res, p)
    assert any("MSPM-NEGATIVE-DENOMINATOR" in w for w in u["warnings"])
    assert u["additional_capitalized_to_inventory"] >= Decimal("0")


def test_huge_but_valid_decimals_warn_instead_of_crashing():
    """FUZZ: quantizing a ratio x pool product near 1e26 crashed with a bare
    InvalidOperation BEFORE any warning could fire."""
    from financial_tools.cap263a.engines.inventory import compute_mspm
    res = {"rows": [], "bucket_totals": {}, "mixed_total": Decimal("0"),
           "deductible_total": Decimal("0")}
    p = EntityProfile(avg_gross_receipts=Decimal("75000000"), method="MSPM",
                      pre_production_471=Decimal("0.01"),
                      pre_production_additional_263A=Decimal("1e12"),
                      pre_production_471_on_hand=Decimal("1e12"))
    u = compute_mspm(res, p)     # must not raise
    assert any("ABSORPTION RATIO" in w for w in u["warnings"])


def test_commissions_survive_the_twelve_month_rule():
    """FUZZ conservation break: a commission on a 12-month-rule item was
    neither deducted, capitalized, nor flagged — $1,000 vanished."""
    from financial_tools.cap263a.engines.intangibles import compute_263a4_5
    from financial_tools.cap263a.model import IntangibleItem
    it = IntangibleItem(item_id="I1", amount=Decimal("100"),
                        benefit_start=date(2026, 1, 1),
                        benefit_end=date(2026, 6, 30), payment_year=2026,
                        facilitative_commissions=Decimal("1000"))
    out = compute_263a4_5([], [it], [])
    assert out["capitalized_total"] == Decimal("1000")
    assert out["deductible_total"] == Decimal("100")


def test_negative_interest_incurred_floored_and_warned():
    """FUZZ: -$100 of nontraced interest produced a negative WAIR, negative
    capitalization, and $99,900 of INVENTED deductible remainder."""
    unit = _unit("U", [Decimal("1000000")], [QDATES[0]])
    bad = DebtInstrument(debt_id="L", principal=Decimal("1000"),
                         interest_incurred=Decimal("-100"))
    out = compute_263af([unit], [bad])
    assert any("NEGATIVE-INTEREST-INCURRED" in w for w in out["warnings"])
    assert out["total_capitalized"] >= Decimal("0")
    assert out["consumption"]["nontraced_remaining_deductible"] == Decimal("0.00")


def test_proration_penny_plug_conserves_to_the_cap():
    """FUZZ: independent per-unit quantization drifted Σ excess a cent PAST
    total available — capitalized interest with no interest source."""
    units = [_unit(f"U{i}", [Decimal("1000000")], [QDATES[0]]) for i in range(3)]
    nt = DebtInstrument(debt_id="L", principal=Decimal("1000"),
                        interest_incurred=Decimal("100.01"))
    out = compute_263af(units, [nt])
    assert out["prorated"] is True
    assert out["total_excess"] == Decimal("100.01")


def test_sca_negative_pool_amount_blocked():
    """FUZZ: the round-2 block covered negative DRIVER VALUES only — a
    negative pool AMOUNT still booked -$1,000 of 'capitalized' (and negative
    APE into Phase D) silently."""
    from financial_tools.cap263a.engines.sca import compute_sca
    from financial_tools.cap263a.model import CostPool, SelfConstructedAsset
    out = compute_sca(
        [CostPool(pool_id="P", amount=Decimal("-1000"), driver="headcount",
                  targets={"A0": Decimal("1")})],
        [SelfConstructedAsset(asset_id="A0", sscm_eligible=True)],
        Decimal("0.5"))
    assert any("NEGATIVE-POOL-AMOUNT" in w for w in out["warnings"])
    assert out["per_asset"]["A0"]["additional_263a"] == Decimal("0")


def test_ppa_negative_fmv_cannot_fabricate_goodwill():
    """FUZZ: a -$100 Class I FMV on a zero-consideration deal fabricated
    +$100 of Class VII goodwill while the sum still tied."""
    from financial_tools.cap263a.engines.purchase_price_allocation import \
        compute_1060_allocation
    from financial_tools.cap263a.model import PurchasePriceAllocation
    out = compute_1060_allocation(PurchasePriceAllocation(
        transaction_id="T", aggregate_consideration=Decimal("0"),
        class_fmv={"I": Decimal("-100")}))
    assert out["class_vii_residual"] == Decimal("0")
    assert any("NEGATIVE-CLASS-FMV" in w for w in out["warnings"])


def test_two_year_catchup_no_subcent_deductions():
    """FUZZ: an odd-cent 2022-2024 basis split into 50.005/50.005."""
    from financial_tools.cap263a.engines.re_capitalization import compute_174
    out = compute_174([], catchup_method="two_year",
                      remaining_2022_2024_basis=Decimal("100.01"))
    c = out["catchup"]
    # banker's rounding parks the odd cent in the following year
    assert c["current_year_deduction"] == Decimal("50.00")
    assert c["following_year_deduction"] == Decimal("50.01")
    assert (c["current_year_deduction"]
            + c["following_year_deduction"]) == Decimal("100.01")


# ---------------------------------------------------------------------------
# Round 4 — the fix-symmetry sweep (docs/TAX_DECISIONS.md §18): the four
# residual gaps found when auditing round 3's fixes for sibling paths.
# ---------------------------------------------------------------------------

def test_entity_profile_rejects_nan_and_implausible_magnitude():
    """ROUND 4: EntityProfile had its OWN Decimal coercion that accepted
    Decimal('NaN') and Decimal('1e400'), bypassing the round-3 model-layer
    guards entirely."""
    with pytest.raises(ValueError, match="non-finite"):
        EntityProfile(ending_inventory_471=Decimal("NaN"))
    with pytest.raises(ValueError, match="implausible"):
        EntityProfile(ending_inventory_471=Decimal("1e400"))


def test_negative_transaction_cost_routed_to_sme_not_deducted():
    """ROUND 4: a -$50,000 fee flowed silently into deductible_total."""
    from financial_tools.cap263a.engines.intangibles import compute_263a4_5
    from financial_tools.cap263a.model import TransactionCostItem
    out = compute_263a4_5([TransactionCostItem(item_id="T1",
                                               amount=Decimal("-50000"))], [], [])
    assert out["deductible_total"] == Decimal("0")
    assert any("NEGATIVE-AMOUNT" in w for w in out["warnings"])


def test_negative_59e_amount_not_scheduled():
    """ROUND 4: a negative qualified expenditure was scheduled silently with
    a negative amortization base."""
    from financial_tools.cap263a.engines.qualified_expenditures import compute_59e
    from financial_tools.cap263a.model import QualifiedExpenditureElection
    out = compute_59e([QualifiedExpenditureElection(
        item_id="Q", category="idc", amount=Decimal("-120000"),
        elected=True, election_year=2026)],
        entity_type="sole_prop", individual_amt_exposure=True)
    assert out["items"] == []
    assert any("NEGATIVE-AMOUNT" in w for w in out["warnings"])


def test_negative_ape_snapshot_warns():
    """ROUND 4: a negative cumulative-APE snapshot floored harmlessly but
    silently — a data error worth naming."""
    unit = _unit("U", [Decimal("-500000")], [QDATES[0]])
    nt = DebtInstrument(debt_id="N", principal=Decimal("1000000"),
                        interest_incurred=Decimal("50000"))
    out = compute_263af([unit], [nt])
    assert any("NEGATIVE-APE" in w for w in out["warnings"])


def test_nonzero_m1_tie_check_is_visually_flagged(tmp_path):
    """The M-1 tie-check cell — whose sole purpose is to be un-missable —
    rendered a nonzero value as a plain unmarked number."""
    from openpyxl import load_workbook
    from financial_tools.cap263a.report import CapitalizationReport
    from financial_tools.cap263a.analysis import analyze
    from financial_tools.cap263a.model import TBLine, BookTaxDifference
    from financial_tools.cap263a.engines.tax_basis_tb import compute_tax_basis_tb
    tb = [TBLine("6100", "Depreciation", "100", "Plant", amount=Decimal("500000")),
          TBLine("6100", "Depreciation", "200", "Office", amount=Decimal("100000"))]
    # acct-only BTD across two CCs -> refused -> intentional tie failure
    out = compute_tax_basis_tb(tb, [BookTaxDifference(
        acct_num="6100", adjustment=Decimal("-50000"))])
    assert out["m1_reconciliation"]["tie_check"] != 0
    r = analyze(out["tax_lines"], EntityProfile(avg_gross_receipts=Decimal("75000000")))
    r["tax_basis_tb"] = out
    wb = load_workbook(CapitalizationReport().generate(r, str(tmp_path / "m1.xlsx")))
    ws = wb["Tax-Basis TB"]
    flagged = [c for row_ in ws.iter_rows() for c in row_
               if isinstance(c.value, str) and "FAILED" in c.value]
    assert flagged, "nonzero tie-check must be visually flagged"
