"""Single-source-of-truth guarantees for the =PY() / exported delivery.

These are the tests the original tool lacked — they assert the exported taxonomy
and the assembled engine blob classify IDENTICALLY to the canonical CLI engine,
so the four-engine divergence problem cannot recur.
"""

from openpyxl import Workbook

from financial_tools.cap263a import classify, get_taxonomy
from financial_tools.cap263a.export_reference import (
    export_reference_sheets, taxonomy_from_workbook)
from financial_tools.cap263a.build_pyexcel import build_engine_source
from financial_tools.cap263a.tests.test_engine import GOLDEN


def test_exported_reference_roundtrips():
    wb = Workbook()
    wb.remove(wb.active)
    export_reference_sheets(wb, hidden=False)
    rebuilt = taxonomy_from_workbook(wb)
    canon = get_taxonomy()
    assert len(rebuilt.categories) == len(canon.categories)
    for acct, cc, code, tier1 in GOLDEN:
        r = classify(acct_desc=acct, cc_desc=cc, tax=rebuilt)
        assert r.code == code, f"roundtrip {acct}/{cc}: {r.code} != {code}"


def test_roundtrip_preserves_empty_lists_and_empty_strings():
    """A blank Excel cell reads back as None, not ''/[] — str(None) == "None"
    used to survive the "|"-split and truthy-filter, corrupting every
    originally-empty keywords/cc_clues list into ["None"] and polluting the
    keyword/clue index for every category that legitimately has none."""
    wb = Workbook()
    wb.remove(wb.active)
    export_reference_sheets(wb, hidden=False)
    rebuilt = taxonomy_from_workbook(wb)
    canon = get_taxonomy()
    for c in rebuilt.categories:
        assert "none" not in [k.lower() for k in c["cc_clues"]], c["code"]
        assert "none" not in [k.lower() for k in c["keywords"]], c["code"]
    canon_by_code = {c["code"]: c for c in canon.categories}
    for c in rebuilt.categories:
        orig = canon_by_code[c["code"]]
        if orig.get("tier3", "") == "":
            assert c["tier3"] == "", f"{c['code']}: tier3 lost empty-string identity"
        if orig.get("tier2", "") == "":
            assert c["tier2"] == "", f"{c['code']}: tier2 lost empty-string identity"


def test_roundtrip_preserves_note_field():
    """The category `note` audit-trail field (documents prior tax-fix
    rationale on EX-BID/NO-INTCAP/NO-OFFICER) used to be silently dropped
    on every export/reload round-trip — not read into _CAT_COLS at all."""
    wb = Workbook()
    wb.remove(wb.active)
    export_reference_sheets(wb, hidden=False)
    rebuilt = taxonomy_from_workbook(wb)
    officer = next(c for c in rebuilt.categories if c["code"] == "NO-OFFICER")
    assert officer.get("note")


def test_assembled_engine_source_is_valid_and_matches():
    src = build_engine_source(with_bootstrap=False)
    compile(src, "<pyexcel>", "exec")          # must be valid Python
    ns = {}
    exec(src, ns)                               # runs the shim + inlined modules

    # rebuild taxonomy from an exported workbook via the assembled code's loader
    wb = Workbook(); wb.remove(wb.active)
    export_reference_sheets(wb, hidden=False)

    def body(name):
        return [list(r) for r in wb[name].iter_rows(min_row=2, values_only=True)]
    data = ns["reference_data_from_rows"](body("_Categories"), body("_CCZones"),
                                          body("_CCReclass"), body("_GenericMap"),
                                          body("_Lexicon"))
    tax = ns["Taxonomy"].from_data(**data)

    for acct, cc, code, tier1 in GOLDEN:
        assembled = ns["classify"](acct_desc=acct, cc_desc=cc, tax=tax)
        canon = classify(acct_desc=acct, cc_desc=cc)
        assert assembled.code == canon.code, f"{acct}/{cc}: {assembled.code} != {canon.code}"
        assert assembled.confidence == canon.confidence


def test_get_taxonomy_fallback_does_not_raise():
    """Regression: the inlined taxonomy.py's own get_taxonomy() calls
    Taxonomy(data=None), which needs the file-loading _load() the shim strips
    out — if that definition were left standing (instead of being rebound
    after inlining) it would silently shadow the shim's version and raise
    NameError the moment classify() is called without an explicit tax=
    (the implicit `tax = tax or get_taxonomy()` fallback in engine.py)."""
    src = build_engine_source(with_bootstrap=False)
    ns = {}
    exec(compile(src, "<pyexcel>", "exec"), ns)
    ns["_TAX"] = get_taxonomy()
    assert ns["get_taxonomy"]() is ns["_TAX"]
    r = ns["classify"](acct_desc="Direct labor", cc_desc="Production")   # no tax= kwarg
    assert r.code == "DL-PROD"
