"""Tangible §263(a) engine wired into the desktop pipeline: readers
tri-state ingestion, run_engagement dispatch, workbook tab, double-count
reconcile, and open-question passthrough."""

from decimal import Decimal

from openpyxl import load_workbook

from financial_tools.cap263a.analysis import EntityProfile
from financial_tools.cap263a.engines.tangible_263a import TangibleExpenditure
from financial_tools.cap263a.model import EngagementData, TBLine
from financial_tools.cap263a.pipeline import CapitalizationPipeline
from financial_tools.cap263a.readers import read_engagement


# ---------------------------------------------------------------------------
# readers: tri-state parsing + routing
# ---------------------------------------------------------------------------

_CSV = """item id,description,amount,invoice cost,unit of property,building,building basis,materials and supplies,betterment,adaptation,restoration,routine maintenance,book capitalized
T1,Roof work,50000,,BLDG-1,yes,900000,no,yes,no,,,
T2,Filters,900,800,BLDG-1,no,,yes,,,,yes,no
"""


def test_readers_tangible_tri_state_csv(tmp_path):
    """Blank repair-regs fact cells must land as None (not established),
    NEVER False; yes/no land True/False; blank invoice cost stays None
    (unknown), never $0."""
    p = tmp_path / "tangible.csv"
    p.write_text(_CSV)
    data = read_engagement(tangible_path=p)
    assert len(data.tangible_items) == 2
    t1, t2 = data.tangible_items

    assert isinstance(t1, TangibleExpenditure)
    assert t1.item_id == "T1" and t1.amount == Decimal("50000")
    assert t1.betterment is True
    assert t1.adaptation is False
    assert t1.restoration is None                      # blank -> None
    assert t1.routine_maintenance_expected_more_than_once is None
    assert t1.invoice_or_item_cost is None             # blank -> None, not $0
    assert t1.is_building is True
    assert t1.building_unadjusted_basis == Decimal("900000")
    assert t1.unit_of_property == "BLDG-1"

    assert t2.is_material_or_supply is True
    assert t2.invoice_or_item_cost == Decimal("800")
    assert t2.betterment is None
    assert t2.routine_maintenance_expected_more_than_once is True
    assert t2.book_capitalized is False


def test_readers_tangible_filename_routing(tmp_path):
    """A directory file named repairs.csv routes to the tangible schedule
    by hint; 'Intangible Assets' must never match the 'tangible' hint
    (word-bounded matching)."""
    d = tmp_path / "engagement"
    d.mkdir()
    (d / "repairs.csv").write_text(_CSV)
    data = read_engagement(d)
    assert len(data.tangible_items) == 2

    from financial_tools.cap263a.readers import _hint_matches
    assert _hint_matches("Intangible Assets") == ["intangibles"]
    assert "tangible" in _hint_matches("263A Tangible Property")


def test_readers_tangible_json_key(tmp_path):
    """The engagement-JSON 'tangible' schedule array parses too."""
    import json
    doc = {"tangible": [{"item id": "J1", "description": "Repave lot",
                         "amount": "12000", "betterment": "no",
                         "adaptation": "no", "restoration": "no"}]}
    p = tmp_path / "engagement.json"
    p.write_text(json.dumps(doc))
    data = read_engagement(p)
    assert len(data.tangible_items) == 1
    j1 = data.tangible_items[0]
    assert j1.item_id == "J1" and j1.betterment is False
    assert j1.restoration is False


# ---------------------------------------------------------------------------
# pipeline: engine dispatch, warnings, workbook tab
# ---------------------------------------------------------------------------

def _engagement():
    """One TB line that classifies into §263(a) Mandatory (so the
    double-count reconcile fires) + two tangible items: a mandatory
    improvement and an all-BAR-None open-question item."""
    return EngagementData(
        tb_lines=[TBLine("6100", "Facilitative transaction cost", "", "",
                         amount=Decimal("100000"))],
        tangible_items=[
            TangibleExpenditure(item_id="T1", description="Roof betterment",
                                amount=Decimal("50000"), betterment=True,
                                adaptation=False, restoration=False),
            TangibleExpenditure(item_id="T2", description="Mystery work",
                                amount=Decimal("10000")),
        ])


def test_pipeline_runs_tangible_engine_end_to_end(tmp_path):
    """run_engagement on an EngagementData with tangible_items: the engine
    result lands in result['tangible_263a'] (50,000 improvement + 10,000
    conservative-pending = 60,000 capitalized), its warnings roll into
    all_warnings, open questions arrive prefixed 'OPEN QUESTION
    [tangible]: ', the §263(a)-tangible double-count reconcile fires
    against the nonzero §263(a) Mandatory bucket, and the workbook renders
    the 'Tangible §263(a)' tab."""
    pipe = CapitalizationPipeline(output_dir=str(tmp_path / "out"))
    profile = EntityProfile(entity_name="Tangible Co",
                            avg_gross_receipts=Decimal("75000000"))
    r = pipe.run_engagement(_engagement(), profile)

    tang = r["tangible_263a"]
    assert tang["capitalized_total"] == Decimal("60000")
    treatments = {i["item_id"]: i["treatment"] for i in tang["items"]}
    assert treatments == {"T1": "improvement_capitalize",
                          "T2": "open_question_capitalize_pending"}

    assert r["bucket_totals"]["§263(a) Mandatory"] == Decimal("100000")
    assert any(w.startswith("DOUBLE-COUNT-RECONCILE [§263(a) tangible]")
               and "$100,000" in w and "$60,000.00" in w
               for w in r["all_warnings"])
    open_qs = [w for w in r["all_warnings"]
               if w.startswith("OPEN QUESTION [tangible]: ")]
    assert open_qs and any("T2" in q for q in open_qs)

    wb = load_workbook(r["_output_path"])
    assert "Tangible §263(a)" in wb.sheetnames
    ws = wb["Tangible §263(a)"]
    texts = [str(c.value) for row in ws.iter_rows() for c in row if c.value]
    assert any("improvement_capitalize" in t for t in texts)
    assert any(t.startswith("? ") for t in texts)      # open questions render
    nums = [c.value for row in ws.iter_rows() for c in row
            if isinstance(c.value, (int, float))]
    assert 60000.0 in nums                              # capitalized total


def test_pipeline_tangible_explicit_elections_kwarg(tmp_path):
    """tangible_elections passes engine kwargs through: with the de minimis
    election on and an $800 invoice cost under the $5,000 AFS ceiling the
    item deducts under §1.263(a)-1(f)."""
    data = EngagementData(
        tb_lines=[TBLine("7000", "Advertising", "", "",
                         amount=Decimal("1000"))],
        tangible_items=[
            TangibleExpenditure(item_id="D1", description="Small tools",
                                amount=Decimal("800"),
                                invoice_or_item_cost=Decimal("800"))])
    pipe = CapitalizationPipeline(output_dir=str(tmp_path / "out"))
    r = pipe.run_engagement(
        data, EntityProfile(avg_gross_receipts=Decimal("75000000")),
        tangible_elections={"de_minimis_election": True},
        generate_workbook=False)
    tang = r["tangible_263a"]
    assert tang["items"][0]["treatment"] == "de_minimis_deduct"
    assert tang["deductible_total"] == Decimal("800")
    # bucket is zero -> no tangible double-count warning
    assert not any(w.startswith("DOUBLE-COUNT-RECONCILE [§263(a) tangible]")
                   for w in r["all_warnings"])


def test_pipeline_interview_derives_tangible_elections(tmp_path):
    """Interview flags feed the engine: Q4.1c (DE_MINIMIS_SAFE_HARBOR_
    ELECTION) turns de_minimis_election on without any explicit kwarg —
    the $800 item deducts. An explicit conflicting kwarg wins and draws
    the TANGIBLE-ELECTION-OVERRIDE warning (same seam as re_options)."""
    answers = {"Q0.2": 2026, "Q0.4": "75000000",
               "Q4.1b": "yes", "Q4.1c": "yes"}
    data = EngagementData(
        tb_lines=[TBLine("7000", "Advertising", "", "",
                         amount=Decimal("1000"))],
        tangible_items=[
            TangibleExpenditure(item_id="D1", description="Small tools",
                                amount=Decimal("800"),
                                invoice_or_item_cost=Decimal("800"))])
    pipe = CapitalizationPipeline(output_dir=str(tmp_path / "out"))
    r = pipe.run_engagement(data, answers=answers, generate_workbook=False)
    assert r["tangible_263a"]["items"][0]["treatment"] == "de_minimis_deduct"

    r2 = pipe.run_engagement(data, answers=answers,
                             tangible_elections={"de_minimis_election": False},
                             generate_workbook=False)
    assert r2["tangible_263a"]["items"][0]["treatment"] != "de_minimis_deduct"
    assert any(w.startswith("TANGIBLE-ELECTION-OVERRIDE")
               for w in r2["all_warnings"])


def test_pipeline_no_tangible_items_no_engine(tmp_path):
    """Empty tangible_items: the engine never runs and no tab renders —
    the legacy shape is untouched."""
    data = EngagementData(
        tb_lines=[TBLine("7000", "Advertising", "", "",
                         amount=Decimal("1000"))])
    pipe = CapitalizationPipeline(output_dir=str(tmp_path / "out"))
    r = pipe.run_engagement(
        data, EntityProfile(avg_gross_receipts=Decimal("75000000")))
    assert "tangible_263a" not in r
    wb = load_workbook(r["_output_path"])
    assert "Tangible §263(a)" not in wb.sheetnames
