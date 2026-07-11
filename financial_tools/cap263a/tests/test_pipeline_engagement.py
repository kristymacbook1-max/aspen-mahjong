"""End-to-end Runtime Pipeline: uploads in -> every engine -> one result."""

import json
from decimal import Decimal

from financial_tools.cap263a.analysis import EntityProfile
from financial_tools.cap263a.pipeline import CapitalizationPipeline


def _engagement_json(tmp_path):
    doc = {
        "trial_balance": [
            {"acct_num": "5000", "acct_desc": "Direct labor", "cc_num": "100",
             "cc_desc": "Production", "amount": "1000000"},
            {"acct_num": "6000", "acct_desc": "Warehouse storage and handling",
             "cc_num": "200", "cc_desc": "Warehouse", "amount": "300000"},
            {"acct_num": "7000", "acct_desc": "Advertising", "cc_num": "400",
             "cc_desc": "Marketing", "amount": "200000"},
        ],
        "btd": [
            {"acct_num": "5000", "cc_num": "100", "description": "Comp accrual",
             "adjustment": "-40000"},
        ],
        "fixed_assets": [
            {"asset_id": "FA1", "description": "Plant", "asset_type": "real",
             "cost": "2500000"},
        ],
        "cip": [
            {"project_id": "P1", "description": "Expansion",
             "linked_asset_id": "FA1", "is_real_property": "yes"},
        ],
        "cip_snapshots": [
            {"project_id": "P1", "measurement_date": "2026-03-31", "cumulative_ape": "3500000"},
            {"project_id": "P1", "measurement_date": "2026-06-30", "cumulative_ape": "5000000"},
            {"project_id": "P1", "measurement_date": "2026-09-30", "cumulative_ape": "6500000"},
            {"project_id": "P1", "measurement_date": "2026-12-31", "cumulative_ape": "8000000"},
        ],
        "debt": [
            {"debt_id": "L1", "principal": "3000000", "rate": "0.06",
             "interest_incurred": "180000", "traced_to": "P1"},
            {"debt_id": "L2", "principal": "2800000", "interest_incurred": "200000"},
        ],
        "re": [
            {"re_id": "R1", "description": "Foreign lab", "amount": "150000",
             "domestic": "no", "tax_year": "2026"},
            {"re_id": "R2", "description": "US software dev", "amount": "500000",
             "domestic": "yes", "tax_year": "2026"},
        ],
        "startup": [
            {"pool_id": "S1", "kind": "startup", "total": "52000",
             "business_commencement": "2026-01-01"},
        ],
        "qualified_expenditures": [
            {"item_id": "Q1", "category": "idc", "amount": "120000",
             "elected": "yes", "election_year": "2026"},
        ],
        "ppa": [
            {"transaction id": "T1", "aggregate consideration": "1000000",
             "class": "IV", "fmv": "300000"},
            {"transaction id": "T1", "class": "V", "fmv": "500000"},
        ],
    }
    p = tmp_path / "engagement.json"
    p.write_text(json.dumps(doc))
    return p


def test_full_pipeline_runs_every_engine(tmp_path):
    """One engagement upload drives every engine: tax-basis TB applied before
    classification (labor 1,000,000 - 40,000 = 960,000 in the §471 bucket),
    §263A(f) hits the golden 376,428.57, foreign R&E amortizes, the startup
    pool computes 3,000 first-year deduction, the §59(e) item is gated OUT
    for a C-corp with no individual AMT exposure, and §1060 leaves a 200,000
    Class VII residual."""
    pipe = CapitalizationPipeline(output_dir=str(tmp_path / "out"))
    profile = EntityProfile(entity_name="E2E Test Co",
                            avg_gross_receipts=Decimal("75000000"),
                            ending_inventory_471=Decimal("500000"))
    r = pipe.run_engagement(_engagement_json(tmp_path), profile)

    # Step 2: tax-basis TB materialized and used downstream
    assert r["tax_basis_tb"]["m1_reconciliation"]["tie_check"] == Decimal("0")
    assert r["bucket_totals"]["Inventory §471"] == Decimal("960000")

    # Phase D golden (the plan's corrected fixture)
    interest = r["interest_263af"]
    assert interest["total_capitalized"] == Decimal("376428.57")
    assert interest["per_unit"]["P1"]["traced_interest"] == Decimal("180000")

    # Phase F: foreign capitalized (mid-year 15-yr), domestic expensed
    re_out = r["re_174"]
    assert re_out["capitalized_total"] == Decimal("150000")
    foreign_items = [i for i in re_out["amortizable_items"]]
    assert foreign_items and foreign_items[0].first_year_amortization() == Decimal("5000.00")

    # Phase G: startup 52,000 -> 3,000 first-year deduction, remainder amortizes
    su = r["intangibles_263a45"]["startup_items"][0]
    assert su["first_year_deduction"] == Decimal("3000")

    # Phase H: gated out for a C-corp with no individual-AMT exposure
    assert r["qualified_59e"]["items"] == []
    assert any("C corp" in w or "CAMT" in w for w in r["qualified_59e"]["warnings"])

    # §1060: 1,000,000 - 300,000 (IV) - 500,000 (V) = 200,000 Class VII
    ppa = r["ppa_1060"][0]
    assert ppa["class_vii_residual"] == Decimal("200000")

    # consolidated outputs
    assert r["basis_amortization"], "AmortizableItem rows must roll up"
    assert isinstance(r["all_warnings"], list)
    assert r["_output_path"].endswith(".xlsx")


def test_engagement_workbook_renders_engine_tabs(tmp_path):
    """The run_engagement workbook must carry the engine outputs — the three
    engagement tabs render with the golden §263A(f) total and the Basis &
    Amortization rows; the legacy TB-only path keeps exactly five tabs (pinned
    by test_report.py)."""
    from openpyxl import load_workbook
    pipe = CapitalizationPipeline(output_dir=str(tmp_path / "out"))
    profile = EntityProfile(entity_name="Tabs Co",
                            avg_gross_receipts=Decimal("75000000"))
    r = pipe.run_engagement(_engagement_json(tmp_path), profile)
    wb = load_workbook(r["_output_path"])
    for tab in ("Tax-Basis TB", "Basis & Amortization", "Engine Results"):
        assert tab in wb.sheetnames, wb.sheetnames
    ba = wb["Basis & Amortization"]
    cells = [str(c.value) for row in ba.iter_rows() for c in row if c.value is not None]
    assert any(v == "376428.57" or v == "376428.57" for v in
               [f"{c}" for c in cells]) or \
        any(abs(float(c.value or 0) - 376428.57) < 0.01
            for row in ba.iter_rows() for c in row
            if isinstance(c.value, (int, float)))
    tbtab = wb["Tax-Basis TB"]
    vals = [c.value for row in tbtab.iter_rows() for c in row
            if isinstance(c.value, (int, float))]
    assert 0 in vals or 0.0 in vals          # the tie-check row renders as 0


def test_exempt_entity_skips_263af_but_not_174(tmp_path):
    """§263A(i) exemption turns off §263A(f) but NOT §174 (a non-§263A
    provision) — the pipeline must gate them differently."""
    pipe = CapitalizationPipeline(output_dir=str(tmp_path / "out"))
    profile = EntityProfile(avg_gross_receipts=Decimal("1000000"))
    r = pipe.run_engagement(_engagement_json(tmp_path), profile,
                            generate_workbook=False)
    assert "interest_263af" not in r
    assert r["re_174"]["capitalized_total"] == Decimal("150000")


def test_59e_runs_for_individual_amt_exposure(tmp_path):
    pipe = CapitalizationPipeline(output_dir=str(tmp_path / "out"))
    profile = EntityProfile(entity_type="partnership",
                            avg_gross_receipts=Decimal("75000000"))
    r = pipe.run_engagement(_engagement_json(tmp_path), profile,
                            individual_amt_exposure=True,
                            generate_workbook=False)
    items = r["qualified_59e"]["items"]
    assert items and items[0]["category"] == "idc"


def test_mspm_and_srm_results_render_to_workbook(tmp_path):
    """RED-TEAM REGRESSION: the Summary tab indexed SPM-only unicap keys
    (sec471_pool/absorption_ratio) — the workbook generator crashed with
    KeyError for EVERY non-exempt MSPM or SRM client while 201 tests stayed
    green, because no test ever drove analyze() -> engine -> report."""
    from openpyxl import load_workbook
    from financial_tools.cap263a.model import TBLine
    from financial_tools.cap263a.analysis import analyze
    from financial_tools.cap263a.report import CapitalizationReport
    lines = [TBLine("5000", "Direct labor", "100", "Production",
                    amount=Decimal("100000"))]
    mspm = EntityProfile(avg_gross_receipts=Decimal("75000000"), method="MSPM",
                         pre_production_471=Decimal("2500000"),
                         production_471=Decimal("7500000"),
                         pre_production_additional_263A=Decimal("200000"),
                         production_additional_263A=Decimal("800000"),
                         pre_production_471_on_hand=Decimal("1000000"),
                         production_471_on_hand=Decimal("2000000"),
                         beginning_DM_not_yet_in_production=Decimal("400000"),
                         ending_DM_not_yet_in_production=Decimal("800000"),
                         DM_purchased_during_year=Decimal("1900000"))
    r = analyze(lines, mspm)
    assert r["unicap"]["additional_capitalized_to_inventory"] == Decimal("284400.00")
    wb = load_workbook(CapitalizationReport().generate(
        r, str(tmp_path / "mspm.xlsx")))
    summ = wb["Summary Dashboard"]
    texts = [str(c.value) for row in summ.iter_rows() for c in row if c.value]
    assert any("MSPM" in t for t in texts)

    srm = EntityProfile(avg_gross_receipts=Decimal("75000000"), method="SRM",
                        produces=False, acquires_for_resale=True,
                        purchasing_costs=Decimal("60000"),
                        current_year_471_costs=Decimal("2000000"),
                        storage_handling_costs=Decimal("105000"),
                        beginning_inventory_471=Decimal("400000"),
                        ending_inventory_471=Decimal("500000"))
    r2 = analyze(lines, srm)
    assert r2["unicap"]["additional_capitalized_to_inventory"] == Decimal("36875.00")
    CapitalizationReport().generate(r2, str(tmp_path / "srm.xlsx"))


def test_sca_pipeline_integration_and_exemption_gate(tmp_path):
    """RED-TEAM: run_engagement's SCA path was never tested, and it ran SCA
    for §263A(i)-exempt taxpayers (whose unicap dict has no mixed_alloc_ratio
    key, silently allocating at ratio 0). SCA is a §263A mechanic — the
    exemption must skip it, with a visible note."""
    from financial_tools.cap263a.model import CostPool, SelfConstructedAsset
    pools = [CostPool(pool_id="HR", description="HR support", amount=Decimal("100000"),
                      driver="headcount", is_mixed_service=True,
                      targets={"A1": Decimal("75"), "A2": Decimal("25")})]
    assets = [SelfConstructedAsset(asset_id="A1", book_cost=Decimal("0"), sscm_eligible=True),
              SelfConstructedAsset(asset_id="A2", book_cost=Decimal("0"), sscm_eligible=True)]
    pipe = CapitalizationPipeline(output_dir=str(tmp_path / "out"))
    big = EntityProfile(avg_gross_receipts=Decimal("75000000"),
                        mixed_alloc_ratio=Decimal("0.6"))
    r = pipe.run_engagement(_engagement_json(tmp_path), big,
                            sca_pools=pools, sca_assets=assets,
                            generate_workbook=False)
    per = r["sca"]["per_asset"]
    assert per["A1"]["mixed_263a"] == Decimal("45000.00")
    assert per["A2"]["mixed_263a"] == Decimal("15000.00")

    small = EntityProfile(avg_gross_receipts=Decimal("1000000"))
    r2 = pipe.run_engagement(_engagement_json(tmp_path), small,
                             sca_pools=pools, sca_assets=assets,
                             generate_workbook=False)
    assert "sca" not in r2
    assert any("SCA skipped" in w for w in r2["all_warnings"])


def test_bool_in_money_field_raises_clear_error():
    """RED-TEAM: a JSON `true` in a money field crashed with a bare
    InvalidOperation deep inside Decimal(); it must name the field."""
    import pytest
    with pytest.raises(TypeError, match="ending_inventory_471"):
        EntityProfile(ending_inventory_471=True)
