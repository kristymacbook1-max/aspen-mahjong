"""Phase 1 Expense Recognition Excel Report Generator.

Produces a 9-tab workbook from Phase 1 analysis results:
  1. Executive Summary
  2. Expense Categories
  3. Economic Performance
  4. Prepaid Expenses
  5. Compensation Analysis
  6. Interest & R&D
  7. Related Party
  8. Accrued Liabilities
  9. Phase 2 Roadmap
"""

import os
import sys
from decimal import Decimal
from openpyxl import Workbook

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', '..'))
from revenue_recognition.utils.excel_styles import (
    FONT_TITLE, FONT_SECTION_HEADER, FONT_COLUMN_HEADER, FONT_BODY,
    FONT_BODY_BOLD, FILL_HEADER, FILL_SECTION, FILL_ALT_ROW,
    FILL_HIGHLIGHT_GREEN, FILL_HIGHLIGHT_RED, FILL_HIGHLIGHT_ORANGE,
    FILL_HIGHLIGHT_BLUE, THIN_BORDER, ALIGN_CENTER, ALIGN_LEFT,
    ALIGN_RIGHT, ALIGN_WRAP, FMT_CURRENCY, FMT_PERCENT,
    apply_header_row, apply_section_header, apply_title,
    auto_fit_columns, write_data_row, format_currency_cell,
    get_risk_font, get_risk_fill, COLORS,
)
from revenue_recognition.utils.currency_helpers import to_decimal, format_currency, format_millions


class Phase1ExpenseReport:
    """Generates the Phase 1 expense recognition Excel workbook."""

    def generate(self, results: dict, output_path: str) -> str:
        self._r = results
        self._wb = Workbook()

        # Remove default sheet
        self._wb.remove(self._wb.active)

        self._create_executive_summary()
        self._create_expense_categories()
        self._create_economic_performance()
        self._create_prepaid_expenses()
        self._create_compensation()
        self._create_interest_rd()
        self._create_related_party()
        self._create_accrued_liabilities()
        self._create_phase2_roadmap()

        os.makedirs(os.path.dirname(output_path) or '.', exist_ok=True)
        self._wb.save(output_path)
        return output_path

    # ------------------------------------------------------------------
    def _create_executive_summary(self):
        ws = self._wb.create_sheet("Executive Summary")
        company = self._r["company"]
        opps = self._r["total_opportunities"]

        apply_title(ws, 1, 1, 6, f"Expense Recognition Analysis — {company.name}")
        row = 3

        # Company info
        apply_section_header(ws, row, 1, 6, "Company Information")
        row += 1
        info = [
            ("Company", company.name),
            ("Ticker", company.ticker),
            ("Entity Type", self._r["entity_config"].label),
            ("Fiscal Year", str(company.tax_year)),
            ("Total Assets", format_millions(company.total_assets)),
            ("Total Revenue", format_millions(company.total_revenue)),
            ("Industry", company.industry),
        ]
        for label, val in info:
            ws.cell(row=row, column=1, value=label).font = FONT_BODY_BOLD
            ws.cell(row=row, column=2, value=val).font = FONT_BODY
            row += 1

        row += 1
        apply_section_header(ws, row, 1, 6, "Opportunity Summary")
        row += 1
        headers = ["Category", "Book-Tax Difference", "Estimated Tax Impact"]
        for i, h in enumerate(headers, 1):
            ws.cell(row=row, column=i, value=h)
        apply_header_row(ws, row, 1, len(headers))
        row += 1

        for item in opps["line_items"]:
            ws.cell(row=row, column=1, value=item["category"]).font = FONT_BODY
            c = ws.cell(row=row, column=2, value=float(item["amount"]))
            format_currency_cell(c)
            c2 = ws.cell(row=row, column=3, value=float(item["tax_impact"]))
            format_currency_cell(c2)
            for col in range(1, 4):
                ws.cell(row=row, column=col).border = THIN_BORDER
            row += 1

        # Totals
        ws.cell(row=row, column=1, value="TOTAL").font = FONT_BODY_BOLD
        c = ws.cell(row=row, column=2, value=float(opps["total_book_tax_differences"]))
        format_currency_cell(c)
        c.font = FONT_BODY_BOLD
        c2 = ws.cell(row=row, column=3, value=float(opps["estimated_tax_impact"]))
        format_currency_cell(c2)
        c2.font = FONT_BODY_BOLD
        row += 2

        # Top findings
        apply_section_header(ws, row, 1, 6, "Top Findings")
        row += 1
        for i, f in enumerate(self._r["findings"][:8], 1):
            ws.cell(row=row, column=1, value=f["area"]).font = FONT_BODY_BOLD
            ws.cell(row=row, column=2, value=f["description"]).font = FONT_BODY
            ws.cell(row=row, column=3, value=f["risk_level"].upper()).font = get_risk_font(f["risk_level"])
            row += 1

        auto_fit_columns(ws)

    # ------------------------------------------------------------------
    def _create_expense_categories(self):
        ws = self._wb.create_sheet("Expense Categories")
        apply_title(ws, 1, 1, 8, "Expense Categories Analysis")

        headers = [
            "Category", "Book Amount (CY)", "Book Amount (PY)", "Change %",
            "Tax Treatment", "IRC Section", "EP Category", "Timing Difference",
            "Recurring Item Eligible", "Method Change", "Risk",
        ]
        row = 3
        for i, h in enumerate(headers, 1):
            ws.cell(row=row, column=i, value=h)
        apply_header_row(ws, row, 1, len(headers))
        row += 1

        for cat in self._r["expense_categories"]:
            cy = to_decimal(cat.book_amount_current)
            py = to_decimal(cat.book_amount_prior)
            from revenue_recognition.utils.currency_helpers import pct_change
            chg = pct_change(cy, py)
            data = [
                cat.name, float(cy), float(py),
                chg if chg is not None else "N/A",
                cat.tax_treatment, cat.irc_section,
                cat.economic_performance_category, float(cat.timing_difference),
                "Yes" if cat.is_recurring_item_eligible else "No",
                cat.method_change_opportunity, cat.risk_level.upper(),
            ]
            alt = (row % 2 == 0)
            write_data_row(ws, row, data, alternate=alt)
            format_currency_cell(ws.cell(row=row, column=2))
            format_currency_cell(ws.cell(row=row, column=3))
            format_currency_cell(ws.cell(row=row, column=8))
            ws.cell(row=row, column=11).font = get_risk_font(cat.risk_level)
            row += 1

        auto_fit_columns(ws)

    # ------------------------------------------------------------------
    def _create_economic_performance(self):
        ws = self._wb.create_sheet("Economic Performance")
        apply_title(ws, 1, 1, 7, "§461(h) Economic Performance Analysis")
        row = 3

        for section_name, items in [
            ("Payment Liabilities — Deductible When PAID (§461(h)(2)(C))",
             self._r["economic_performance"]["payment_liabilities"]),
            ("Service Liabilities — Deductible When Services PROVIDED (§461(h)(2)(A)(i))",
             self._r["economic_performance"]["service_liabilities"]),
            ("Property Liabilities — Deductible When Property PROVIDED (§461(h)(2)(A)(ii))",
             self._r["economic_performance"]["property_liabilities"]),
            ("Special Rules — Workers' Comp, Torts, Contested (§461(f))",
             self._r["economic_performance"]["special_liabilities"]),
        ]:
            apply_section_header(ws, row, 1, 7, section_name)
            row += 1
            if not items:
                ws.cell(row=row, column=1, value="No items identified").font = FONT_BODY
                row += 2
                continue

            headers = ["Expense", "Book Amount", "Tax Treatment", "IRC Section",
                       "Timing Difference", "Risk"]
            for i, h in enumerate(headers, 1):
                ws.cell(row=row, column=i, value=h)
            apply_header_row(ws, row, 1, len(headers))
            row += 1

            for item in items:
                write_data_row(ws, row, [
                    item["name"],
                    float(to_decimal(item["book_amount"])),
                    item["tax_treatment"],
                    item.get("irc_section", "§461(h)"),
                    float(to_decimal(item["timing_difference"])),
                    item["risk_level"].upper(),
                ], alternate=(row % 2 == 0))
                format_currency_cell(ws.cell(row=row, column=2))
                format_currency_cell(ws.cell(row=row, column=5))
                ws.cell(row=row, column=6).font = get_risk_font(item["risk_level"])
                row += 1
            row += 1

        auto_fit_columns(ws)

    # ------------------------------------------------------------------
    def _create_prepaid_expenses(self):
        ws = self._wb.create_sheet("Prepaid Expenses")
        apply_title(ws, 1, 1, 7, "Prepaid Expenses — 12-Month Rule Analysis")

        row = 3
        headers = ["Expense", "Amount", "Period (Months)", "Qualifies 12-Mo Rule",
                    "Current Treatment", "Optimal Treatment", "Savings", "Authority"]
        for i, h in enumerate(headers, 1):
            ws.cell(row=row, column=i, value=h)
        apply_header_row(ws, row, 1, len(headers))
        row += 1

        for item in self._r["prepaid_expenses"]["items"]:
            write_data_row(ws, row, [
                item["name"],
                float(to_decimal(item["amount"])),
                item["period_months"],
                "Yes" if item["qualifies_12_month"] else "No",
                item["current_treatment"],
                item["optimal_treatment"],
                float(to_decimal(item["savings"])),
                item["authority"],
            ], alternate=(row % 2 == 0))
            format_currency_cell(ws.cell(row=row, column=2))
            format_currency_cell(ws.cell(row=row, column=7))
            row += 1

        row += 1
        ws.cell(row=row, column=1, value="Total Savings Opportunity").font = FONT_BODY_BOLD
        c = ws.cell(row=row, column=7, value=float(self._r["prepaid_expenses"]["total_savings"]))
        format_currency_cell(c)
        c.font = FONT_BODY_BOLD

        auto_fit_columns(ws)

    # ------------------------------------------------------------------
    def _create_compensation(self):
        ws = self._wb.create_sheet("Compensation Analysis")
        apply_title(ws, 1, 1, 8, "Compensation & Benefits Analysis")
        row = 3

        apply_section_header(ws, row, 1, 8, "§162(m) Executive Compensation Analysis")
        row += 1
        headers = ["Employee/Group", "Type", "Book Expense", "Tax Deductible",
                    "§162(m) Disallowed", "Timing Diff", "Deferred Comp Rules", "Risk"]
        for i, h in enumerate(headers, 1):
            ws.cell(row=row, column=i, value=h)
        apply_header_row(ws, row, 1, len(headers))
        row += 1

        comp = self._r["compensation"]
        for item in comp["items"]:
            write_data_row(ws, row, [
                item["name"], item["employee_type"],
                float(to_decimal(item["book_expense"])),
                float(to_decimal(item["tax_deductible"])),
                float(to_decimal(item["section_162m_disallowed"])),
                float(to_decimal(item["timing_difference"])),
                item["deferred_comp_rules"],
                item["risk_level"].upper(),
            ], alternate=(row % 2 == 0))
            format_currency_cell(ws.cell(row=row, column=3))
            format_currency_cell(ws.cell(row=row, column=4))
            format_currency_cell(ws.cell(row=row, column=5))
            format_currency_cell(ws.cell(row=row, column=6))
            ws.cell(row=row, column=8).font = get_risk_font(item["risk_level"])
            row += 1

        row += 1
        ws.cell(row=row, column=1, value="Total §162(m) Disallowed").font = FONT_BODY_BOLD
        c = ws.cell(row=row, column=5, value=float(comp["total_162m_disallowed"]))
        format_currency_cell(c)
        c.font = FONT_BODY_BOLD

        auto_fit_columns(ws)

    # ------------------------------------------------------------------
    def _create_interest_rd(self):
        ws = self._wb.create_sheet("Interest & R&D")
        apply_title(ws, 1, 1, 5, "§163(j) Interest Limitation & §174 R&D Capitalization")
        row = 3

        # §163(j) section
        apply_section_header(ws, row, 1, 5, "§163(j) Business Interest Limitation")
        row += 1
        il = self._r["interest_limitation"]
        lines = [
            ("Total Business Interest Expense", il["total_interest"]),
            ("Business Interest Income", il["business_interest_income"]),
            ("Floor Plan Financing Interest", il["floor_plan_interest"]),
            ("Adjusted Taxable Income (ATI)", il["adjusted_taxable_income"]),
            ("30% of ATI", il["thirty_pct_ati"]),
            ("Deductible Amount", il["deductible_amount"]),
            ("Disallowed Amount", il["disallowed_amount"]),
            ("Carryforward", il["carryforward"]),
        ]
        for label, val in lines:
            ws.cell(row=row, column=1, value=label).font = FONT_BODY
            c = ws.cell(row=row, column=2, value=float(to_decimal(val)))
            format_currency_cell(c)
            ws.cell(row=row, column=1).border = THIN_BORDER
            c.border = THIN_BORDER
            row += 1

        if il["is_limited"]:
            ws.cell(row=row, column=1, value="⚠ Interest is LIMITED under §163(j)").font = get_risk_font("high")
        else:
            ws.cell(row=row, column=1, value="Interest is fully deductible").font = get_risk_font("low")
        row += 2

        # §174 section
        apply_section_header(ws, row, 1, 5, "§174 Research & Development Capitalization")
        row += 1
        rd = self._r["rd_capitalization"]
        rd_lines = [
            ("Total R&D Expense (Book)", rd["total_rd"]),
            ("Domestic R&D", rd["domestic_rd"]),
            ("Foreign R&D", rd["foreign_rd"]),
            ("Domestic Annual Amortization (5 yr)", rd["domestic_annual_amortization"]),
            ("Foreign Annual Amortization (15 yr)", rd["foreign_annual_amortization"]),
            ("Total Tax Amortization", rd["total_amortization"]),
            ("Book-Tax Difference", rd["book_tax_difference"]),
        ]
        for label, val in rd_lines:
            ws.cell(row=row, column=1, value=label).font = FONT_BODY
            c = ws.cell(row=row, column=2, value=float(to_decimal(val)))
            format_currency_cell(c)
            ws.cell(row=row, column=1).border = THIN_BORDER
            c.border = THIN_BORDER
            row += 1

        auto_fit_columns(ws)

    # ------------------------------------------------------------------
    def _create_related_party(self):
        ws = self._wb.create_sheet("Related Party")
        apply_title(ws, 1, 1, 7, "§267 Related Party Expense Analysis")
        row = 3

        headers = ["Payee", "Relationship", "Amount", "Payee Year End",
                    "Includible by Payee", "Deferred Amount", "Risk"]
        for i, h in enumerate(headers, 1):
            ws.cell(row=row, column=i, value=h)
        apply_header_row(ws, row, 1, len(headers))
        row += 1

        rp = self._r["related_party"]
        for item in rp["items"]:
            write_data_row(ws, row, [
                item["payee_name"], item["relationship"],
                float(to_decimal(item["amount"])), item["payee_year_end"],
                float(to_decimal(item["includible_by_payee"])),
                float(to_decimal(item["deferred"])),
                item["risk_level"].upper(),
            ], alternate=(row % 2 == 0))
            format_currency_cell(ws.cell(row=row, column=3))
            format_currency_cell(ws.cell(row=row, column=5))
            format_currency_cell(ws.cell(row=row, column=6))
            ws.cell(row=row, column=7).font = get_risk_font(item["risk_level"])
            row += 1

        row += 1
        ws.cell(row=row, column=1, value="Total Deferred").font = FONT_BODY_BOLD
        c = ws.cell(row=row, column=6, value=float(rp["total_deferred"]))
        format_currency_cell(c)
        c.font = FONT_BODY_BOLD

        auto_fit_columns(ws)

    # ------------------------------------------------------------------
    def _create_accrued_liabilities(self):
        ws = self._wb.create_sheet("Accrued Liabilities")
        apply_title(ws, 1, 1, 8, "Accrued Liabilities & Reserves Analysis")
        row = 3

        headers = ["Liability", "Book Balance", "Tax Deductible",
                    "Timing Difference", "All-Events Met", "EP Met",
                    "Recurring Exception", "Timing Rule", "Risk"]
        for i, h in enumerate(headers, 1):
            ws.cell(row=row, column=i, value=h)
        apply_header_row(ws, row, 1, len(headers))
        row += 1

        reserves = self._r["reserves"]
        for item in reserves["items"]:
            write_data_row(ws, row, [
                item["name"],
                float(to_decimal(item["book_balance"])),
                float(to_decimal(item["tax_deductible"])),
                float(to_decimal(item["timing_difference"])),
                "Yes" if item["all_events_met"] else "No",
                "Yes" if item["ep_met"] else "No",
                "Yes" if item["recurring_exception"] else "No",
                item["timing_rule"],
                item["risk_level"].upper(),
            ], alternate=(row % 2 == 0))
            format_currency_cell(ws.cell(row=row, column=2))
            format_currency_cell(ws.cell(row=row, column=3))
            format_currency_cell(ws.cell(row=row, column=4))
            ws.cell(row=row, column=9).font = get_risk_font(item["risk_level"])
            row += 1

        row += 1
        ws.cell(row=row, column=1, value="Total Timing Difference").font = FONT_BODY_BOLD
        c = ws.cell(row=row, column=4, value=float(reserves["total_timing_difference"]))
        format_currency_cell(c)
        c.font = FONT_BODY_BOLD

        auto_fit_columns(ws)

    # ------------------------------------------------------------------
    def _create_phase2_roadmap(self):
        ws = self._wb.create_sheet("Phase 2 Roadmap")
        apply_title(ws, 1, 1, 5, "Phase 2 Recommendations & Data Requirements")
        row = 3

        headers = ["Priority", "Analysis Area", "Description", "Data Needed"]
        for i, h in enumerate(headers, 1):
            ws.cell(row=row, column=i, value=h)
        apply_header_row(ws, row, 1, len(headers))
        row += 1

        for item in self._r["phase2_roadmap"]:
            write_data_row(ws, row, [
                item["priority"], item["area"],
                item["description"], item["data_needed"],
            ], alternate=(row % 2 == 0))
            row += 1

        auto_fit_columns(ws)
