"""Phase 3 Expense Recognition Excel Report Generator.

Produces an 8-tab workbook from Phase 3 analysis results:
  1. Executive Summary
  2. Contract Portfolio
  3. Economic Performance Detail
  4. Expense Schedule
  5. Reserve Analysis
  6. Compensation Detail
  7. Book-Tax Reconciliation
  8. Implementation Plan
"""

import os
import sys
from decimal import Decimal
from openpyxl import Workbook

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', '..'))
from revenue_recognition.utils.excel_styles import (
    FONT_TITLE, FONT_SECTION_HEADER, FONT_COLUMN_HEADER, FONT_BODY,
    FONT_BODY_BOLD, FILL_SECTION, FILL_ALT_ROW, FILL_HIGHLIGHT_BLUE,
    FILL_HIGHLIGHT_GREEN, FILL_HIGHLIGHT_RED, THIN_BORDER, ALIGN_CENTER,
    ALIGN_LEFT, ALIGN_RIGHT, ALIGN_WRAP, FMT_CURRENCY,
    apply_header_row, apply_section_header, apply_title,
    auto_fit_columns, write_data_row, format_currency_cell,
    get_risk_font, get_risk_fill,
)
from revenue_recognition.utils.currency_helpers import to_decimal, format_currency


class Phase3ExpenseReport:
    """Generates the Phase 3 expense recognition Excel workbook."""

    def generate(self, results: dict, output_path: str) -> str:
        self._r = results
        self._wb = Workbook()
        self._wb.remove(self._wb.active)

        self._create_executive_summary()
        self._create_contract_portfolio()
        self._create_ep_detail()
        self._create_expense_schedule()
        self._create_reserve_analysis()
        self._create_compensation_detail()
        self._create_book_tax_reconciliation()
        self._create_implementation_plan()

        os.makedirs(os.path.dirname(output_path) or '.', exist_ok=True)
        self._wb.save(output_path)
        return output_path

    # ------------------------------------------------------------------
    def _create_executive_summary(self):
        ws = self._wb.create_sheet("Executive Summary")
        apply_title(ws, 1, 1, 6, f"Phase 3 Expense Analysis — {self._r['company_name']}")
        row = 3

        apply_section_header(ws, row, 1, 6, "Key Findings")
        row += 1
        recon = self._r["book_tax_reconciliation"]
        metrics = [
            ("Total Book Expenses Analyzed", recon["total_book"]),
            ("Total Tax Deductions", recon["total_tax"]),
            ("Total Book-Tax Difference", recon["total_difference"]),
            ("Contracts Analyzed", len(self._r["contract_results"])),
            ("Reserves Analyzed", self._r["reserves"]["count"]),
            ("Compensation Contracts", self._r["compensation"]["count"]),
        ]
        for label, val in metrics:
            ws.cell(row=row, column=1, value=label).font = FONT_BODY_BOLD
            if isinstance(val, (int, float)):
                ws.cell(row=row, column=2, value=val).font = FONT_BODY
            else:
                c = ws.cell(row=row, column=2, value=float(to_decimal(val)))
                format_currency_cell(c)
            row += 1

        row += 1
        apply_section_header(ws, row, 1, 6, "Priority Findings")
        row += 1
        for f in self._r["findings"][:8]:
            ws.cell(row=row, column=1, value=f["area"]).font = FONT_BODY_BOLD
            ws.cell(row=row, column=2, value=f["description"]).font = FONT_BODY
            c = ws.cell(row=row, column=3, value=float(to_decimal(f["book_tax_difference"])))
            format_currency_cell(c)
            ws.cell(row=row, column=4, value=f["risk_level"].upper()).font = get_risk_font(f["risk_level"])
            row += 1

        auto_fit_columns(ws)

    # ------------------------------------------------------------------
    def _create_contract_portfolio(self):
        ws = self._wb.create_sheet("Contract Portfolio")
        apply_title(ws, 1, 1, 9, "Vendor Contract Portfolio")
        row = 3

        headers = ["Contract ID", "Vendor", "Description", "Type", "Total Value",
                    "Annual Amount", "Related Party", "Book Expense", "Tax Deduction",
                    "Timing Diff", "Line Items"]
        for i, h in enumerate(headers, 1):
            ws.cell(row=row, column=i, value=h)
        apply_header_row(ws, row, 1, len(headers))
        row += 1

        for cr in self._r["contract_results"]:
            c = cr["contract"]
            write_data_row(ws, row, [
                c["id"], c["vendor"], c["description"], c["type"],
                float(to_decimal(c["total_value"])),
                float(to_decimal(c["annual_amount"])),
                "Yes" if c["is_related_party"] else "No",
                float(to_decimal(cr["total_book_expense"])),
                float(to_decimal(cr["total_tax_deduction"])),
                float(to_decimal(cr["total_timing_difference"])),
                cr["line_item_count"],
            ], alternate=(row % 2 == 0))
            for col in [5, 6, 8, 9, 10]:
                format_currency_cell(ws.cell(row=row, column=col))
            row += 1

        auto_fit_columns(ws)

    # ------------------------------------------------------------------
    def _create_ep_detail(self):
        ws = self._wb.create_sheet("Economic Performance Detail")
        apply_title(ws, 1, 1, 8, "§461(h) Economic Performance Testing — Per Line Item")
        row = 3

        for cr in self._r["contract_results"]:
            vendor = cr["contract"]["vendor"]
            apply_section_header(ws, row, 1, 8, f"{vendor} ({cr['contract']['id']})")
            row += 1

            if not cr["ep_results"]:
                ws.cell(row=row, column=1, value="No line items").font = FONT_BODY
                row += 2
                continue

            headers = ["Line ID", "Amount", "EP Category", "Rule Applied",
                       "EP Met", "EP Date", "Tax Year", "Timing Diff"]
            for i, h in enumerate(headers, 1):
                ws.cell(row=row, column=i, value=h)
            apply_header_row(ws, row, 1, len(headers))
            row += 1

            for ep in cr["ep_results"]:
                ep_date_str = str(ep["ep_date"]) if ep["ep_date"] else "N/A"
                write_data_row(ws, row, [
                    ep["line_id"],
                    float(to_decimal(ep["amount"])),
                    ep["ep_category"],
                    ep["rule_applied"],
                    "Yes" if ep["ep_met"] else "No",
                    ep_date_str,
                    ep["tax_deduction_year"] or "TBD",
                    float(to_decimal(ep["timing_difference"])),
                ], alternate=(row % 2 == 0))
                format_currency_cell(ws.cell(row=row, column=2))
                format_currency_cell(ws.cell(row=row, column=8))
                row += 1
            row += 1

        auto_fit_columns(ws)

    # ------------------------------------------------------------------
    def _create_expense_schedule(self):
        ws = self._wb.create_sheet("Expense Schedule")
        apply_title(ws, 1, 1, 6, "Expense Timing Schedule — Book vs Tax")
        row = 3

        headers = ["Line ID", "Description", "Amount", "Book Accrual Date",
                    "Tax Deduction Date", "Timing Match"]
        for i, h in enumerate(headers, 1):
            ws.cell(row=row, column=i, value=h)
        apply_header_row(ws, row, 1, len(headers))
        row += 1

        for item in self._r["expense_schedule"]:
            book_str = str(item["book_accrual_date"]) if item["book_accrual_date"] else "N/A"
            tax_str = str(item["tax_deduction_date"]) if item["tax_deduction_date"] else "TBD"
            write_data_row(ws, row, [
                item["line_id"], item["description"],
                float(to_decimal(item["amount"])),
                book_str, tax_str,
                "Yes" if item["timing_matches"] else "NO — MISMATCH",
            ], alternate=(row % 2 == 0))
            format_currency_cell(ws.cell(row=row, column=3))
            if not item["timing_matches"]:
                ws.cell(row=row, column=6).font = get_risk_font("high")
            row += 1

        auto_fit_columns(ws)

    # ------------------------------------------------------------------
    def _create_reserve_analysis(self):
        ws = self._wb.create_sheet("Reserve Analysis")
        apply_title(ws, 1, 1, 9, "Reserve & Contingency Analysis")
        row = 3

        headers = ["Reserve", "Type", "Book Balance", "Tax Deductible",
                    "Timing Diff", "All-Events", "EP Met", "EP Category", "EP Rule"]
        for i, h in enumerate(headers, 1):
            ws.cell(row=row, column=i, value=h)
        apply_header_row(ws, row, 1, len(headers))
        row += 1

        for res in self._r["reserves"]["items"]:
            write_data_row(ws, row, [
                res["name"], res["reserve_type"],
                float(to_decimal(res["book_balance"])),
                float(to_decimal(res["tax_deductible"])),
                float(to_decimal(res["timing_difference"])),
                res["all_events_analysis"],
                "Yes" if res["ep_met"] else "No",
                res["ep_category"], res["ep_rule"],
            ], alternate=(row % 2 == 0))
            format_currency_cell(ws.cell(row=row, column=3))
            format_currency_cell(ws.cell(row=row, column=4))
            format_currency_cell(ws.cell(row=row, column=5))
            row += 1

        row += 1
        ws.cell(row=row, column=1, value="Total Timing Difference").font = FONT_BODY_BOLD
        c = ws.cell(row=row, column=5, value=float(self._r["reserves"]["total_timing_difference"]))
        format_currency_cell(c)
        c.font = FONT_BODY_BOLD

        auto_fit_columns(ws)

    # ------------------------------------------------------------------
    def _create_compensation_detail(self):
        ws = self._wb.create_sheet("Compensation Detail")
        apply_title(ws, 1, 1, 10, "Compensation Contract Analysis")
        row = 3

        headers = ["Employee", "Title", "Covered Employee", "Base Salary",
                    "Bonus", "Equity Comp", "Deferred Comp", "Total Book",
                    "§162(m) Disallowed", "Tax Deductible"]
        for i, h in enumerate(headers, 1):
            ws.cell(row=row, column=i, value=h)
        apply_header_row(ws, row, 1, len(headers))
        row += 1

        for c in self._r["compensation"]["items"]:
            write_data_row(ws, row, [
                c["employee"], c["title"],
                "Yes" if c["is_covered"] else "No",
                float(to_decimal(c["base_salary"])),
                float(to_decimal(c["bonus"])),
                float(to_decimal(c["equity_comp"])),
                float(to_decimal(c["deferred_comp"])),
                float(to_decimal(c["total_book"])),
                float(to_decimal(c["section_162m_disallowed"])),
                float(to_decimal(c["tax_deductible"])),
            ], alternate=(row % 2 == 0))
            for col in [4, 5, 6, 7, 8, 9, 10]:
                format_currency_cell(ws.cell(row=row, column=col))
            if to_decimal(c["section_162m_disallowed"]) > 0:
                ws.cell(row=row, column=9).font = get_risk_font("high")
            row += 1

        row += 1
        ws.cell(row=row, column=1, value="Total §162(m) Disallowed").font = FONT_BODY_BOLD
        c = ws.cell(row=row, column=9, value=float(self._r["compensation"]["total_162m_disallowed"]))
        format_currency_cell(c)
        c.font = FONT_BODY_BOLD

        auto_fit_columns(ws)

    # ------------------------------------------------------------------
    def _create_book_tax_reconciliation(self):
        ws = self._wb.create_sheet("Book-Tax Reconciliation")
        apply_title(ws, 1, 1, 5, "Aggregate Book-Tax Reconciliation by Category")
        row = 3

        headers = ["Category", "Book Amount", "Tax Amount", "Difference"]
        for i, h in enumerate(headers, 1):
            ws.cell(row=row, column=i, value=h)
        apply_header_row(ws, row, 1, len(headers))
        row += 1

        recon = self._r["book_tax_reconciliation"]
        for item in recon["items"]:
            write_data_row(ws, row, [
                item["category"],
                float(to_decimal(item["book_amount"])),
                float(to_decimal(item["tax_amount"])),
                float(to_decimal(item["difference"])),
            ], alternate=(row % 2 == 0))
            for col in [2, 3, 4]:
                format_currency_cell(ws.cell(row=row, column=col))
            row += 1

        row += 1
        ws.cell(row=row, column=1, value="TOTAL").font = FONT_BODY_BOLD
        for col, val in [(2, recon["total_book"]), (3, recon["total_tax"]), (4, recon["total_difference"])]:
            c = ws.cell(row=row, column=col, value=float(to_decimal(val)))
            format_currency_cell(c)
            c.font = FONT_BODY_BOLD

        auto_fit_columns(ws)

    # ------------------------------------------------------------------
    def _create_implementation_plan(self):
        ws = self._wb.create_sheet("Implementation Plan")
        apply_title(ws, 1, 1, 6, "Implementation Plan & Method Changes")
        row = 3

        headers = ["Priority", "Action", "Description", "IRC Section",
                    "Form/Document", "§481(a) Adj", "Deadline"]
        for i, h in enumerate(headers, 1):
            ws.cell(row=row, column=i, value=h)
        apply_header_row(ws, row, 1, len(headers))
        row += 1

        for item in self._r["implementation_plan"]:
            write_data_row(ws, row, [
                item["priority"], item["action"], item["description"],
                item["irc_section"], item["form"],
                "Yes" if item["section_481a"] else "No",
                item["deadline"],
            ], alternate=(row % 2 == 0))
            row += 1

        auto_fit_columns(ws)
