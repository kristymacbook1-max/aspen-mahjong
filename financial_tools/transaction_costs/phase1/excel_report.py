"""Phase 1 Transaction Cost Excel Report Generator.

Produces a 7-tab workbook:
  1. Executive Summary
  2. Transaction Portfolio
  3. Cost Classification
  4. Debt Issuance Costs
  5. Abandoned Transactions
  6. Findings
  7. Phase 2 Roadmap
"""

import os, sys
from decimal import Decimal
from openpyxl import Workbook

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', '..'))
from revenue_recognition.utils.excel_styles import (
    FONT_TITLE, FONT_BODY, FONT_BODY_BOLD, THIN_BORDER, FMT_CURRENCY,
    apply_header_row, apply_section_header, apply_title,
    auto_fit_columns, write_data_row, format_currency_cell, get_risk_font,
)
from revenue_recognition.utils.currency_helpers import to_decimal, format_currency


class Phase1TransactionCostReport:

    def generate(self, results: dict, output_path: str) -> str:
        self._r = results
        self._wb = Workbook()
        self._wb.remove(self._wb.active)

        self._create_executive_summary()
        self._create_transactions()
        self._create_cost_classification()
        self._create_debt_issuance()
        self._create_abandoned()
        self._create_findings()
        self._create_roadmap()

        os.makedirs(os.path.dirname(output_path) or '.', exist_ok=True)
        self._wb.save(output_path)
        return output_path

    def _create_executive_summary(self):
        ws = self._wb.create_sheet("Executive Summary")
        apply_title(ws, 1, 1, 5, f"Transaction Cost Analysis — {self._r['company'].name}")
        row = 3

        cc = self._r["cost_classification"]
        info = [
            ("Transactions Analyzed", len(self._r["transaction_results"])),
            ("Total Costs to Capitalize", format_currency(cc["total_capitalize"])),
            ("Total Costs to Deduct", format_currency(cc["total_deduct"])),
            ("Costs Requiring Allocation", format_currency(cc["total_allocate"])),
            ("Classification Mismatches", len(cc["mismatches"])),
        ]
        for label, val in info:
            ws.cell(row=row, column=1, value=label).font = FONT_BODY_BOLD
            ws.cell(row=row, column=2, value=str(val)).font = FONT_BODY
            row += 1

        row += 1
        apply_section_header(ws, row, 1, 5, "Key Findings")
        row += 1
        for f in self._r["findings"][:8]:
            ws.cell(row=row, column=1, value=f["area"]).font = FONT_BODY_BOLD
            ws.cell(row=row, column=2, value=f["description"]).font = FONT_BODY
            ws.cell(row=row, column=3, value=f["risk_level"].upper()).font = get_risk_font(f["risk_level"])
            row += 1
        auto_fit_columns(ws)

    def _create_transactions(self):
        ws = self._wb.create_sheet("Transaction Portfolio")
        apply_title(ws, 1, 1, 7, "Transaction Portfolio Overview")
        row = 3

        headers = ["ID", "Description", "Type", "Value", "Status", "Total Costs", "Capitalize", "Deduct"]
        for i, h in enumerate(headers, 1):
            ws.cell(row=row, column=i, value=h)
        apply_header_row(ws, row, 1, len(headers))
        row += 1

        for tr in self._r["transaction_results"]:
            t = tr["transaction"]
            write_data_row(ws, row, [
                t.transaction_id, t.description, t.transaction_type,
                float(to_decimal(t.total_value)), t.status,
                float(to_decimal(tr["total_costs"])),
                float(to_decimal(tr["amount_to_capitalize"])),
                float(to_decimal(tr["amount_to_deduct"])),
            ], alternate=(row % 2 == 0))
            for c in [4, 6, 7, 8]:
                format_currency_cell(ws.cell(row=row, column=c))
            row += 1
        auto_fit_columns(ws)

    def _create_cost_classification(self):
        ws = self._wb.create_sheet("Cost Classification")
        apply_title(ws, 1, 1, 7, "Cost Classification Under Reg. 1.263(a)-5")
        row = 3

        for section, items in [
            ("CAPITALIZE — Facilitative Costs", self._r["cost_classification"]["capitalize"]),
            ("DEDUCT — Non-Facilitative / Integration", self._r["cost_classification"]["deduct"]),
            ("ALLOCATE — Requires Further Analysis", self._r["cost_classification"]["allocate"]),
        ]:
            apply_section_header(ws, row, 1, 7, section)
            row += 1
            if not items:
                ws.cell(row=row, column=1, value="None").font = FONT_BODY
                row += 2
                continue

            headers = ["Description", "Vendor", "Amount", "Type", "Rule", "Mismatch"]
            for i, h in enumerate(headers, 1):
                ws.cell(row=row, column=i, value=h)
            apply_header_row(ws, row, 1, len(headers))
            row += 1

            for item in items:
                write_data_row(ws, row, [
                    item["description"], item["vendor"],
                    float(to_decimal(item["amount"])), item["cost_type"],
                    item["rule"], "YES" if item["mismatch"] else "",
                ], alternate=(row % 2 == 0))
                format_currency_cell(ws.cell(row=row, column=3))
                if item["mismatch"]:
                    ws.cell(row=row, column=6).font = get_risk_font("high")
                row += 1
            row += 1
        auto_fit_columns(ws)

    def _create_debt_issuance(self):
        ws = self._wb.create_sheet("Debt Issuance Costs")
        apply_title(ws, 1, 1, 5, "Debt Issuance Cost Amortization")
        row = 3

        headers = ["Description", "Amount", "Instrument", "Term (Yrs)", "Annual Amortization"]
        for i, h in enumerate(headers, 1):
            ws.cell(row=row, column=i, value=h)
        apply_header_row(ws, row, 1, len(headers))
        row += 1

        for item in self._r["debt_issuance"]["items"]:
            write_data_row(ws, row, [
                item["description"], float(to_decimal(item["amount"])),
                item["instrument"], item["term"],
                float(to_decimal(item["annual_amortization"])),
            ], alternate=(row % 2 == 0))
            format_currency_cell(ws.cell(row=row, column=2))
            format_currency_cell(ws.cell(row=row, column=5))
            row += 1
        auto_fit_columns(ws)

    def _create_abandoned(self):
        ws = self._wb.create_sheet("Abandoned Transactions")
        apply_title(ws, 1, 1, 4, "Abandoned Transaction — Loss Deduction")
        row = 3

        if not self._r["abandoned_transactions"]["items"]:
            ws.cell(row=row, column=1, value="No abandoned transactions identified.").font = FONT_BODY
        else:
            headers = ["Transaction", "Total Costs", "Treatment", "Authority"]
            for i, h in enumerate(headers, 1):
                ws.cell(row=row, column=i, value=h)
            apply_header_row(ws, row, 1, len(headers))
            row += 1
            for item in self._r["abandoned_transactions"]["items"]:
                write_data_row(ws, row, [
                    item["transaction"].description,
                    float(to_decimal(item["total_costs"])),
                    item["treatment"], item["rule"],
                ], alternate=(row % 2 == 0))
                format_currency_cell(ws.cell(row=row, column=2))
                row += 1
        auto_fit_columns(ws)

    def _create_findings(self):
        ws = self._wb.create_sheet("Findings")
        apply_title(ws, 1, 1, 5, "Transaction Cost Findings")
        row = 3

        headers = ["Area", "Description", "Amount", "Risk", "Recommendation"]
        for i, h in enumerate(headers, 1):
            ws.cell(row=row, column=i, value=h)
        apply_header_row(ws, row, 1, len(headers))
        row += 1

        for f in self._r["findings"]:
            write_data_row(ws, row, [
                f["area"], f["description"], float(to_decimal(f["amount"])),
                f["risk_level"].upper(), f["recommendation"],
            ], alternate=(row % 2 == 0))
            format_currency_cell(ws.cell(row=row, column=3))
            ws.cell(row=row, column=4).font = get_risk_font(f["risk_level"])
            row += 1
        auto_fit_columns(ws)

    def _create_roadmap(self):
        ws = self._wb.create_sheet("Phase 2 Roadmap")
        apply_title(ws, 1, 1, 4, "Phase 2 Recommendations")
        row = 3

        headers = ["Priority", "Area", "Description", "Data Needed"]
        for i, h in enumerate(headers, 1):
            ws.cell(row=row, column=i, value=h)
        apply_header_row(ws, row, 1, len(headers))
        row += 1

        for item in self._r["phase2_roadmap"]:
            write_data_row(ws, row, [
                item["priority"], item["area"], item["description"], item["data_needed"],
            ], alternate=(row % 2 == 0))
            row += 1
        auto_fit_columns(ws)
