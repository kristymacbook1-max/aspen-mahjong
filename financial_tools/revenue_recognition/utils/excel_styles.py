"""Shared Excel styling constants and helper functions for revenue recognition workbooks."""

from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, numbers


# --- Color Palette ---
COLORS = {
    "dark_blue": "1F3864",
    "medium_blue": "2E75B6",
    "light_blue": "D6E4F0",
    "green": "548235",
    "light_green": "E2EFDA",
    "red": "C00000",
    "light_red": "FCE4EC",
    "orange": "ED7D31",
    "light_orange": "FFF2CC",
    "gray": "808080",
    "light_gray": "F2F2F2",
    "white": "FFFFFF",
    "black": "000000",
}

# --- Fonts ---
FONT_TITLE = Font(name="Calibri", size=16, bold=True, color=COLORS["dark_blue"])
FONT_SECTION_HEADER = Font(name="Calibri", size=12, bold=True, color=COLORS["white"])
FONT_COLUMN_HEADER = Font(name="Calibri", size=10, bold=True, color=COLORS["dark_blue"])
FONT_BODY = Font(name="Calibri", size=10, color=COLORS["black"])
FONT_BODY_BOLD = Font(name="Calibri", size=10, bold=True, color=COLORS["black"])
FONT_CURRENCY = Font(name="Calibri", size=10, color=COLORS["black"])
FONT_PERCENT = Font(name="Calibri", size=10, color=COLORS["black"])
FONT_LINK = Font(name="Calibri", size=10, color=COLORS["medium_blue"], underline="single")
FONT_FLAG_HIGH = Font(name="Calibri", size=10, bold=True, color=COLORS["red"])
FONT_FLAG_MEDIUM = Font(name="Calibri", size=10, bold=True, color=COLORS["orange"])
FONT_FLAG_LOW = Font(name="Calibri", size=10, color=COLORS["green"])

# --- Fills ---
FILL_HEADER = PatternFill(start_color=COLORS["dark_blue"], end_color=COLORS["dark_blue"], fill_type="solid")
FILL_SECTION = PatternFill(start_color=COLORS["medium_blue"], end_color=COLORS["medium_blue"], fill_type="solid")
FILL_ALT_ROW = PatternFill(start_color=COLORS["light_gray"], end_color=COLORS["light_gray"], fill_type="solid")
FILL_HIGHLIGHT_GREEN = PatternFill(start_color=COLORS["light_green"], end_color=COLORS["light_green"], fill_type="solid")
FILL_HIGHLIGHT_RED = PatternFill(start_color=COLORS["light_red"], end_color=COLORS["light_red"], fill_type="solid")
FILL_HIGHLIGHT_ORANGE = PatternFill(start_color=COLORS["light_orange"], end_color=COLORS["light_orange"], fill_type="solid")
FILL_HIGHLIGHT_BLUE = PatternFill(start_color=COLORS["light_blue"], end_color=COLORS["light_blue"], fill_type="solid")

# --- Borders ---
THIN_BORDER = Border(
    left=Side(style="thin", color=COLORS["gray"]),
    right=Side(style="thin", color=COLORS["gray"]),
    top=Side(style="thin", color=COLORS["gray"]),
    bottom=Side(style="thin", color=COLORS["gray"]),
)

BOTTOM_BORDER = Border(
    bottom=Side(style="medium", color=COLORS["dark_blue"]),
)

# --- Alignments ---
ALIGN_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
ALIGN_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")
ALIGN_WRAP = Alignment(horizontal="left", vertical="top", wrap_text=True)

# --- Number Formats ---
FMT_CURRENCY = '"$"#,##0'
FMT_CURRENCY_DECIMAL = '"$"#,##0.00'
FMT_PERCENT = "0.0%"
FMT_NUMBER = "#,##0"
FMT_DATE = "MM/DD/YYYY"


def apply_header_row(ws, row, col_start, col_end):
    """Apply header styling to a range of cells in a row."""
    for col in range(col_start, col_end + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = FONT_COLUMN_HEADER
        cell.fill = FILL_HIGHLIGHT_BLUE
        cell.alignment = ALIGN_CENTER
        cell.border = THIN_BORDER


def apply_section_header(ws, row, col_start, col_end, title):
    """Write and style a section header spanning multiple columns."""
    ws.merge_cells(start_row=row, start_column=col_start, end_row=row, end_column=col_end)
    cell = ws.cell(row=row, column=col_start)
    cell.value = title
    cell.font = FONT_SECTION_HEADER
    cell.fill = FILL_SECTION
    cell.alignment = ALIGN_LEFT


def apply_title(ws, row, col_start, col_end, title):
    """Write and style a workbook/sheet title spanning multiple columns."""
    ws.merge_cells(start_row=row, start_column=col_start, end_row=row, end_column=col_end)
    cell = ws.cell(row=row, column=col_start)
    cell.value = title
    cell.font = FONT_TITLE
    cell.alignment = ALIGN_LEFT


def auto_fit_columns(ws, min_width=10, max_width=50):
    """Auto-fit column widths based on content."""
    for column_cells in ws.columns:
        max_length = 0
        col_letter = column_cells[0].column_letter
        for cell in column_cells:
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))
        adjusted = min(max(max_length + 2, min_width), max_width)
        ws.column_dimensions[col_letter].width = adjusted


def write_data_row(ws, row, data, start_col=1, alternate=False):
    """Write a row of data with optional alternate-row shading."""
    for i, value in enumerate(data):
        cell = ws.cell(row=row, column=start_col + i, value=value)
        cell.font = FONT_BODY
        cell.border = THIN_BORDER
        cell.alignment = ALIGN_LEFT
        if alternate:
            cell.fill = FILL_ALT_ROW


def format_currency_cell(cell):
    """Apply currency formatting to a cell."""
    cell.number_format = FMT_CURRENCY
    cell.font = FONT_CURRENCY
    cell.alignment = ALIGN_RIGHT


def format_percent_cell(cell):
    """Apply percentage formatting to a cell."""
    cell.number_format = FMT_PERCENT
    cell.font = FONT_PERCENT
    cell.alignment = ALIGN_RIGHT


def get_risk_font(level):
    """Return the appropriate font for a risk level."""
    level = (level or "").lower()
    if level == "high":
        return FONT_FLAG_HIGH
    elif level == "medium":
        return FONT_FLAG_MEDIUM
    return FONT_FLAG_LOW


def get_risk_fill(level):
    """Return the appropriate fill for a risk level."""
    level = (level or "").lower()
    if level == "high":
        return FILL_HIGHLIGHT_RED
    elif level == "medium":
        return FILL_HIGHLIGHT_ORANGE
    return FILL_HIGHLIGHT_GREEN
