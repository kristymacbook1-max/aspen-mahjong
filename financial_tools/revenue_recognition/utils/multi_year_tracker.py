"""Multi-year persistence layer for tracking revenue recognition analysis across years.

Stores and retrieves analysis results in per-company JSON files, computes
year-over-year trends, tracks temporary difference reversals, and generates
multi-year Excel trend workbooks.
"""

import json
import os
from dataclasses import asdict, is_dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any, Dict, List, Optional

from openpyxl import Workbook

from .currency_helpers import to_decimal, round_currency, format_currency
from .excel_styles import (
    FONT_TITLE,
    FONT_SECTION_HEADER,
    FONT_COLUMN_HEADER,
    FONT_BODY,
    FONT_BODY_BOLD,
    FONT_CURRENCY,
    FILL_HEADER,
    FILL_SECTION,
    FILL_ALT_ROW,
    FILL_HIGHLIGHT_BLUE,
    FILL_HIGHLIGHT_GREEN,
    FILL_HIGHLIGHT_RED,
    THIN_BORDER,
    ALIGN_CENTER,
    ALIGN_LEFT,
    ALIGN_RIGHT,
    FMT_CURRENCY,
    FMT_PERCENT,
    FMT_NUMBER,
    apply_header_row,
    apply_section_header,
    apply_title,
    auto_fit_columns,
    write_data_row,
    format_currency_cell,
    format_percent_cell,
)


# ---------------------------------------------------------------------------
# JSON serialization helpers
# ---------------------------------------------------------------------------

class _AnalysisEncoder(json.JSONEncoder):
    """Custom JSON encoder that handles Decimal, date/datetime, and dataclass objects."""

    def default(self, obj: Any) -> Any:
        if isinstance(obj, Decimal):
            return {"__decimal__": str(obj)}
        if isinstance(obj, (date, datetime)):
            return {"__date__": obj.isoformat()}
        if is_dataclass(obj) and not isinstance(obj, type):
            return {"__dataclass__": type(obj).__name__, **asdict(obj)}
        return super().default(obj)


def _analysis_decoder(obj: Dict) -> Any:
    """Object hook for json.load that restores Decimal and date values."""
    if "__decimal__" in obj:
        try:
            return Decimal(obj["__decimal__"])
        except InvalidOperation:
            return Decimal("0")
    if "__date__" in obj:
        raw = obj["__date__"]
        for fmt in ("%Y-%m-%d", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%dT%H:%M:%S.%f"):
            try:
                parsed = datetime.strptime(raw, fmt)
                return parsed.date() if "T" not in raw else parsed
            except ValueError:
                continue
        return raw
    # Strip the __dataclass__ marker but keep the rest of the dict
    obj.pop("__dataclass__", None)
    return obj


# ---------------------------------------------------------------------------
# MultiYearTracker
# ---------------------------------------------------------------------------

class MultiYearTracker:
    """Persists and tracks revenue recognition analysis across multiple years."""

    def __init__(self, company_name: str, storage_dir: str = "data/tracking"):
        self.company_name = company_name
        self.storage_dir = storage_dir
        self._file_path = os.path.join(
            storage_dir,
            self._safe_filename(company_name) + ".json",
        )
        self._data: Dict = self._load_file()

    # ------------------------------------------------------------------
    # Public API
    # ------------------------------------------------------------------

    def save_year(self, tax_year: int, phase: str, results: dict) -> None:
        """Save analysis results for a specific year and phase.

        Args:
            tax_year: The tax year (e.g. 2024).
            phase: One of 'phase1', 'phase2', or 'phase3'.
            results: The analysis results dict produced by the corresponding analyzer.
        """
        year_key = str(tax_year)
        if year_key not in self._data:
            self._data[year_key] = {}
        self._data[year_key][phase] = results
        self._persist()

    def load_year(self, tax_year: int, phase: str) -> dict:
        """Load previously saved results for a year and phase.

        Args:
            tax_year: The tax year.
            phase: One of 'phase1', 'phase2', or 'phase3'.

        Returns:
            The stored results dict, or an empty dict if not found.
        """
        year_key = str(tax_year)
        return self._data.get(year_key, {}).get(phase, {})

    def get_available_years(self) -> List[int]:
        """Return all years with saved data, sorted ascending."""
        years = []
        for key in self._data:
            try:
                years.append(int(key))
            except ValueError:
                continue
        return sorted(years)

    def compute_trends(
        self,
        metric_name: str,
        years: Optional[List[int]] = None,
    ) -> List[dict]:
        """Track a specific metric across years.

        Searches phase3, phase2, then phase1 metrics for the requested metric.

        Args:
            metric_name: Key name of the metric (e.g. 'total_book_revenue').
            years: Years to include. Defaults to all available years.

        Returns:
            List of {year, value, yoy_change} dicts sorted by year.
        """
        years = self._resolve_years(years)
        results: List[dict] = []
        prev_value = None

        for year in years:
            value = self._find_metric(year, metric_name)
            yoy = None
            if prev_value is not None and value is not None and prev_value != 0:
                yoy = float((to_decimal(value) - to_decimal(prev_value)) / abs(to_decimal(prev_value)))
            results.append({
                "year": year,
                "value": value,
                "yoy_change": yoy,
            })
            if value is not None:
                prev_value = value

        return results

    def compute_reversal_tracking(
        self,
        years: Optional[List[int]] = None,
    ) -> List[dict]:
        """Track how temporary book-tax differences reverse over time.

        For each year shows opening balance, new differences created,
        differences reversed, and closing balance. This helps validate
        DTA/DTL movements.

        Args:
            years: Years to include. Defaults to all available years.

        Returns:
            List of dicts per year with opening_balance, new_differences,
            reversed_differences, and closing_balance.
        """
        years = self._resolve_years(years)
        results: List[dict] = []
        prior_closing = Decimal("0")

        for year in years:
            total_temp = to_decimal(
                self._find_metric(year, "total_temporary_differences")
            )
            # total_temp represents the net new temporary difference for the year.
            # Prior year's closing balance reverses as current-year taxable income
            # catches up to book income (or vice versa).
            # Heuristic: the reversal equals the prior closing balance that is
            # assumed to reverse in the immediately following year (one-year deferral
            # pattern common for deferred revenue under sec 451(c)).
            reversed_amount = prior_closing
            new_differences = total_temp
            closing = new_differences  # after full reversal of prior year balance

            results.append({
                "year": year,
                "opening_balance": prior_closing,
                "new_differences": new_differences,
                "reversed_differences": reversed_amount,
                "closing_balance": closing,
            })
            prior_closing = closing

        return results

    def generate_multi_year_summary(
        self,
        years: Optional[List[int]] = None,
    ) -> dict:
        """Generate a comprehensive multi-year summary.

        Includes revenue growth trend, book-tax difference trend,
        deferred revenue trend, tax savings per year, DTA/DTL balance trend,
        and cumulative section 481(a) adjustment tracking.

        Args:
            years: Years to include. Defaults to all available years.

        Returns:
            Dict with keyed trend lists.
        """
        years = self._resolve_years(years)

        revenue_trend = self.compute_trends("total_book_revenue", years)
        # Fall back to phase1 total revenue if phase3 metric not available
        if all(r["value"] is None for r in revenue_trend):
            revenue_trend = self.compute_trends("total_revenue_tb", years)

        book_tax_trend = self.compute_trends("total_book_tax_difference", years)
        deferred_trend = self.compute_trends("total_deferred_revenue", years)
        tax_savings_trend = self.compute_trends("total_tax_savings_opportunity", years)
        dta_dtl_trend = self.compute_trends("estimated_tax_impact_temporary", years)
        reversal_tracking = self.compute_reversal_tracking(years)

        # Cumulative 481(a) adjustment tracking
        cumulative_481a = self._compute_cumulative_481a(years)

        return {
            "company_name": self.company_name,
            "years": years,
            "revenue_trend": revenue_trend,
            "book_tax_difference_trend": book_tax_trend,
            "deferred_revenue_trend": deferred_trend,
            "tax_savings_per_year": tax_savings_trend,
            "dta_dtl_balance_trend": dta_dtl_trend,
            "reversal_tracking": reversal_tracking,
            "cumulative_481a_tracking": cumulative_481a,
        }

    def generate_trend_excel(self, output_dir: str = "output") -> str:
        """Generate an Excel workbook showing multi-year trends.

        Tabs:
        - Revenue Trend
        - Book-Tax Differences
        - Deferred Revenue Trend
        - Tax Savings Summary
        - DTA/DTL Tracking

        Args:
            output_dir: Directory for the output file.

        Returns:
            Absolute path to the generated Excel file.
        """
        os.makedirs(output_dir, exist_ok=True)

        summary = self.generate_multi_year_summary()
        years = summary["years"]
        if not years:
            # Nothing to write — return path to an empty workbook
            wb = Workbook()
            ws = wb.active
            ws.title = "No Data"
            ws.cell(row=1, column=1, value="No multi-year data available.")
            filepath = os.path.join(
                output_dir,
                f"{self._safe_filename(self.company_name)}_trends.xlsx",
            )
            wb.save(filepath)
            return os.path.abspath(filepath)

        wb = Workbook()
        # Remove default sheet — we'll create our own
        wb.remove(wb.active)

        self._write_revenue_trend_sheet(wb, summary, years)
        self._write_book_tax_sheet(wb, summary, years)
        self._write_deferred_revenue_sheet(wb, summary, years)
        self._write_tax_savings_sheet(wb, summary, years)
        self._write_dta_dtl_sheet(wb, summary, years)

        filepath = os.path.join(
            output_dir,
            f"{self._safe_filename(self.company_name)}_trends.xlsx",
        )
        wb.save(filepath)
        return os.path.abspath(filepath)

    # ------------------------------------------------------------------
    # Excel sheet builders
    # ------------------------------------------------------------------

    def _write_revenue_trend_sheet(
        self, wb: Workbook, summary: dict, years: List[int]
    ) -> None:
        ws = wb.create_sheet("Revenue Trend")
        row = 1
        apply_title(ws, row, 1, len(years) + 2, f"{self.company_name} — Revenue Trend")
        row += 2

        # Header row
        headers = ["Metric"] + [str(y) for y in years]
        for ci, h in enumerate(headers, start=1):
            cell = ws.cell(row=row, column=ci, value=h)
            cell.font = FONT_COLUMN_HEADER
            cell.fill = FILL_HIGHLIGHT_BLUE
            cell.alignment = ALIGN_CENTER
            cell.border = THIN_BORDER
        row += 1

        # Total revenue row
        row = self._write_trend_row(
            ws, row, "Total Revenue",
            summary["revenue_trend"], years, currency=True,
        )

        # Revenue by stream — pull from phase1 per-year data
        stream_names = self._collect_revenue_stream_names(years)
        for stream_name in stream_names:
            stream_values = []
            for year in years:
                val = self._find_stream_amount(year, stream_name)
                stream_values.append(val)
            row = self._write_values_row(
                ws, row, f"  {stream_name}", stream_values, years, currency=True,
            )

        # YoY growth row
        row = self._write_trend_row(
            ws, row, "YoY Growth",
            summary["revenue_trend"], years, growth=True,
        )

        auto_fit_columns(ws)

    def _write_book_tax_sheet(
        self, wb: Workbook, summary: dict, years: List[int]
    ) -> None:
        ws = wb.create_sheet("Book-Tax Differences")
        row = 1
        apply_title(ws, row, 1, len(years) + 2, f"{self.company_name} — Book-Tax Differences")
        row += 2

        headers = ["Metric"] + [str(y) for y in years]
        for ci, h in enumerate(headers, start=1):
            cell = ws.cell(row=row, column=ci, value=h)
            cell.font = FONT_COLUMN_HEADER
            cell.fill = FILL_HIGHLIGHT_BLUE
            cell.alignment = ALIGN_CENTER
            cell.border = THIN_BORDER
        row += 1

        row = self._write_trend_row(
            ws, row, "Net Book-Tax Difference",
            summary["book_tax_difference_trend"], years, currency=True,
        )

        # Temporary vs permanent
        temp_trend = self.compute_trends("total_temporary_differences", years)
        perm_trend = self.compute_trends("total_permanent_differences", years)
        row = self._write_trend_row(ws, row, "Temporary Differences", temp_trend, years, currency=True)
        row = self._write_trend_row(ws, row, "Permanent Differences", perm_trend, years, currency=True)

        # Reversal tracking section
        row += 1
        apply_section_header(ws, row, 1, len(years) + 1, "Reversal Tracking")
        row += 1

        reversal_metrics = [
            ("Opening Balance", "opening_balance"),
            ("New Differences", "new_differences"),
            ("Reversed", "reversed_differences"),
            ("Closing Balance", "closing_balance"),
        ]
        for label, key in reversal_metrics:
            vals = [entry.get(key) for entry in summary["reversal_tracking"]]
            row = self._write_values_row(ws, row, label, vals, years, currency=True)

        auto_fit_columns(ws)

    def _write_deferred_revenue_sheet(
        self, wb: Workbook, summary: dict, years: List[int]
    ) -> None:
        ws = wb.create_sheet("Deferred Revenue Trend")
        row = 1
        apply_title(ws, row, 1, len(years) + 2, f"{self.company_name} — Deferred Revenue Trend")
        row += 2

        headers = ["Metric"] + [str(y) for y in years]
        for ci, h in enumerate(headers, start=1):
            cell = ws.cell(row=row, column=ci, value=h)
            cell.font = FONT_COLUMN_HEADER
            cell.fill = FILL_HIGHLIGHT_BLUE
            cell.alignment = ALIGN_CENTER
            cell.border = THIN_BORDER
        row += 1

        row = self._write_trend_row(
            ws, row, "Deferred Revenue Balance",
            summary["deferred_revenue_trend"], years, currency=True,
        )
        row = self._write_trend_row(
            ws, row, "YoY Change",
            summary["deferred_revenue_trend"], years, growth=True,
        )

        # Additions and recognition from rollforward data
        additions_trend = self.compute_trends("deferred_rev_additions", years)
        recognition_trend = self.compute_trends("deferred_rev_recognized", years)
        if any(r["value"] is not None for r in additions_trend):
            row = self._write_trend_row(ws, row, "Additions", additions_trend, years, currency=True)
        if any(r["value"] is not None for r in recognition_trend):
            row = self._write_trend_row(ws, row, "Recognition", recognition_trend, years, currency=True)

        auto_fit_columns(ws)

    def _write_tax_savings_sheet(
        self, wb: Workbook, summary: dict, years: List[int]
    ) -> None:
        ws = wb.create_sheet("Tax Savings Summary")
        row = 1
        apply_title(ws, row, 1, len(years) + 2, f"{self.company_name} — Tax Savings Summary")
        row += 2

        headers = ["Metric"] + [str(y) for y in years]
        for ci, h in enumerate(headers, start=1):
            cell = ws.cell(row=row, column=ci, value=h)
            cell.font = FONT_COLUMN_HEADER
            cell.fill = FILL_HIGHLIGHT_BLUE
            cell.alignment = ALIGN_CENTER
            cell.border = THIN_BORDER
        row += 1

        row = self._write_trend_row(
            ws, row, "Tax Savings Achieved",
            summary["tax_savings_per_year"], years, currency=True,
        )

        # Cumulative savings
        cumulative = []
        running = Decimal("0")
        for entry in summary["tax_savings_per_year"]:
            val = to_decimal(entry.get("value"))
            running += val
            cumulative.append(running)
        row = self._write_values_row(ws, row, "Cumulative Savings", cumulative, years, currency=True)

        # 481(a) tracking
        if summary.get("cumulative_481a_tracking"):
            row += 1
            apply_section_header(ws, row, 1, len(years) + 1, "Cumulative Section 481(a) Adjustment")
            row += 1
            for label_key in [
                ("Annual 481(a) Spread", "annual_spread"),
                ("Cumulative Recognized", "cumulative_recognized"),
                ("Remaining Balance", "remaining_balance"),
            ]:
                label, key = label_key
                vals = [entry.get(key) for entry in summary["cumulative_481a_tracking"]]
                row = self._write_values_row(ws, row, label, vals, years, currency=True)

        auto_fit_columns(ws)

    def _write_dta_dtl_sheet(
        self, wb: Workbook, summary: dict, years: List[int]
    ) -> None:
        ws = wb.create_sheet("DTA-DTL Tracking")
        row = 1
        apply_title(ws, row, 1, len(years) + 2, f"{self.company_name} — DTA/DTL Balance Tracking")
        row += 2

        headers = ["Metric"] + [str(y) for y in years]
        for ci, h in enumerate(headers, start=1):
            cell = ws.cell(row=row, column=ci, value=h)
            cell.font = FONT_COLUMN_HEADER
            cell.fill = FILL_HIGHLIGHT_BLUE
            cell.alignment = ALIGN_CENTER
            cell.border = THIN_BORDER
        row += 1

        row = self._write_trend_row(
            ws, row, "DTA/DTL from Temp Differences",
            summary["dta_dtl_balance_trend"], years, currency=True,
        )
        row = self._write_trend_row(
            ws, row, "YoY Movement",
            summary["dta_dtl_balance_trend"], years, growth=True,
        )

        auto_fit_columns(ws)

    # ------------------------------------------------------------------
    # Row-writing helpers
    # ------------------------------------------------------------------

    def _write_trend_row(
        self,
        ws,
        row: int,
        label: str,
        trend: List[dict],
        years: List[int],
        currency: bool = False,
        growth: bool = False,
    ) -> int:
        """Write a single data row from a trend list. Returns next row number."""
        cell = ws.cell(row=row, column=1, value=label)
        cell.font = FONT_BODY_BOLD
        cell.border = THIN_BORDER

        for ci, entry in enumerate(trend, start=2):
            if growth:
                raw = entry.get("yoy_change")
            else:
                raw = entry.get("value")

            value = self._to_excel_value(raw)
            cell = ws.cell(row=row, column=ci, value=value)
            cell.border = THIN_BORDER

            if currency and value is not None:
                format_currency_cell(cell)
            elif growth and value is not None:
                format_percent_cell(cell)
            else:
                cell.font = FONT_BODY
                cell.alignment = ALIGN_RIGHT

        return row + 1

    def _write_values_row(
        self,
        ws,
        row: int,
        label: str,
        values: list,
        years: List[int],
        currency: bool = False,
    ) -> int:
        """Write a row of raw values. Returns next row number."""
        cell = ws.cell(row=row, column=1, value=label)
        cell.font = FONT_BODY_BOLD
        cell.border = THIN_BORDER

        for ci, val in enumerate(values, start=2):
            value = self._to_excel_value(val)
            cell = ws.cell(row=row, column=ci, value=value)
            cell.border = THIN_BORDER
            if currency and value is not None:
                format_currency_cell(cell)
            else:
                cell.font = FONT_BODY
                cell.alignment = ALIGN_RIGHT

        return row + 1

    # ------------------------------------------------------------------
    # Internal helpers
    # ------------------------------------------------------------------

    def _resolve_years(self, years: Optional[List[int]]) -> List[int]:
        if years is not None:
            return sorted(years)
        return self.get_available_years()

    def _find_metric(self, year: int, metric_name: str) -> Any:
        """Search through phase3 -> phase2 -> phase1 metrics for a value."""
        year_key = str(year)
        year_data = self._data.get(year_key, {})
        for phase in ("phase3", "phase2", "phase1"):
            phase_data = year_data.get(phase, {})
            metrics = phase_data.get("metrics", {})
            if metric_name in metrics:
                return metrics[metric_name]
        return None

    def _find_stream_amount(self, year: int, stream_name: str) -> Any:
        """Find a revenue stream amount from phase1 data for a given year."""
        year_key = str(year)
        year_data = self._data.get(year_key, {})
        phase1 = year_data.get("phase1", {})
        streams = phase1.get("revenue_streams", [])
        for stream in streams:
            name = stream.get("name", "") if isinstance(stream, dict) else getattr(stream, "name", "")
            if name == stream_name:
                if isinstance(stream, dict):
                    return stream.get("amount_current")
                return getattr(stream, "amount_current", None)
        return None

    def _collect_revenue_stream_names(self, years: List[int]) -> List[str]:
        """Collect unique revenue stream names across all years."""
        names_seen = []
        for year in years:
            year_key = str(year)
            phase1 = self._data.get(year_key, {}).get("phase1", {})
            for stream in phase1.get("revenue_streams", []):
                name = stream.get("name", "") if isinstance(stream, dict) else getattr(stream, "name", "")
                if name and name not in names_seen:
                    names_seen.append(name)
        return names_seen

    def _compute_cumulative_481a(self, years: List[int]) -> List[dict]:
        """Track cumulative section 481(a) adjustment recognition.

        Looks for a '481a_total_adjustment' and '481a_spread_period' metric
        in each year's data. The adjustment is typically spread over 4 years
        (for positive adjustments) per IRC section 481(a).
        """
        results: List[dict] = []
        # Collect all 481(a) adjustments with their spread periods
        pending_spreads: List[dict] = []

        for year in years:
            total_adj = self._find_metric(year, "481a_total_adjustment")
            spread_period = self._find_metric(year, "481a_spread_period")

            if total_adj is not None and to_decimal(total_adj) != 0:
                period = int(spread_period) if spread_period else 4
                pending_spreads.append({
                    "start_year": year,
                    "total": to_decimal(total_adj),
                    "period": period,
                    "annual": round_currency(to_decimal(total_adj) / Decimal(str(period))),
                })

            annual_spread = Decimal("0")
            for spread in pending_spreads:
                year_offset = year - spread["start_year"]
                if 0 <= year_offset < spread["period"]:
                    annual_spread += spread["annual"]

            cumulative_recognized = Decimal("0")
            for spread in pending_spreads:
                years_elapsed = min(year - spread["start_year"] + 1, spread["period"])
                if years_elapsed > 0:
                    cumulative_recognized += spread["annual"] * Decimal(str(years_elapsed))

            total_all_adjustments = sum(s["total"] for s in pending_spreads)
            remaining = total_all_adjustments - cumulative_recognized

            results.append({
                "year": year,
                "annual_spread": annual_spread,
                "cumulative_recognized": cumulative_recognized,
                "remaining_balance": remaining,
            })

        return results

    @staticmethod
    def _to_excel_value(raw: Any) -> Any:
        """Convert a raw value to something openpyxl can write."""
        if raw is None:
            return None
        if isinstance(raw, Decimal):
            return float(raw)
        if isinstance(raw, (int, float)):
            return raw
        # Try numeric conversion
        try:
            return float(Decimal(str(raw)))
        except (InvalidOperation, ValueError):
            return raw

    @staticmethod
    def _safe_filename(name: str) -> str:
        """Convert a company name to a safe filename fragment."""
        safe = "".join(c if c.isalnum() or c in (" ", "-", "_") else "_" for c in name)
        return safe.strip().replace(" ", "_").lower()

    def _load_file(self) -> dict:
        """Load existing tracking data from disk, or return empty dict."""
        if os.path.isfile(self._file_path):
            with open(self._file_path, "r", encoding="utf-8") as f:
                return json.load(f, object_hook=_analysis_decoder)
        return {}

    def _persist(self) -> None:
        """Write current tracking data to disk."""
        os.makedirs(os.path.dirname(self._file_path), exist_ok=True)
        with open(self._file_path, "w", encoding="utf-8") as f:
            json.dump(self._data, f, cls=_AnalysisEncoder, indent=2)
