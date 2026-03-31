"""Generate Excel workbook for technical authority analysis results."""

import os
from collections import Counter
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Font

from ..utils.excel_styles import (
    apply_title,
    apply_section_header,
    apply_header_row,
    write_data_row,
    auto_fit_columns,
    format_currency_cell,
    FONT_BODY,
    FONT_BODY_BOLD,
    FONT_TITLE,
    FILL_HIGHLIGHT_BLUE,
    FILL_HIGHLIGHT_GREEN,
    FILL_HIGHLIGHT_ORANGE,
    FILL_HIGHLIGHT_RED,
    THIN_BORDER,
    ALIGN_LEFT,
    ALIGN_WRAP,
    ALIGN_CENTER,
    get_risk_font,
    get_risk_fill,
)
from ..utils.currency_helpers import to_decimal

CONFIDENCE_FILLS = {
    "should": FILL_HIGHLIGHT_GREEN,
    "more_likely_than_not": FILL_HIGHLIGHT_ORANGE,
    "reasonable_basis": FILL_HIGHLIGHT_RED,
}

CONFIDENCE_LABELS = {
    "should": "Should",
    "more_likely_than_not": "More Likely Than Not",
    "reasonable_basis": "Reasonable Basis",
}


class TechnicalAuthorityReport:
    """Generates a multi-tab Excel workbook for a technical authority memo."""

    def __init__(self, memo, output_dir="output"):
        self.memo = memo
        self.output_dir = output_dir

    def generate(self) -> str:
        """Generate the Excel workbook and return the file path."""
        os.makedirs(self.output_dir, exist_ok=True)
        wb = Workbook()

        self._build_executive_summary(wb.active)
        self._build_technical_positions(wb.create_sheet("Technical Positions"))
        self._build_authority_analysis(wb.create_sheet("Authority Analysis"))
        self._build_technical_memo(wb.create_sheet("Technical Memo"))
        self._build_filing_requirements(wb.create_sheet("Filing Requirements"))

        safe_name = self.memo.taxpayer_name.replace(" ", "_").replace("/", "_")
        filename = f"technical_authority_{safe_name}_{self.memo.tax_year}.xlsx"
        filepath = os.path.join(self.output_dir, filename)
        wb.save(filepath)
        return filepath

    def _build_executive_summary(self, ws):
        ws.title = "Executive Summary"
        memo = self.memo
        apply_title(ws, 1, 1, 4, "Technical Authority Analysis — Executive Summary")
        ws.row_dimensions[1].height = 30

        info = [
            ("Taxpayer:", memo.taxpayer_name),
            ("Tax Year:", str(memo.tax_year)),
            ("Prepared By:", memo.prepared_by),
            ("Date:", memo.date),
        ]
        row = 3
        for label, value in info:
            ws.cell(row=row, column=1, value=label).font = FONT_BODY_BOLD
            ws.cell(row=row, column=2, value=value).font = FONT_BODY
            row += 1

        row += 1
        apply_section_header(ws, row, 1, 4, "Overall Recommendation")
        row += 1
        cell = ws.cell(row=row, column=1, value=memo.overall_recommendation)
        cell.font = FONT_BODY
        cell.alignment = ALIGN_WRAP
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        ws.row_dimensions[row].height = 60

        row += 2
        apply_section_header(ws, row, 1, 4, "Positions by Confidence Level")
        row += 1
        counts = Counter(p.confidence_level for p in memo.positions)
        headers = ["Confidence Level", "Count"]
        ws.cell(row=row, column=1, value=headers[0])
        ws.cell(row=row, column=2, value=headers[1])
        apply_header_row(ws, row, 1, 2)
        row += 1
        for key in ("should", "more_likely_than_not", "reasonable_basis"):
            label = CONFIDENCE_LABELS.get(key, key)
            cnt = counts.get(key, 0)
            write_data_row(ws, row, [label, cnt])
            ws.cell(row=row, column=1).fill = CONFIDENCE_FILLS.get(key, FILL_HIGHLIGHT_BLUE)
            row += 1

        row += 1
        apply_section_header(ws, row, 1, 4, "Total Estimated Tax Impact")
        row += 1
        total = sum((to_decimal(p.tax_impact) for p in memo.positions), Decimal("0"))
        cell = ws.cell(row=row, column=1, value=float(total))
        format_currency_cell(cell)
        cell.font = Font(name="Calibri", size=14, bold=True)

        auto_fit_columns(ws)

    def _build_technical_positions(self, ws):
        headers = [
            "Position ID", "Title", "Issue", "Conclusion", "Confidence Level",
            "Tax Impact", "Risk Assessment", "Disclosure Required",
            "Method Change Required", "Recommended Action",
        ]
        apply_title(ws, 1, 1, len(headers), "Technical Positions")
        ws.row_dimensions[1].height = 30

        row = 3
        for col_idx, h in enumerate(headers, 1):
            ws.cell(row=row, column=col_idx, value=h)
        apply_header_row(ws, row, 1, len(headers))

        for pos in self.memo.positions:
            row += 1
            yn = lambda v: "Yes" if v else "No"
            data = [
                pos.position_id,
                pos.title,
                pos.issue,
                pos.conclusion,
                CONFIDENCE_LABELS.get(pos.confidence_level, pos.confidence_level),
                float(to_decimal(pos.tax_impact)),
                pos.risk_assessment,
                yn(pos.disclosure_required),
                yn(pos.method_change_required),
                pos.recommended_action,
            ]
            write_data_row(ws, row, data, alternate=(row % 2 == 0))

            # Wrap long text columns
            for col in (3, 4, 10):
                ws.cell(row=row, column=col).alignment = ALIGN_WRAP

            # Confidence level coloring
            conf_cell = ws.cell(row=row, column=5)
            fill = CONFIDENCE_FILLS.get(pos.confidence_level)
            if fill:
                conf_cell.fill = fill
            conf_cell.alignment = ALIGN_CENTER

            # Currency format for tax impact
            format_currency_cell(ws.cell(row=row, column=6))

            # Risk styling
            risk_cell = ws.cell(row=row, column=7)
            risk_cell.font = get_risk_font(pos.risk_assessment)
            risk_cell.fill = get_risk_fill(pos.risk_assessment)

            ws.row_dimensions[row].height = 45

        auto_fit_columns(ws)

    def _build_authority_analysis(self, ws):
        apply_title(ws, 1, 1, 4, "Authority Analysis")
        ws.row_dimensions[1].height = 30
        row = 3

        for pos in self.memo.positions:
            apply_section_header(ws, row, 1, 4, pos.title)
            row += 2

            # Supporting authorities
            ws.cell(row=row, column=1, value="Supporting Authorities").font = FONT_BODY_BOLD
            row += 1
            if pos.authorities_supporting:
                for i, cite in enumerate(pos.authorities_supporting, 1):
                    cell = ws.cell(row=row, column=1, value=f"{i}. {cite}")
                    cell.font = FONT_BODY
                    cell.alignment = ALIGN_WRAP
                    cell.border = THIN_BORDER
                    cell.fill = FILL_HIGHLIGHT_GREEN
                    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
                    row += 1
            else:
                ws.cell(row=row, column=1, value="None identified").font = FONT_BODY
                row += 1

            row += 1
            # Adverse authorities
            ws.cell(row=row, column=1, value="Adverse Authorities").font = FONT_BODY_BOLD
            row += 1
            if pos.authorities_adverse:
                for i, cite in enumerate(pos.authorities_adverse, 1):
                    cell = ws.cell(row=row, column=1, value=f"{i}. {cite}")
                    cell.font = FONT_BODY
                    cell.alignment = ALIGN_WRAP
                    cell.border = THIN_BORDER
                    cell.fill = FILL_HIGHLIGHT_RED
                    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
                    row += 1
            else:
                ws.cell(row=row, column=1, value="None identified").font = FONT_BODY
                row += 1

            row += 2

        auto_fit_columns(ws)

    def _build_technical_memo(self, ws):
        ws.column_dimensions["A"].width = 100
        row = 1

        # Title
        cell = ws.cell(row=row, column=1, value=self.memo.title)
        cell.font = FONT_TITLE
        ws.row_dimensions[row].height = 36
        row += 2

        # Header block
        header_lines = [
            ("To:", self.memo.taxpayer_name),
            ("From:", self.memo.prepared_by),
            ("Date:", self.memo.date),
            ("Re:", f"Technical Authority Analysis — Tax Year {self.memo.tax_year}"),
        ]
        for label, value in header_lines:
            cell = ws.cell(row=row, column=1, value=f"{label}  {value}")
            cell.font = FONT_BODY_BOLD
            cell.alignment = ALIGN_LEFT
            row += 1

        row += 1
        # ISSUES
        cell = ws.cell(row=row, column=1, value="ISSUES")
        cell.font = Font(name="Calibri", size=12, bold=True)
        row += 1
        for i, issue in enumerate(self.memo.issues, 1):
            cell = ws.cell(row=row, column=1, value=f"{i}. {issue}")
            cell.font = FONT_BODY
            cell.alignment = ALIGN_WRAP
            ws.row_dimensions[row].height = 30
            row += 1

        row += 1
        # FACTS
        cell = ws.cell(row=row, column=1, value="FACTS")
        cell.font = Font(name="Calibri", size=12, bold=True)
        row += 1
        cell = ws.cell(row=row, column=1, value=self.memo.facts)
        cell.font = FONT_BODY
        cell.alignment = ALIGN_WRAP
        ws.row_dimensions[row].height = 80
        row += 2

        # Per-position analysis
        for pos in self.memo.positions:
            cell = ws.cell(row=row, column=1, value=f"ANALYSIS — {pos.title}")
            cell.font = Font(name="Calibri", size=12, bold=True)
            row += 1
            cell = ws.cell(row=row, column=1, value=pos.analysis)
            cell.font = FONT_BODY
            cell.alignment = ALIGN_WRAP
            ws.row_dimensions[row].height = 120
            row += 2

            cell = ws.cell(row=row, column=1, value="CONCLUSION")
            cell.font = Font(name="Calibri", size=11, bold=True)
            row += 1
            cell = ws.cell(row=row, column=1, value=pos.conclusion)
            cell.font = FONT_BODY
            cell.alignment = ALIGN_WRAP
            ws.row_dimensions[row].height = 45
            row += 2

        # Overall recommendation
        cell = ws.cell(row=row, column=1, value="OVERALL RECOMMENDATION")
        cell.font = Font(name="Calibri", size=12, bold=True)
        row += 1
        cell = ws.cell(row=row, column=1, value=self.memo.overall_recommendation)
        cell.font = FONT_BODY
        cell.alignment = ALIGN_WRAP
        ws.row_dimensions[row].height = 60
        row += 2

        # Filing requirements
        cell = ws.cell(row=row, column=1, value="FILING REQUIREMENTS")
        cell.font = Font(name="Calibri", size=12, bold=True)
        row += 1
        for req in self.memo.filing_requirements:
            cell = ws.cell(row=row, column=1, value=f"\u2022 {req}")
            cell.font = FONT_BODY
            cell.alignment = ALIGN_WRAP
            ws.row_dimensions[row].height = 25
            row += 1

    def _build_filing_requirements(self, ws):
        headers = ["Requirement", "Description", "Deadline", "Related Position"]
        apply_title(ws, 1, 1, len(headers), "Filing Requirements")
        ws.row_dimensions[1].height = 30

        row = 3
        for col_idx, h in enumerate(headers, 1):
            ws.cell(row=row, column=col_idx, value=h)
        apply_header_row(ws, row, 1, len(headers))

        entries = []
        for pos in self.memo.positions:
            if pos.method_change_required:
                entries.append((
                    "Form 3115",
                    "Application for Change in Accounting Method",
                    f"Filed with tax return for {self.memo.tax_year}",
                    pos.position_id,
                ))
            if pos.form_8275_needed:
                entries.append((
                    "Form 8275",
                    "Disclosure Statement — adequate disclosure of position",
                    f"Filed with tax return for {self.memo.tax_year}",
                    pos.position_id,
                ))
            if pos.disclosure_required and not pos.form_8275_needed:
                entries.append((
                    "Form 8275-R",
                    "Regulation Disclosure Statement — contrary to regulation",
                    f"Filed with tax return for {self.memo.tax_year}",
                    pos.position_id,
                ))

        for entry in entries:
            row += 1
            write_data_row(ws, row, list(entry), alternate=(row % 2 == 0))
            ws.row_dimensions[row].height = 30

        auto_fit_columns(ws)
