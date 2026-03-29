"""Phase 1: Excel report generator for 10-K revenue recognition analysis."""

import os
from datetime import datetime
from openpyxl import Workbook

from ..utils.excel_styles import (
    apply_title, apply_section_header, apply_header_row,
    write_data_row, auto_fit_columns, format_currency_cell, format_percent_cell,
    get_risk_font, get_risk_fill,
    FONT_BODY, FONT_BODY_BOLD, FONT_LINK, FONT_TITLE,
    FILL_HIGHLIGHT_BLUE, FILL_HIGHLIGHT_GREEN, FILL_ALT_ROW,
    THIN_BORDER, ALIGN_CENTER, ALIGN_LEFT, ALIGN_RIGHT, ALIGN_WRAP,
    FMT_CURRENCY, FMT_PERCENT,
)


class Phase1ExcelReport:
    """Generates the Phase 1 Excel workbook from analyzer results."""

    def __init__(self, analysis_results: dict, output_dir: str = "output"):
        self.results = analysis_results
        self.output_dir = output_dir
        self.wb = Workbook()

    def generate(self) -> str:
        """Generate the full workbook and return the file path."""
        self._create_executive_summary()
        self._create_company_overview()
        self._create_revenue_streams()
        self._create_key_metrics()
        self._create_opportunities()
        self._create_risks()
        self._create_deferred_revenue()
        self._create_phase2_roadmap()

        # Remove default sheet if extra
        if "Sheet" in self.wb.sheetnames and len(self.wb.sheetnames) > 1:
            del self.wb["Sheet"]

        os.makedirs(self.output_dir, exist_ok=True)
        company_name = self.results["company"].name.replace(" ", "_")
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"RevRec_Phase1_{company_name}_{timestamp}.xlsx"
        filepath = os.path.join(self.output_dir, filename)
        self.wb.save(filepath)
        return filepath

    def _create_executive_summary(self):
        """Create the executive summary sheet."""
        ws = self.wb.active
        ws.title = "Executive Summary"
        company = self.results["company"]
        metrics = self.results["metrics"]

        apply_title(ws, 1, 1, 6, f"Revenue Recognition Analysis — {company.name}")
        ws.cell(row=2, column=1, value=f"Phase 1: 10-K Public Information Review")
        ws.cell(row=2, column=1).font = FONT_BODY
        ws.cell(row=3, column=1, value=f"Generated: {datetime.now().strftime('%B %d, %Y')}")
        ws.cell(row=3, column=1).font = FONT_BODY

        # Summary stats
        row = 5
        apply_section_header(ws, row, 1, 6, "Analysis Summary")
        row += 1

        summary_data = [
            ("Overall Risk Assessment", metrics.get("overall_risk_assessment", "N/A")),
            ("Opportunities Identified", metrics.get("opportunity_count", 0)),
            ("Risks Identified", metrics.get("risk_count", 0)),
            ("High Priority Items", metrics.get("high_priority_items", 0)),
            ("Revenue Growth (YoY)", self._fmt_pct(metrics.get("revenue_growth_yoy"))),
            ("Effective Tax Rate", self._fmt_pct(metrics.get("effective_tax_rate"))),
            ("Deferred Rev / Revenue", self._fmt_pct(metrics.get("deferred_rev_to_revenue"))),
            ("Days Sales Outstanding", f"{metrics.get('dso', 0):.0f} days" if metrics.get("dso") else "N/A"),
        ]

        for label, value in summary_data:
            ws.cell(row=row, column=1, value=label).font = FONT_BODY_BOLD
            cell = ws.cell(row=row, column=2, value=value)
            cell.font = FONT_BODY
            if label == "Overall Risk Assessment":
                cell.font = get_risk_font(str(value))
                cell.fill = get_risk_fill(str(value))
            row += 1

        # Top opportunities
        row += 1
        apply_section_header(ws, row, 1, 6, "Top Opportunities for Phase 2 Investigation")
        row += 1
        headers = ["#", "Category", "Finding", "Risk Level", "Phase 2 Action"]
        for i, h in enumerate(headers):
            ws.cell(row=row, column=i + 1, value=h)
        apply_header_row(ws, row, 1, len(headers))
        row += 1

        high_opps = [o for o in self.results["opportunities"] if o["risk_level"] == "High"]
        for idx, opp in enumerate(high_opps[:5], 1):
            data = [idx, opp["category"], opp["finding"], opp["risk_level"], opp["phase2_action"]]
            write_data_row(ws, row, data, alternate=(idx % 2 == 0))
            ws.cell(row=row, column=4).font = get_risk_font(opp["risk_level"])
            row += 1

        auto_fit_columns(ws)
        ws.sheet_properties.tabColor = "1F3864"

    def _create_company_overview(self):
        """Create company overview sheet."""
        ws = self.wb.create_sheet("Company Overview")
        company = self.results["company"]

        apply_title(ws, 1, 1, 4, "Company Overview")
        row = 3

        fields = [
            ("Company Name", company.name),
            ("Ticker", company.ticker or "N/A"),
            ("CIK", company.cik or "N/A"),
            ("SIC Code", company.sic_code or "N/A"),
            ("Industry", company.industry),
            ("Fiscal Year End", f"Month {company.fiscal_year_end_month}"),
            ("Filing Date", company.filing_date or "N/A"),
            ("Currency", company.reporting_currency),
        ]

        for label, value in fields:
            ws.cell(row=row, column=1, value=label).font = FONT_BODY_BOLD
            ws.cell(row=row, column=2, value=value).font = FONT_BODY
            row += 1

        # Auditor info
        row += 1
        apply_section_header(ws, row, 1, 4, "Audit Information")
        row += 1
        ws.cell(row=row, column=1, value="Auditor").font = FONT_BODY_BOLD
        ws.cell(row=row, column=2, value=self.results.get("company", {}).name if hasattr(self.results, "get") else "").font = FONT_BODY
        row += 1
        ws.cell(row=row, column=1, value="Opinion Type").font = FONT_BODY_BOLD

        # Disclosures summary
        row += 2
        apply_section_header(ws, row, 1, 4, "Revenue Recognition Policy Summary")
        row += 1
        disc = self.results["disclosures"]
        ws.cell(row=row, column=1, value=disc.asc606_policy_summary or "Not provided").font = FONT_BODY
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        ws.cell(row=row, column=1).alignment = ALIGN_WRAP
        ws.row_dimensions[row].height = 60

        auto_fit_columns(ws)

    def _create_revenue_streams(self):
        """Create revenue streams detail sheet."""
        ws = self.wb.create_sheet("Revenue Streams")
        apply_title(ws, 1, 1, 8, "Revenue Stream Analysis")

        row = 3
        headers = [
            "Revenue Stream", "Recognition Method",
            "Current Year", "Prior Year", "2 Years Prior",
            "YoY Growth", "Customer Concentration", "Notes"
        ]
        for i, h in enumerate(headers):
            ws.cell(row=row, column=i + 1, value=h)
        apply_header_row(ws, row, 1, len(headers))
        row += 1

        for idx, stream in enumerate(self.results["revenue_streams"]):
            growth = None
            if stream.amount_prior and stream.amount_prior != 0:
                growth = float((stream.amount_current - stream.amount_prior) / abs(stream.amount_prior))

            data = [
                stream.name,
                stream.recognition_method,
                float(stream.amount_current),
                float(stream.amount_prior),
                float(stream.amount_two_years_prior),
                growth,
                stream.customer_concentration_pct,
                stream.notes,
            ]
            write_data_row(ws, row, data, alternate=(idx % 2 == 0))

            # Format currency columns
            for col in [3, 4, 5]:
                format_currency_cell(ws.cell(row=row, column=col))
            # Format percent columns
            for col in [6, 7]:
                if ws.cell(row=row, column=col).value is not None:
                    format_percent_cell(ws.cell(row=row, column=col))

            row += 1

        # Geographic breakdown if available
        row += 1
        has_geo = any(s.geographic_breakdown for s in self.results["revenue_streams"])
        if has_geo:
            apply_section_header(ws, row, 1, 8, "Geographic Breakdown")
            row += 1
            for stream in self.results["revenue_streams"]:
                if stream.geographic_breakdown:
                    ws.cell(row=row, column=1, value=stream.name).font = FONT_BODY_BOLD
                    row += 1
                    for region, amount in stream.geographic_breakdown.items():
                        ws.cell(row=row, column=2, value=region).font = FONT_BODY
                        cell = ws.cell(row=row, column=3, value=float(amount))
                        format_currency_cell(cell)
                        row += 1

        auto_fit_columns(ws)

    def _create_key_metrics(self):
        """Create key financial metrics sheet."""
        ws = self.wb.create_sheet("Key Metrics")
        metrics = self.results["metrics"]

        apply_title(ws, 1, 1, 4, "Key Financial Metrics")

        row = 3
        apply_section_header(ws, row, 1, 4, "Revenue & Growth Metrics")
        row += 1
        headers = ["Metric", "Value", "Benchmark", "Assessment"]
        for i, h in enumerate(headers):
            ws.cell(row=row, column=i + 1, value=h)
        apply_header_row(ws, row, 1, 4)
        row += 1

        metric_rows = [
            ("Revenue Growth (YoY)", metrics.get("revenue_growth_yoy"),
             "Industry avg", self._assess_metric("revenue_growth_yoy", metrics)),
            ("Revenue Growth (Prior Year)", metrics.get("revenue_growth_2yr"),
             "Industry avg", "Reference"),
            ("Gross Margin (Current)", metrics.get("gross_margin_current"),
             "Prior year", self._assess_metric("gross_margin", metrics)),
            ("Gross Margin (Prior)", metrics.get("gross_margin_prior"),
             "—", "Reference"),
            ("AR-to-Revenue Ratio", metrics.get("ar_to_revenue"),
             "< 25%", self._assess_metric("ar_to_revenue", metrics)),
            ("AR Growth (YoY)", metrics.get("ar_growth_yoy"),
             "≈ Revenue growth", self._assess_metric("ar_growth", metrics)),
            ("Days Sales Outstanding", metrics.get("dso"),
             "Industry avg", "Reference"),
            ("Deferred Rev / Revenue", metrics.get("deferred_rev_to_revenue"),
             "Varies", "Reference"),
            ("Deferred Rev Growth", metrics.get("deferred_rev_growth"),
             "< 15%", self._assess_metric("deferred_rev_growth", metrics)),
            ("Effective Tax Rate", metrics.get("effective_tax_rate"),
             "21% statutory", self._assess_metric("etr", metrics)),
            ("RPO / Revenue", metrics.get("rpo_to_revenue"),
             "> 20% notable", "Reference"),
        ]

        for label, value, benchmark, assessment in metric_rows:
            ws.cell(row=row, column=1, value=label).font = FONT_BODY_BOLD
            cell = ws.cell(row=row, column=2)
            if value is not None:
                if "Growth" in label or "Margin" in label or "Ratio" in label or "Rate" in label or "RPO" in label:
                    cell.value = value
                    format_percent_cell(cell)
                elif "DSO" in label or "Days" in label:
                    cell.value = f"{value:.0f} days"
                    cell.font = FONT_BODY
                else:
                    cell.value = value
                    cell.font = FONT_BODY
            else:
                cell.value = "N/A"
                cell.font = FONT_BODY
            ws.cell(row=row, column=3, value=benchmark).font = FONT_BODY
            ws.cell(row=row, column=4, value=assessment).font = FONT_BODY
            ws.cell(row=row, column=1).border = THIN_BORDER
            ws.cell(row=row, column=2).border = THIN_BORDER
            ws.cell(row=row, column=3).border = THIN_BORDER
            ws.cell(row=row, column=4).border = THIN_BORDER
            row += 1

        auto_fit_columns(ws)

    def _create_opportunities(self):
        """Create opportunities detail sheet."""
        ws = self.wb.create_sheet("Opportunities")
        apply_title(ws, 1, 1, 6, "Revenue Recognition Opportunities")

        row = 3
        headers = ["#", "Category", "Finding", "Detail", "Risk Level", "Phase 2 Action"]
        for i, h in enumerate(headers):
            ws.cell(row=row, column=i + 1, value=h)
        apply_header_row(ws, row, 1, len(headers))
        row += 1

        # Sort: High first, then Medium, then Low
        priority_order = {"High": 0, "Medium": 1, "Low": 2}
        sorted_opps = sorted(
            self.results["opportunities"],
            key=lambda x: priority_order.get(x["risk_level"], 3)
        )

        for idx, opp in enumerate(sorted_opps, 1):
            data = [
                idx,
                opp["category"],
                opp["finding"],
                opp["detail"],
                opp["risk_level"],
                opp["phase2_action"],
            ]
            write_data_row(ws, row, data, alternate=(idx % 2 == 0))
            # Style risk level
            risk_cell = ws.cell(row=row, column=5)
            risk_cell.font = get_risk_font(opp["risk_level"])
            risk_cell.fill = get_risk_fill(opp["risk_level"])
            # Wrap detail column
            ws.cell(row=row, column=4).alignment = ALIGN_WRAP
            ws.cell(row=row, column=6).alignment = ALIGN_WRAP
            ws.row_dimensions[row].height = 45
            row += 1

        auto_fit_columns(ws, max_width=60)
        ws.column_dimensions["D"].width = 50
        ws.column_dimensions["F"].width = 40

    def _create_risks(self):
        """Create risks detail sheet."""
        ws = self.wb.create_sheet("Risks")
        apply_title(ws, 1, 1, 6, "Revenue Recognition Risks")

        row = 3
        headers = ["#", "Category", "Finding", "Detail", "Risk Level", "Phase 2 Action"]
        for i, h in enumerate(headers):
            ws.cell(row=row, column=i + 1, value=h)
        apply_header_row(ws, row, 1, len(headers))
        row += 1

        priority_order = {"High": 0, "Medium": 1, "Low": 2}
        sorted_risks = sorted(
            self.results["risks"],
            key=lambda x: priority_order.get(x["risk_level"], 3)
        )

        for idx, risk in enumerate(sorted_risks, 1):
            data = [
                idx,
                risk["category"],
                risk["finding"],
                risk["detail"],
                risk["risk_level"],
                risk["phase2_action"],
            ]
            write_data_row(ws, row, data, alternate=(idx % 2 == 0))
            risk_cell = ws.cell(row=row, column=5)
            risk_cell.font = get_risk_font(risk["risk_level"])
            risk_cell.fill = get_risk_fill(risk["risk_level"])
            ws.cell(row=row, column=4).alignment = ALIGN_WRAP
            ws.cell(row=row, column=6).alignment = ALIGN_WRAP
            ws.row_dimensions[row].height = 45
            row += 1

        auto_fit_columns(ws, max_width=60)
        ws.column_dimensions["D"].width = 50
        ws.column_dimensions["F"].width = 40

    def _create_deferred_revenue(self):
        """Create deferred revenue analysis sheet."""
        ws = self.wb.create_sheet("Deferred Revenue")
        dr = self.results["deferred_revenue"]

        apply_title(ws, 1, 1, 4, "Deferred Revenue Analysis")

        row = 3
        apply_section_header(ws, row, 1, 4, "Balance Summary")
        row += 1

        fields = [
            ("Current Year Balance", float(dr.current_balance)),
            ("Prior Year Balance", float(dr.prior_balance)),
            ("Current Portion", float(dr.current_portion)),
            ("Non-Current Portion", float(dr.noncurrent_portion)),
            ("Recognized from Opening", float(dr.revenue_recognized_from_opening)),
        ]

        for label, value in fields:
            ws.cell(row=row, column=1, value=label).font = FONT_BODY_BOLD
            cell = ws.cell(row=row, column=2, value=value)
            format_currency_cell(cell)
            ws.cell(row=row, column=1).border = THIN_BORDER
            cell.border = THIN_BORDER
            row += 1

        # YoY change
        row += 1
        apply_section_header(ws, row, 1, 4, "Year-over-Year Analysis")
        row += 1
        if dr.prior_balance and dr.prior_balance != 0:
            change = float(dr.current_balance - dr.prior_balance)
            pct = float((dr.current_balance - dr.prior_balance) / abs(dr.prior_balance))
            ws.cell(row=row, column=1, value="Change ($)").font = FONT_BODY_BOLD
            cell = ws.cell(row=row, column=2, value=change)
            format_currency_cell(cell)
            row += 1
            ws.cell(row=row, column=1, value="Change (%)").font = FONT_BODY_BOLD
            cell = ws.cell(row=row, column=2, value=pct)
            format_percent_cell(cell)
            row += 1

        # Tax implications
        row += 1
        apply_section_header(ws, row, 1, 4, "Tax Implications (IRC §451(c))")
        row += 1
        tax_notes = [
            "Under IRC §451(c), taxpayers may elect to defer advance payments to the next tax year.",
            "Book deferred revenue may differ from tax deferred amounts — quantify in Phase 2.",
            f"Current deferred balance of ${float(dr.current_balance):,.0f} represents potential deferral.",
            "Phase 2 will compare book vs. tax treatment using trial balance and tax return.",
        ]
        for note in tax_notes:
            ws.cell(row=row, column=1, value=note).font = FONT_BODY
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
            ws.cell(row=row, column=1).alignment = ALIGN_WRAP
            row += 1

        if dr.notes:
            row += 1
            ws.cell(row=row, column=1, value="Additional Notes:").font = FONT_BODY_BOLD
            row += 1
            ws.cell(row=row, column=1, value=dr.notes).font = FONT_BODY
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)

        auto_fit_columns(ws)

    def _create_phase2_roadmap(self):
        """Create Phase 2 investigation roadmap sheet."""
        ws = self.wb.create_sheet("Phase 2 Roadmap")
        apply_title(ws, 1, 1, 6, "Phase 2 Investigation Roadmap")
        ws.cell(row=2, column=1, value="Documents needed: Trial Balance, Tax Return (1120/1065), Work Papers").font = FONT_BODY

        row = 4
        apply_section_header(ws, row, 1, 6, "Required Documents & Data")
        row += 1
        documents = [
            ("Trial Balance", "General ledger detail with account-level balances",
             "Map revenue accounts; reconcile to 10-K; identify timing differences"),
            ("Tax Return", "Form 1120/1065 with all schedules",
             "Compare book vs. tax revenue; identify M-1/M-3 adjustments"),
            ("Tax Return Work Papers", "Supporting schedules and computations",
             "Revenue recognition adjustments; deferred revenue rollforward; "
             "method of accounting documentation"),
            ("Revenue Detail", "Monthly/quarterly revenue by stream",
             "Cutoff testing; seasonality analysis; trend anomalies"),
            ("AR Aging", "Aged accounts receivable detail",
             "Collection analysis; reserve adequacy; recognition timing"),
        ]

        headers = ["Document", "Description", "Analysis Purpose"]
        for i, h in enumerate(headers):
            ws.cell(row=row, column=i + 1, value=h)
        apply_header_row(ws, row, 1, 3)
        row += 1

        for idx, (doc, desc, purpose) in enumerate(documents):
            write_data_row(ws, row, [doc, desc, purpose], alternate=(idx % 2 == 0))
            ws.cell(row=row, column=3).alignment = ALIGN_WRAP
            ws.row_dimensions[row].height = 35
            row += 1

        # Action items by priority
        row += 1
        apply_section_header(ws, row, 1, 6, "Phase 2 Action Items (by Priority)")
        row += 1

        headers = ["Priority", "Action", "Source Finding", "Expected Outcome"]
        for i, h in enumerate(headers):
            ws.cell(row=row, column=i + 1, value=h)
        apply_header_row(ws, row, 1, 4)
        row += 1

        # Collect all phase 2 actions from opportunities and risks
        all_items = []
        for item in self.results["opportunities"] + self.results["risks"]:
            all_items.append({
                "priority": item["risk_level"],
                "action": item["phase2_action"],
                "source": item["finding"],
                "outcome": "Quantify book-tax difference and identify adjustment opportunity",
            })

        priority_order = {"High": 0, "Medium": 1, "Low": 2}
        all_items.sort(key=lambda x: priority_order.get(x["priority"], 3))

        # Deduplicate actions
        seen_actions = set()
        for item in all_items:
            if item["action"] in seen_actions:
                continue
            seen_actions.add(item["action"])
            data = [item["priority"], item["action"], item["source"], item["outcome"]]
            write_data_row(ws, row, data, alternate=(len(seen_actions) % 2 == 0))
            risk_cell = ws.cell(row=row, column=1)
            risk_cell.font = get_risk_font(item["priority"])
            risk_cell.fill = get_risk_fill(item["priority"])
            ws.cell(row=row, column=2).alignment = ALIGN_WRAP
            ws.cell(row=row, column=3).alignment = ALIGN_WRAP
            ws.row_dimensions[row].height = 35
            row += 1

        auto_fit_columns(ws, max_width=50)
        ws.column_dimensions["B"].width = 45
        ws.column_dimensions["C"].width = 40

    # --- Helper methods ---

    @staticmethod
    def _fmt_pct(value):
        if value is None:
            return "N/A"
        return f"{value:.1%}"

    @staticmethod
    def _assess_metric(metric_key, metrics):
        """Return a simple assessment string for a metric."""
        if metric_key == "revenue_growth_yoy":
            val = metrics.get("revenue_growth_yoy")
            if val is None:
                return "N/A"
            return "Strong" if val > 0.10 else "Moderate" if val > 0 else "Declining"
        elif metric_key == "gross_margin":
            curr = metrics.get("gross_margin_current")
            prior = metrics.get("gross_margin_prior")
            if curr is None or prior is None:
                return "N/A"
            diff = curr - prior
            if abs(diff) < 0.03:
                return "Stable"
            return "Improving" if diff > 0 else "Declining"
        elif metric_key == "ar_to_revenue":
            val = metrics.get("ar_to_revenue")
            if val is None:
                return "N/A"
            return "Flag — High" if val > 0.25 else "Normal"
        elif metric_key == "ar_growth":
            ar_g = metrics.get("ar_growth_yoy")
            rev_g = metrics.get("revenue_growth_yoy")
            if ar_g is None or rev_g is None:
                return "N/A"
            gap = ar_g - rev_g
            return "Flag — AR outpacing revenue" if gap > 0.10 else "Normal"
        elif metric_key == "deferred_rev_growth":
            val = metrics.get("deferred_rev_growth")
            if val is None:
                return "N/A"
            return "Flag — Investigate" if val > 0.15 else "Normal"
        elif metric_key == "etr":
            val = metrics.get("effective_tax_rate")
            if val is None:
                return "N/A"
            if val > 0.25:
                return "Above statutory — opportunity"
            elif val < 0.15:
                return "Below statutory — review"
            return "Near statutory"
        return "N/A"
