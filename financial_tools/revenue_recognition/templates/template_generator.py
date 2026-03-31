"""Generates blank formatted Excel intake templates for revenue recognition phases.

Each template provides structured data entry worksheets with instructions,
formatting, data validation, and color-coded required/optional fields.
"""

import os
from pathlib import Path

from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation

from ..utils.excel_styles import (
    FILL_HIGHLIGHT_BLUE,
    FILL_HIGHLIGHT_GREEN,
    FONT_TITLE,
    FONT_BODY,
    FONT_BODY_BOLD,
    FONT_COLUMN_HEADER,
    THIN_BORDER,
    ALIGN_CENTER,
    ALIGN_LEFT,
    ALIGN_WRAP,
    FMT_CURRENCY,
    FMT_PERCENT,
    FMT_DATE,
    FMT_NUMBER,
    apply_title,
    auto_fit_columns,
)


# Column definition: (header_text, is_required, width, number_format_or_None)
# Data validation is handled separately per sheet.

_RECOGNITION_METHODS = "point-in-time,over-time-time,over-time-input,over-time-output,subscription,license"
_ACCOUNT_TYPES = "revenue,deferred_revenue,contract_asset,ar,cogs,expense,other"
_TAX_TREATMENTS = "same-as-book,deferred,accelerated,method-difference"
_FORM_TYPES = "1120,1065,1120-S"
_ACCOUNTING_METHODS = "cash,accrual,hybrid"
_YES_NO = "Yes,No"
_TEMP_PERM = "Temporary,Permanent"
_DTA_DTL = "DTA,DTL,N/A"
_WP_CATEGORIES = "revenue_timing,deferred_revenue,variable_consideration,method_change,other"
_PO_TYPES = "product,service,license_right_to_use,license_right_to_access,combined"
_SATISFACTION_PATTERNS = "point-in-time,over-time-time,over-time-input,over-time-output"
_VC_TYPES = "discount,rebate,penalty,refund,bonus,price_concession"
_ESTIMATION_METHODS = "expected_value,most_likely_amount"
_MOD_TREATMENTS = "separate_contract,prospective,cumulative_catchup"
_TAX_METHODS = "accrual,completed_contract,pct_completion,cash"


class TemplateGenerator:
    """Generates blank Excel intake templates for each analysis phase."""

    # ------------------------------------------------------------------
    # Internal helpers
    # ------------------------------------------------------------------

    @staticmethod
    def _ensure_dir(output_dir: str) -> Path:
        path = Path(output_dir)
        path.mkdir(parents=True, exist_ok=True)
        return path

    @staticmethod
    def _add_dropdown(ws, col_letter: str, start_row: int, end_row: int, options: str):
        """Attach a list-based data validation dropdown to a column range."""
        dv = DataValidation(
            type="list",
            formula1=f'"{options}"',
            allow_blank=True,
            showDropDown=False,
        )
        dv.error = "Please select a value from the dropdown list."
        dv.errorTitle = "Invalid Entry"
        dv.prompt = "Select from the list."
        dv.promptTitle = "Allowed Values"
        dv.showInputMessage = True
        dv.showErrorMessage = True
        dv.sqref = f"{col_letter}{start_row}:{col_letter}{end_row}"
        ws.add_data_validation(dv)

    @staticmethod
    def _write_headers(ws, row: int, columns: list[tuple]):
        """Write column headers with required (blue) / optional (green) fills.

        *columns* is a list of (header_text, is_required, width, number_format).
        """
        for col_idx, (header, required, width, _) in enumerate(columns, start=1):
            cell = ws.cell(row=row, column=col_idx, value=header)
            cell.font = FONT_COLUMN_HEADER
            cell.fill = FILL_HIGHLIGHT_BLUE if required else FILL_HIGHLIGHT_GREEN
            cell.alignment = ALIGN_CENTER
            cell.border = THIN_BORDER
            ws.column_dimensions[cell.column_letter].width = width

    @staticmethod
    def _apply_body_format(ws, header_row: int, columns: list[tuple], data_rows: int = 100):
        """Pre-format body cells with borders, alignment, and number formats."""
        for r in range(header_row + 1, header_row + 1 + data_rows):
            for col_idx, (_, _, _, fmt) in enumerate(columns, start=1):
                cell = ws.cell(row=r, column=col_idx)
                cell.font = FONT_BODY
                cell.border = THIN_BORDER
                cell.alignment = ALIGN_LEFT
                if fmt:
                    cell.number_format = fmt

    @staticmethod
    def _write_instructions(ws, title: str, lines: list[str]):
        """Populate an Instructions tab with a title and numbered guidance lines."""
        apply_title(ws, 1, 1, 8, title)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=8)

        row = 3
        legend_items = [
            ("Color Legend:", None, True),
            ("  Blue column headers = Required fields", FILL_HIGHLIGHT_BLUE, False),
            ("  Green column headers = Optional fields", FILL_HIGHLIGHT_GREEN, False),
        ]
        for text, fill, bold in legend_items:
            cell = ws.cell(row=row, column=1, value=text)
            cell.font = FONT_BODY_BOLD if bold else FONT_BODY
            if fill:
                cell.fill = fill
            row += 1

        row += 1
        for idx, line in enumerate(lines, start=1):
            cell = ws.cell(row=row, column=1, value=f"{idx}. {line}")
            cell.font = FONT_BODY
            cell.alignment = ALIGN_WRAP
            row += 1

        ws.column_dimensions["A"].width = 100

    # ------------------------------------------------------------------
    # Phase 1
    # ------------------------------------------------------------------

    def generate_phase1_template(self, output_dir: str = "output/templates") -> str:
        """Generate the Phase 1 (10-K public data) intake template."""
        dest = self._ensure_dir(output_dir)
        wb = Workbook()

        # ---- Instructions ----
        ws_inst = wb.active
        ws_inst.title = "Instructions"
        self._write_instructions(ws_inst, "Phase 1 — 10-K Public Data Intake Template", [
            "Fill in company profile information on the 'Company Profile' tab.",
            "Enter each distinct revenue stream on the 'Revenue Streams' tab (one per row).",
            "Enter deferred revenue balances on the 'Deferred Revenue' tab.",
            "Enter contract asset data on the 'Contract Assets' tab.",
            "Complete all disclosure fields on the 'Disclosures' tab.",
            "Enter income statement and balance sheet figures on their respective tabs.",
            "List any restatement history or revenue risk factors on the 'Risk Factors' tab.",
            "Blue-highlighted headers are required; green are optional but recommended.",
            "Use dropdown menus where provided — do not enter free-form text in those cells.",
            "Currency amounts should be entered as numbers (no $ signs or commas).",
        ])

        # ---- Company Profile ----
        ws_cp = wb.create_sheet("Company Profile")
        cp_cols = [
            ("Company Name", True, 30, None),
            ("Ticker", False, 12, None),
            ("CIK", False, 14, None),
            ("SIC Code", False, 12, None),
            ("Industry", True, 25, None),
            ("FY End Month (1-12)", True, 18, FMT_NUMBER),
            ("Filing Date", False, 16, FMT_DATE),
            ("Currency", True, 12, None),
            ("Auditor", False, 25, None),
            ("Opinion Type", False, 18, None),
        ]
        self._write_headers(ws_cp, 1, cp_cols)
        self._apply_body_format(ws_cp, 1, cp_cols, data_rows=1)
        self._add_dropdown(ws_cp, "J", 2, 2, "unqualified,qualified,adverse,disclaimer")

        # ---- Revenue Streams ----
        ws_rs = wb.create_sheet("Revenue Streams")
        rs_cols = [
            ("Name", True, 25, None),
            ("Description", False, 35, None),
            ("Recognition Method", True, 22, None),
            ("Amount Current", True, 18, FMT_CURRENCY),
            ("Amount Prior", True, 18, FMT_CURRENCY),
            ("Amount 2Yr Prior", False, 18, FMT_CURRENCY),
            ("Customer Concentration %", False, 22, FMT_PERCENT),
            ("Notes", False, 30, None),
            ("Geo - US/Domestic", False, 18, FMT_CURRENCY),
            ("Geo - EMEA", False, 18, FMT_CURRENCY),
            ("Geo - APAC", False, 18, FMT_CURRENCY),
            ("Geo - LATAM", False, 18, FMT_CURRENCY),
            ("Geo - Other", False, 18, FMT_CURRENCY),
        ]
        self._write_headers(ws_rs, 1, rs_cols)
        self._apply_body_format(ws_rs, 1, rs_cols, data_rows=50)
        self._add_dropdown(ws_rs, "C", 2, 51, _RECOGNITION_METHODS)

        # ---- Deferred Revenue ----
        ws_dr = wb.create_sheet("Deferred Revenue")
        dr_cols = [
            ("Current Balance", True, 18, FMT_CURRENCY),
            ("Prior Balance", True, 18, FMT_CURRENCY),
            ("Current Portion", True, 18, FMT_CURRENCY),
            ("Noncurrent Portion", True, 18, FMT_CURRENCY),
            ("Revenue Recognized from Opening", True, 25, FMT_CURRENCY),
            ("Notes", False, 35, None),
        ]
        self._write_headers(ws_dr, 1, dr_cols)
        self._apply_body_format(ws_dr, 1, dr_cols, data_rows=1)

        # ---- Contract Assets ----
        ws_ca = wb.create_sheet("Contract Assets")
        ca_cols = [
            ("Current Balance", True, 18, FMT_CURRENCY),
            ("Prior Balance", True, 18, FMT_CURRENCY),
            ("Impairment Losses", False, 18, FMT_CURRENCY),
            ("Notes", False, 35, None),
        ]
        self._write_headers(ws_ca, 1, ca_cols)
        self._apply_body_format(ws_ca, 1, ca_cols, data_rows=1)

        # ---- Disclosures ----
        ws_disc = wb.create_sheet("Disclosures")
        disc_cols = [
            ("Field", True, 35, None),
            ("Value", True, 60, None),
        ]
        self._write_headers(ws_disc, 1, disc_cols)
        disclosure_rows = [
            "ASC 606 Policy Summary",
            "Significant Judgment 1",
            "Significant Judgment 2",
            "Significant Judgment 3",
            "Variable Consideration Type 1",
            "Variable Consideration Type 2",
            "Variable Consideration Type 3",
            "Contract Cost Capitalized",
            "Remaining Performance Obligations",
            "RPO Expected Timing",
            "Disaggregation Dimension 1",
            "Disaggregation Dimension 2",
            "Disaggregation Dimension 3",
            "Significant Change 1",
            "Significant Change 2",
            "Related Party Revenue",
        ]
        for i, label in enumerate(disclosure_rows, start=2):
            cell = ws_disc.cell(row=i, column=1, value=label)
            cell.font = FONT_BODY_BOLD
            cell.border = THIN_BORDER
            cell.alignment = ALIGN_LEFT
            val_cell = ws_disc.cell(row=i, column=2)
            val_cell.font = FONT_BODY
            val_cell.border = THIN_BORDER
            val_cell.alignment = ALIGN_LEFT

        # ---- Income Statement ----
        ws_is = wb.create_sheet("Income Statement")
        is_cols = [
            ("Line Item", True, 30, None),
            ("Current Year", True, 20, FMT_CURRENCY),
            ("Prior Year", True, 20, FMT_CURRENCY),
        ]
        self._write_headers(ws_is, 1, is_cols)
        is_rows = [
            "Total Revenue",
            "Cost of Revenue (COGS)",
            "Operating Income",
            "Net Income",
            "Income Tax Expense",
            "Pretax Income",
        ]
        for i, label in enumerate(is_rows, start=2):
            cell = ws_is.cell(row=i, column=1, value=label)
            cell.font = FONT_BODY_BOLD
            cell.border = THIN_BORDER
            for c in (2, 3):
                bc = ws_is.cell(row=i, column=c)
                bc.font = FONT_BODY
                bc.border = THIN_BORDER
                bc.number_format = FMT_CURRENCY

        # ---- Balance Sheet ----
        ws_bs = wb.create_sheet("Balance Sheet")
        bs_cols = [
            ("Line Item", True, 30, None),
            ("Current Year", True, 20, FMT_CURRENCY),
            ("Prior Year", True, 20, FMT_CURRENCY),
        ]
        self._write_headers(ws_bs, 1, bs_cols)
        bs_rows = [
            "Accounts Receivable",
            "Allowance for Doubtful Accounts",
            "Total Assets",
        ]
        for i, label in enumerate(bs_rows, start=2):
            cell = ws_bs.cell(row=i, column=1, value=label)
            cell.font = FONT_BODY_BOLD
            cell.border = THIN_BORDER
            for c in (2, 3):
                bc = ws_bs.cell(row=i, column=c)
                bc.font = FONT_BODY
                bc.border = THIN_BORDER
                bc.number_format = FMT_CURRENCY

        # ---- Risk Factors ----
        ws_rf = wb.create_sheet("Risk Factors")
        rf_cols = [
            ("Type", True, 20, None),
            ("Description", True, 60, None),
        ]
        self._write_headers(ws_rf, 1, rf_cols)
        self._apply_body_format(ws_rf, 1, rf_cols, data_rows=20)
        self._add_dropdown(ws_rf, "A", 2, 21, "Restatement,Revenue Risk Factor")

        # Save
        filepath = str(dest / "phase1_template.xlsx")
        wb.save(filepath)
        return filepath

    # ------------------------------------------------------------------
    # Phase 2
    # ------------------------------------------------------------------

    def generate_phase2_template(self, output_dir: str = "output/templates") -> str:
        """Generate the Phase 2 (trial balance / tax return) intake template."""
        dest = self._ensure_dir(output_dir)
        wb = Workbook()

        # ---- Instructions ----
        ws_inst = wb.active
        ws_inst.title = "Instructions"
        self._write_instructions(ws_inst, "Phase 2 — Trial Balance & Tax Data Intake Template", [
            "Export your trial balance and enter one row per GL account on the 'Trial Balance' tab.",
            "Map revenue-related accounts to streams on the 'Account Mappings' tab.",
            "Complete the deferred revenue rollforward on its tab (one row per category).",
            "Enter tax return data on the 'Tax Return' tab (single-row entry).",
            "List M-1/M-3 revenue adjustments on the 'M-1 Adjustments' tab.",
            "Document tax work paper items on the 'Work Papers' tab.",
            "Use dropdown menus where provided.",
            "Currency amounts should be entered as numbers without formatting.",
        ])

        # ---- Trial Balance ----
        ws_tb = wb.create_sheet("Trial Balance")
        tb_cols = [
            ("Account Number", True, 18, None),
            ("Account Name", True, 30, None),
            ("Account Type", True, 18, None),
            ("Beginning Balance", True, 18, FMT_CURRENCY),
            ("Ending Balance", True, 18, FMT_CURRENCY),
            ("Debits", False, 18, FMT_CURRENCY),
            ("Credits", False, 18, FMT_CURRENCY),
            ("Department", False, 18, None),
            ("Entity", False, 18, None),
            ("Notes", False, 30, None),
        ]
        self._write_headers(ws_tb, 1, tb_cols)
        self._apply_body_format(ws_tb, 1, tb_cols, data_rows=500)
        self._add_dropdown(ws_tb, "C", 2, 501, _ACCOUNT_TYPES)

        # ---- Account Mappings ----
        ws_am = wb.create_sheet("Account Mappings")
        am_cols = [
            ("Account Number", True, 18, None),
            ("Revenue Stream", True, 25, None),
            ("Recognition Type", True, 20, None),
            ("Tax Treatment", True, 20, None),
            ("Book-Tax Difference", False, 20, FMT_CURRENCY),
            ("Notes", False, 30, None),
        ]
        self._write_headers(ws_am, 1, am_cols)
        self._apply_body_format(ws_am, 1, am_cols, data_rows=200)
        self._add_dropdown(ws_am, "C", 2, 201, _RECOGNITION_METHODS)
        self._add_dropdown(ws_am, "D", 2, 201, _TAX_TREATMENTS)

        # ---- Deferred Revenue Rollforward ----
        ws_drf = wb.create_sheet("Deferred Revenue Rollforward")
        drf_cols = [
            ("Category", True, 25, None),
            ("Opening Balance", True, 18, FMT_CURRENCY),
            ("Additions", True, 18, FMT_CURRENCY),
            ("Recognized", True, 18, FMT_CURRENCY),
            ("Adjustments", False, 18, FMT_CURRENCY),
            ("Closing Balance", True, 18, FMT_CURRENCY),
            ("Notes", False, 30, None),
        ]
        self._write_headers(ws_drf, 1, drf_cols)
        self._apply_body_format(ws_drf, 1, drf_cols, data_rows=20)

        # ---- Tax Return ----
        ws_tr = wb.create_sheet("Tax Return")
        tr_cols = [
            ("Form Type", True, 14, None),
            ("Tax Year", True, 12, FMT_NUMBER),
            ("Gross Receipts", True, 18, FMT_CURRENCY),
            ("Returns & Allowances", False, 20, FMT_CURRENCY),
            ("Net Receipts", True, 18, FMT_CURRENCY),
            ("Book Income", True, 18, FMT_CURRENCY),
            ("Tax Income", True, 18, FMT_CURRENCY),
            ("Accounting Method", True, 18, None),
            ("Section 451(c) Election", False, 22, None),
            ("Section 451(b) AFS", False, 20, None),
            ("Tax Deferred Rev Current", False, 22, FMT_CURRENCY),
            ("Tax Deferred Rev Prior", False, 20, FMT_CURRENCY),
        ]
        self._write_headers(ws_tr, 1, tr_cols)
        self._apply_body_format(ws_tr, 1, tr_cols, data_rows=1)
        self._add_dropdown(ws_tr, "A", 2, 2, _FORM_TYPES)
        self._add_dropdown(ws_tr, "H", 2, 2, _ACCOUNTING_METHODS)
        self._add_dropdown(ws_tr, "I", 2, 2, _YES_NO)
        self._add_dropdown(ws_tr, "J", 2, 2, _YES_NO)

        # ---- M-1 Adjustments ----
        ws_m1 = wb.create_sheet("M-1 Adjustments")
        m1_cols = [
            ("Description", True, 35, None),
            ("Book Amount", True, 18, FMT_CURRENCY),
            ("Tax Amount", True, 18, FMT_CURRENCY),
            ("Type", True, 16, None),
            ("Explanation", False, 40, None),
        ]
        self._write_headers(ws_m1, 1, m1_cols)
        self._apply_body_format(ws_m1, 1, m1_cols, data_rows=50)
        self._add_dropdown(ws_m1, "D", 2, 51, _TEMP_PERM)

        # ---- Work Papers ----
        ws_wp = wb.create_sheet("Work Papers")
        wp_cols = [
            ("Description", True, 35, None),
            ("Category", True, 22, None),
            ("Book Amount", True, 18, FMT_CURRENCY),
            ("Tax Amount", True, 18, FMT_CURRENCY),
            ("Difference", False, 18, FMT_CURRENCY),
            ("Permanent or Temporary", True, 22, None),
            ("DTA or DTL", False, 12, None),
            ("Supporting Reference", False, 25, None),
            ("Notes", False, 30, None),
        ]
        self._write_headers(ws_wp, 1, wp_cols)
        self._apply_body_format(ws_wp, 1, wp_cols, data_rows=100)
        # Pre-populate difference formula: =C{row}-D{row}
        for r in range(2, 102):
            ws_wp.cell(row=r, column=5).value = f"=C{r}-D{r}"
        self._add_dropdown(ws_wp, "B", 2, 101, _WP_CATEGORIES)
        self._add_dropdown(ws_wp, "F", 2, 101, _TEMP_PERM)
        self._add_dropdown(ws_wp, "G", 2, 101, _DTA_DTL)

        # Save
        filepath = str(dest / "phase2_template.xlsx")
        wb.save(filepath)
        return filepath

    # ------------------------------------------------------------------
    # Phase 3
    # ------------------------------------------------------------------

    def generate_phase3_template(self, output_dir: str = "output/templates") -> str:
        """Generate the Phase 3 (contract-level) intake template."""
        dest = self._ensure_dir(output_dir)
        wb = Workbook()

        # ---- Instructions ----
        ws_inst = wb.active
        ws_inst.title = "Instructions"
        self._write_instructions(ws_inst, "Phase 3 — Contract-Level Analysis Intake Template", [
            "Enter each customer contract on the 'Contracts' tab (one per row).",
            "Enter performance obligations on the 'Performance Obligations' tab, referencing Contract ID.",
            "Enter variable consideration elements on the 'Variable Consideration' tab.",
            "Enter contract modifications on the 'Contract Modifications' tab.",
            "Set analysis parameters on the 'Analysis Settings' tab.",
            "Use dropdown menus where provided.",
            "Dates should be entered in MM/DD/YYYY or YYYY-MM-DD format.",
            "Currency amounts should be entered as numbers without formatting.",
        ])

        # ---- Contracts ----
        ws_c = wb.create_sheet("Contracts")
        c_cols = [
            ("Contract ID", True, 16, None),
            ("Customer Name", True, 25, None),
            ("Description", False, 35, None),
            ("Contract Date", True, 16, FMT_DATE),
            ("Start Date", True, 16, FMT_DATE),
            ("End Date", True, 16, FMT_DATE),
            ("Total Transaction Price", True, 22, FMT_CURRENCY),
            ("Currency", False, 12, None),
            ("Tax Method", True, 20, None),
            ("Section 451(c) Applicable", False, 22, None),
            ("Advance Payment Amount", False, 22, FMT_CURRENCY),
            ("Long-Term Contract", False, 18, None),
            ("Notes", False, 30, None),
        ]
        self._write_headers(ws_c, 1, c_cols)
        self._apply_body_format(ws_c, 1, c_cols, data_rows=100)
        self._add_dropdown(ws_c, "I", 2, 101, _TAX_METHODS)
        self._add_dropdown(ws_c, "J", 2, 101, _YES_NO)
        self._add_dropdown(ws_c, "L", 2, 101, _YES_NO)

        # ---- Performance Obligations ----
        ws_po = wb.create_sheet("Performance Obligations")
        po_cols = [
            ("Contract ID", True, 16, None),
            ("Obligation ID", True, 16, None),
            ("Description", True, 30, None),
            ("Type", True, 22, None),
            ("Satisfaction Pattern", True, 22, None),
            ("Standalone Selling Price", True, 22, FMT_CURRENCY),
            ("Allocated Transaction Price", True, 24, FMT_CURRENCY),
            ("Satisfaction Date", False, 16, FMT_DATE),
            ("Service Start", False, 16, FMT_DATE),
            ("Service End", False, 16, FMT_DATE),
            ("Percent Complete", False, 16, FMT_PERCENT),
            ("Costs Incurred", False, 18, FMT_CURRENCY),
            ("Total Estimated Costs", False, 20, FMT_CURRENCY),
            ("Units Delivered", False, 16, FMT_NUMBER),
            ("Total Units", False, 14, FMT_NUMBER),
            ("Book Revenue Recognized", False, 22, FMT_CURRENCY),
            ("Book Revenue Deferred", False, 22, FMT_CURRENCY),
            ("Tax Revenue Recognized", False, 22, FMT_CURRENCY),
            ("Tax Treatment Notes", False, 25, None),
            ("Notes", False, 25, None),
        ]
        self._write_headers(ws_po, 1, po_cols)
        self._apply_body_format(ws_po, 1, po_cols, data_rows=200)
        self._add_dropdown(ws_po, "D", 2, 201, _PO_TYPES)
        self._add_dropdown(ws_po, "E", 2, 201, _SATISFACTION_PATTERNS)

        # ---- Variable Consideration ----
        ws_vc = wb.create_sheet("Variable Consideration")
        vc_cols = [
            ("Contract ID", True, 16, None),
            ("Type", True, 18, None),
            ("Description", True, 30, None),
            ("Estimated Amount", True, 18, FMT_CURRENCY),
            ("Constrained Amount", True, 18, FMT_CURRENCY),
            ("Estimation Method", True, 22, None),
            ("Constraint Rationale", False, 30, None),
            ("Book Treatment", False, 25, None),
            ("Tax Treatment", False, 25, None),
            ("Book-Tax Difference", False, 20, FMT_CURRENCY),
        ]
        self._write_headers(ws_vc, 1, vc_cols)
        self._apply_body_format(ws_vc, 1, vc_cols, data_rows=100)
        self._add_dropdown(ws_vc, "B", 2, 101, _VC_TYPES)
        self._add_dropdown(ws_vc, "F", 2, 101, _ESTIMATION_METHODS)

        # ---- Contract Modifications ----
        ws_cm = wb.create_sheet("Contract Modifications")
        cm_cols = [
            ("Contract ID", True, 16, None),
            ("Modification Date", True, 18, FMT_DATE),
            ("Description", True, 35, None),
            ("Accounting Treatment", True, 22, None),
            ("Price Change", False, 18, FMT_CURRENCY),
            ("Scope Change", False, 25, None),
            ("Impact on Recognition", False, 30, None),
            ("Book-Tax Impact", False, 18, FMT_CURRENCY),
            ("Notes", False, 30, None),
        ]
        self._write_headers(ws_cm, 1, cm_cols)
        self._apply_body_format(ws_cm, 1, cm_cols, data_rows=50)
        self._add_dropdown(ws_cm, "D", 2, 51, _MOD_TREATMENTS)

        # ---- Analysis Settings ----
        ws_as = wb.create_sheet("Analysis Settings")
        as_cols = [
            ("Setting", True, 30, None),
            ("Value", True, 25, None),
        ]
        self._write_headers(ws_as, 1, as_cols)
        settings_rows = [
            ("Reporting Period End", FMT_DATE),
            ("Tax Year", FMT_NUMBER),
            ("Statutory Rate", FMT_PERCENT),
        ]
        for i, (label, fmt) in enumerate(settings_rows, start=2):
            cell = ws_as.cell(row=i, column=1, value=label)
            cell.font = FONT_BODY_BOLD
            cell.border = THIN_BORDER
            cell.alignment = ALIGN_LEFT
            val_cell = ws_as.cell(row=i, column=2)
            val_cell.font = FONT_BODY
            val_cell.border = THIN_BORDER
            val_cell.alignment = ALIGN_LEFT
            if fmt:
                val_cell.number_format = fmt

        # Save
        filepath = str(dest / "phase3_template.xlsx")
        wb.save(filepath)
        return filepath

    # ------------------------------------------------------------------
    # All phases
    # ------------------------------------------------------------------

    def generate_all_templates(self, output_dir: str = "output/templates") -> list[str]:
        """Generate templates for all three phases and return their file paths."""
        return [
            self.generate_phase1_template(output_dir),
            self.generate_phase2_template(output_dir),
            self.generate_phase3_template(output_dir),
        ]
