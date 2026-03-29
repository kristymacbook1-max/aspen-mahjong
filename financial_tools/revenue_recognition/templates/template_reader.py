"""Reads filled-in Excel intake templates and converts them to dataclass objects.

Handles empty cells, currency parsing, date parsing, boolean coercion,
and decimal conversion to produce clean Phase 1/2/3 input objects.
"""

from decimal import Decimal
from typing import Optional

from openpyxl import load_workbook

from ..utils.date_helpers import parse_date
from ..utils.currency_helpers import to_decimal

from ..phase1.ten_k_analyzer import (
    TenKInput,
    CompanyProfile,
    RevenueStream,
    DeferredRevenueData,
    ContractAssetData,
    RevenueDisclosures,
    IncomeStatementData,
    BalanceSheetData,
)
from ..phase2.trial_balance_analyzer import (
    Phase2Input,
    TrialBalanceAccount,
    RevenueAccountMapping,
    DeferredRevenueRollforward,
    TaxReturnData,
    WorkPaperItem,
)
from ..phase3.contract_analyzer import (
    Phase3Input,
    Contract,
    PerformanceObligation,
    VariableConsideration,
    ContractModification,
)


def _parse_bool(value) -> bool:
    """Convert various truthy representations to a Python bool."""
    if value is None:
        return False
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return bool(value)
    text = str(value).strip().lower()
    return text in ("yes", "true", "1", "y")


def _str(value) -> str:
    """Safely convert a cell value to a stripped string."""
    if value is None:
        return ""
    return str(value).strip()


def _opt_float(value) -> Optional[float]:
    """Convert to float or return None if blank."""
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace("%", "").replace(",", "")
    if not text or text == "-":
        return None
    try:
        return float(text)
    except ValueError:
        return None


def _opt_int(value) -> Optional[int]:
    """Convert to int or return None if blank."""
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return int(value)
    text = str(value).strip().replace(",", "")
    if not text or text == "-":
        return None
    try:
        return int(float(text))
    except ValueError:
        return None


def _row_values(ws, row: int, max_col: int) -> list:
    """Return a list of cell values for a given row."""
    return [ws.cell(row=row, column=c).value for c in range(1, max_col + 1)]


def _is_empty_row(values, key_indices: list[int]) -> bool:
    """Check whether all key-column values are blank."""
    for idx in key_indices:
        v = values[idx] if idx < len(values) else None
        if v is not None and _str(v) != "":
            return False
    return True


class TemplateReader:
    """Reads filled-in Excel templates and produces phase input dataclass objects."""

    # ------------------------------------------------------------------
    # Phase 1
    # ------------------------------------------------------------------

    def read_phase1(self, filepath: str) -> TenKInput:
        """Read a Phase 1 template and return a populated TenKInput."""
        wb = load_workbook(filepath, data_only=True)

        company = self._read_company_profile(wb["Company Profile"])
        revenue_streams = self._read_revenue_streams(wb["Revenue Streams"])
        deferred_revenue = self._read_deferred_revenue(wb["Deferred Revenue"])
        contract_assets = self._read_contract_assets(wb["Contract Assets"])
        disclosures = self._read_disclosures(wb["Disclosures"])
        income_statement = self._read_income_statement(wb["Income Statement"])
        balance_sheet = self._read_balance_sheet(wb["Balance Sheet"])
        auditor_name, audit_opinion, restatements, risk_factors = self._read_risk_factors(
            wb["Risk Factors"], company
        )

        # Auditor info may also live on company profile row
        if not auditor_name:
            cp_row = _row_values(wb["Company Profile"], 2, 10)
            auditor_name = _str(cp_row[8]) if len(cp_row) > 8 else ""
            audit_opinion = _str(cp_row[9]) if len(cp_row) > 9 else ""

        wb.close()

        return TenKInput(
            company=company,
            revenue_streams=revenue_streams,
            deferred_revenue=deferred_revenue,
            contract_assets=contract_assets,
            disclosures=disclosures,
            income_statement=income_statement,
            balance_sheet=balance_sheet,
            auditor_name=auditor_name,
            audit_opinion_type=audit_opinion,
            restatements=restatements,
            risk_factors_revenue_related=risk_factors,
        )

    def _read_company_profile(self, ws) -> CompanyProfile:
        vals = _row_values(ws, 2, 10)
        return CompanyProfile(
            name=_str(vals[0]),
            ticker=_str(vals[1]) or None,
            cik=_str(vals[2]) or None,
            sic_code=_str(vals[3]) or None,
            industry=_str(vals[4]),
            fiscal_year_end_month=_opt_int(vals[5]) or 12,
            filing_date=_str(vals[6]) or None,
            reporting_currency=_str(vals[7]) or "USD",
        )

    def _read_revenue_streams(self, ws) -> list:
        streams = []
        for row in range(2, ws.max_row + 1):
            vals = _row_values(ws, row, 13)
            if _is_empty_row(vals, [0]):
                continue
            geo = {}
            geo_labels = ["US/Domestic", "EMEA", "APAC", "LATAM", "Other"]
            for i, label in enumerate(geo_labels):
                amt = to_decimal(vals[8 + i])
                if amt:
                    geo[label] = amt
            streams.append(RevenueStream(
                name=_str(vals[0]),
                description=_str(vals[1]),
                recognition_method=_str(vals[2]),
                amount_current=to_decimal(vals[3]),
                amount_prior=to_decimal(vals[4]),
                amount_two_years_prior=to_decimal(vals[5]),
                customer_concentration_pct=_opt_float(vals[6]),
                notes=_str(vals[7]),
                geographic_breakdown=geo,
            ))
        return streams

    def _read_deferred_revenue(self, ws) -> DeferredRevenueData:
        vals = _row_values(ws, 2, 6)
        return DeferredRevenueData(
            current_balance=to_decimal(vals[0]),
            prior_balance=to_decimal(vals[1]),
            current_portion=to_decimal(vals[2]),
            noncurrent_portion=to_decimal(vals[3]),
            revenue_recognized_from_opening=to_decimal(vals[4]),
            notes=_str(vals[5]),
        )

    def _read_contract_assets(self, ws) -> ContractAssetData:
        vals = _row_values(ws, 2, 4)
        return ContractAssetData(
            current_balance=to_decimal(vals[0]),
            prior_balance=to_decimal(vals[1]),
            impairment_losses=to_decimal(vals[2]),
            notes=_str(vals[3]),
        )

    def _read_disclosures(self, ws) -> RevenueDisclosures:
        # Read key-value pairs from column A (field) / column B (value)
        field_map = {}
        for row in range(2, ws.max_row + 1):
            key = _str(ws.cell(row=row, column=1).value)
            val = ws.cell(row=row, column=2).value
            if key:
                field_map[key] = val

        def _collect_list(prefix: str) -> list[str]:
            items = []
            for k, v in field_map.items():
                if k.startswith(prefix) and v is not None and _str(v):
                    items.append(_str(v))
            return items

        return RevenueDisclosures(
            asc606_policy_summary=_str(field_map.get("ASC 606 Policy Summary", "")),
            significant_judgments=_collect_list("Significant Judgment"),
            variable_consideration_types=_collect_list("Variable Consideration Type"),
            contract_cost_capitalized=to_decimal(field_map.get("Contract Cost Capitalized")),
            remaining_performance_obligations=to_decimal(
                field_map.get("Remaining Performance Obligations")
            ),
            rpo_expected_timing=_str(field_map.get("RPO Expected Timing", "")),
            disaggregation_dimensions=_collect_list("Disaggregation Dimension"),
            significant_changes_noted=_collect_list("Significant Change"),
            related_party_revenue=to_decimal(field_map.get("Related Party Revenue")),
        )

    def _read_income_statement(self, ws) -> IncomeStatementData:
        label_map = {}
        for row in range(2, ws.max_row + 1):
            label = _str(ws.cell(row=row, column=1).value)
            current = ws.cell(row=row, column=2).value
            prior = ws.cell(row=row, column=3).value
            if label:
                label_map[label] = (current, prior)

        def _get(key):
            pair = label_map.get(key, (None, None))
            return to_decimal(pair[0]), to_decimal(pair[1])

        rev_c, rev_p = _get("Total Revenue")
        cogs_c, cogs_p = _get("Cost of Revenue (COGS)")
        oi_c, oi_p = _get("Operating Income")
        ni_c, ni_p = _get("Net Income")
        tax_c, tax_p = _get("Income Tax Expense")
        pti_c, pti_p = _get("Pretax Income")

        return IncomeStatementData(
            total_revenue_current=rev_c,
            total_revenue_prior=rev_p,
            cost_of_revenue_current=cogs_c,
            cost_of_revenue_prior=cogs_p,
            operating_income_current=oi_c,
            operating_income_prior=oi_p,
            net_income_current=ni_c,
            net_income_prior=ni_p,
            income_tax_expense_current=tax_c,
            income_tax_expense_prior=tax_p,
            pretax_income_current=pti_c,
            pretax_income_prior=pti_p,
        )

    def _read_balance_sheet(self, ws) -> BalanceSheetData:
        label_map = {}
        for row in range(2, ws.max_row + 1):
            label = _str(ws.cell(row=row, column=1).value)
            current = ws.cell(row=row, column=2).value
            prior = ws.cell(row=row, column=3).value
            if label:
                label_map[label] = (current, prior)

        def _get(key):
            pair = label_map.get(key, (None, None))
            return to_decimal(pair[0]), to_decimal(pair[1])

        ar_c, ar_p = _get("Accounts Receivable")
        allo_c, allo_p = _get("Allowance for Doubtful Accounts")
        ta_c, ta_p = _get("Total Assets")

        return BalanceSheetData(
            accounts_receivable_current=ar_c,
            accounts_receivable_prior=ar_p,
            allowance_for_doubtful_current=allo_c,
            allowance_for_doubtful_prior=allo_p,
            total_assets_current=ta_c,
            total_assets_prior=ta_p,
        )

    def _read_risk_factors(self, ws, company: CompanyProfile):
        auditor_name = ""
        audit_opinion = ""
        restatements = []
        risk_factors = []
        for row in range(2, ws.max_row + 1):
            vals = _row_values(ws, row, 2)
            if _is_empty_row(vals, [0, 1]):
                continue
            rtype = _str(vals[0]).lower()
            desc = _str(vals[1])
            if "restatement" in rtype:
                restatements.append(desc)
            else:
                risk_factors.append(desc)
        return auditor_name, audit_opinion, restatements, risk_factors

    # ------------------------------------------------------------------
    # Phase 2
    # ------------------------------------------------------------------

    def read_phase2(self, filepath: str) -> Phase2Input:
        """Read a Phase 2 template and return a populated Phase2Input."""
        wb = load_workbook(filepath, data_only=True)

        trial_balance = self._read_trial_balance(wb["Trial Balance"])
        account_mappings = self._read_account_mappings(wb["Account Mappings"])
        deferred_rollforwards = self._read_deferred_rollforwards(
            wb["Deferred Revenue Rollforward"]
        )
        tax_return = self._read_tax_return(wb["Tax Return"])
        m1_adjustments = self._read_m1_adjustments(wb["M-1 Adjustments"])
        work_papers = self._read_work_papers(wb["Work Papers"])

        wb.close()

        # Attach M-1 adjustments to tax return
        tax_return.m1_revenue_adjustments = m1_adjustments

        return Phase2Input(
            trial_balance=trial_balance,
            account_mappings=account_mappings,
            deferred_rollforwards=deferred_rollforwards,
            tax_return=tax_return,
            work_papers=work_papers,
        )

    def _read_trial_balance(self, ws) -> list:
        accounts = []
        for row in range(2, ws.max_row + 1):
            vals = _row_values(ws, row, 10)
            if _is_empty_row(vals, [0, 1]):
                continue
            accounts.append(TrialBalanceAccount(
                account_number=_str(vals[0]),
                account_name=_str(vals[1]),
                account_type=_str(vals[2]),
                beginning_balance=to_decimal(vals[3]),
                ending_balance=to_decimal(vals[4]),
                debits=to_decimal(vals[5]),
                credits=to_decimal(vals[6]),
                department=_str(vals[7]),
                entity=_str(vals[8]),
                notes=_str(vals[9]),
            ))
        return accounts

    def _read_account_mappings(self, ws) -> list:
        mappings = []
        for row in range(2, ws.max_row + 1):
            vals = _row_values(ws, row, 6)
            if _is_empty_row(vals, [0]):
                continue
            mappings.append(RevenueAccountMapping(
                account_number=_str(vals[0]),
                revenue_stream=_str(vals[1]),
                recognition_type=_str(vals[2]),
                tax_treatment=_str(vals[3]),
                book_tax_difference=to_decimal(vals[4]),
                notes=_str(vals[5]),
            ))
        return mappings

    def _read_deferred_rollforwards(self, ws) -> list:
        rollforwards = []
        for row in range(2, ws.max_row + 1):
            vals = _row_values(ws, row, 7)
            if _is_empty_row(vals, [0]):
                continue
            rollforwards.append(DeferredRevenueRollforward(
                category=_str(vals[0]),
                opening_balance=to_decimal(vals[1]),
                additions=to_decimal(vals[2]),
                recognized=to_decimal(vals[3]),
                adjustments=to_decimal(vals[4]),
                closing_balance=to_decimal(vals[5]),
                notes=_str(vals[6]),
            ))
        return rollforwards

    def _read_tax_return(self, ws) -> TaxReturnData:
        vals = _row_values(ws, 2, 12)
        return TaxReturnData(
            form_type=_str(vals[0]) or "1120",
            tax_year=_opt_int(vals[1]) or 0,
            gross_receipts_line=to_decimal(vals[2]),
            returns_and_allowances=to_decimal(vals[3]),
            net_receipts=to_decimal(vals[4]),
            book_income=to_decimal(vals[5]),
            tax_income=to_decimal(vals[6]),
            accounting_method=_str(vals[7]),
            section_451c_election=_parse_bool(vals[8]),
            section_451b_afs=_parse_bool(vals[9]),
            tax_deferred_revenue_current=to_decimal(vals[10]),
            tax_deferred_revenue_prior=to_decimal(vals[11]),
        )

    def _read_m1_adjustments(self, ws) -> list:
        adjustments = []
        for row in range(2, ws.max_row + 1):
            vals = _row_values(ws, row, 5)
            if _is_empty_row(vals, [0]):
                continue
            adjustments.append({
                "description": _str(vals[0]),
                "book_amount": to_decimal(vals[1]),
                "tax_amount": to_decimal(vals[2]),
                "type": _str(vals[3]),
                "explanation": _str(vals[4]),
            })
        return adjustments

    def _read_work_papers(self, ws) -> list:
        items = []
        for row in range(2, ws.max_row + 1):
            vals = _row_values(ws, row, 9)
            if _is_empty_row(vals, [0]):
                continue
            items.append(WorkPaperItem(
                description=_str(vals[0]),
                category=_str(vals[1]),
                book_amount=to_decimal(vals[2]),
                tax_amount=to_decimal(vals[3]),
                difference=to_decimal(vals[4]),
                permanent_or_temporary=_str(vals[5]),
                dta_or_dtl=_str(vals[6]),
                supporting_reference=_str(vals[7]),
                notes=_str(vals[8]),
            ))
        return items

    # ------------------------------------------------------------------
    # Phase 3
    # ------------------------------------------------------------------

    def read_phase3(self, filepath: str) -> Phase3Input:
        """Read a Phase 3 template and return a populated Phase3Input."""
        wb = load_workbook(filepath, data_only=True)

        contracts = self._read_contracts(wb["Contracts"])
        performance_obligations = self._read_performance_obligations(
            wb["Performance Obligations"]
        )
        variable_considerations = self._read_variable_consideration(
            wb["Variable Consideration"]
        )
        modifications = self._read_contract_modifications(wb["Contract Modifications"])
        settings = self._read_analysis_settings(wb["Analysis Settings"])

        wb.close()

        # Attach POs, VCs, and modifications to their parent contracts
        po_by_contract: dict[str, list] = {}
        for po in performance_obligations:
            po_by_contract.setdefault(po._contract_id, []).append(po)

        vc_by_contract: dict[str, list] = {}
        for vc in variable_considerations:
            vc_by_contract.setdefault(vc._contract_id, []).append(vc)

        mod_by_contract: dict[str, list] = {}
        for mod in modifications:
            mod_by_contract.setdefault(mod._contract_id, []).append(mod)

        for contract in contracts:
            contract.performance_obligations = po_by_contract.get(contract.id, [])
            contract.variable_consideration = vc_by_contract.get(contract.id, [])
            contract.modifications = mod_by_contract.get(contract.id, [])

        return Phase3Input(
            contracts=contracts,
            reporting_period_end=settings.get("reporting_period_end"),
            tax_year=settings.get("tax_year", 0),
            statutory_rate=settings.get("statutory_rate", 0.21),
        )

    def _read_contracts(self, ws) -> list:
        contracts = []
        for row in range(2, ws.max_row + 1):
            vals = _row_values(ws, row, 13)
            if _is_empty_row(vals, [0]):
                continue
            contracts.append(Contract(
                id=_str(vals[0]),
                customer_name=_str(vals[1]),
                description=_str(vals[2]),
                contract_date=parse_date(vals[3]),
                start_date=parse_date(vals[4]),
                end_date=parse_date(vals[5]),
                total_transaction_price=to_decimal(vals[6]),
                currency=_str(vals[7]) or "USD",
                tax_method=_str(vals[8]),
                section_451c_applicable=_parse_bool(vals[9]),
                advance_payment_amount=to_decimal(vals[10]),
                long_term_contract=_parse_bool(vals[11]),
                notes=_str(vals[12]),
            ))
        return contracts

    def _read_performance_obligations(self, ws) -> list:
        obligations = []
        for row in range(2, ws.max_row + 1):
            vals = _row_values(ws, row, 20)
            if _is_empty_row(vals, [0, 1]):
                continue
            po = PerformanceObligation(
                id=_str(vals[1]),
                description=_str(vals[2]),
                type=_str(vals[3]),
                satisfaction_pattern=_str(vals[4]),
                standalone_selling_price=to_decimal(vals[5]),
                allocated_transaction_price=to_decimal(vals[6]),
                satisfaction_date=parse_date(vals[7]),
                service_start=parse_date(vals[8]),
                service_end=parse_date(vals[9]),
                percent_complete=_opt_float(vals[10]) or 0.0,
                costs_incurred=to_decimal(vals[11]),
                total_estimated_costs=to_decimal(vals[12]),
                units_delivered=_opt_int(vals[13]) or 0,
                total_units=_opt_int(vals[14]) or 0,
                book_revenue_recognized=to_decimal(vals[15]),
                book_revenue_deferred=to_decimal(vals[16]),
                tax_revenue_recognized=to_decimal(vals[17]),
                tax_treatment_notes=_str(vals[18]),
                notes=_str(vals[19]),
            )
            # Stash contract_id for parent linkage (not a dataclass field)
            po._contract_id = _str(vals[0])
            obligations.append(po)
        return obligations

    def _read_variable_consideration(self, ws) -> list:
        items = []
        for row in range(2, ws.max_row + 1):
            vals = _row_values(ws, row, 10)
            if _is_empty_row(vals, [0]):
                continue
            vc = VariableConsideration(
                type=_str(vals[1]),
                description=_str(vals[2]),
                estimated_amount=to_decimal(vals[3]),
                constrained_amount=to_decimal(vals[4]),
                estimation_method=_str(vals[5]),
                constraint_rationale=_str(vals[6]),
                book_treatment=_str(vals[7]),
                tax_treatment=_str(vals[8]),
                book_tax_difference=to_decimal(vals[9]),
            )
            vc._contract_id = _str(vals[0])
            items.append(vc)
        return items

    def _read_contract_modifications(self, ws) -> list:
        items = []
        for row in range(2, ws.max_row + 1):
            vals = _row_values(ws, row, 9)
            if _is_empty_row(vals, [0]):
                continue
            mod = ContractModification(
                modification_date=parse_date(vals[1]),
                description=_str(vals[2]),
                accounting_treatment=_str(vals[3]),
                price_change=to_decimal(vals[4]),
                scope_change=_str(vals[5]),
                impact_on_recognition=_str(vals[6]),
                book_tax_impact=to_decimal(vals[7]),
                notes=_str(vals[8]),
            )
            mod._contract_id = _str(vals[0])
            items.append(mod)
        return items

    def _read_analysis_settings(self, ws) -> dict:
        settings = {}
        for row in range(2, ws.max_row + 1):
            key = _str(ws.cell(row=row, column=1).value)
            val = ws.cell(row=row, column=2).value
            if not key:
                continue
            if key == "Reporting Period End":
                settings["reporting_period_end"] = parse_date(val)
            elif key == "Tax Year":
                settings["tax_year"] = _opt_int(val) or 0
            elif key == "Statutory Rate":
                rate = _opt_float(val)
                settings["statutory_rate"] = rate if rate is not None else 0.21
        return settings
