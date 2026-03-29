"""Phase 2: Excel report generator for trial balance & tax return analysis."""

import os
from datetime import datetime
from decimal import Decimal
from openpyxl import Workbook

from ..utils.excel_styles import (
    apply_title, apply_section_header, apply_header_row,
    write_data_row, auto_fit_columns, format_currency_cell, format_percent_cell,
    get_risk_font, get_risk_fill,
    FONT_BODY, FONT_BODY_BOLD, FONT_TITLE,
    FILL_HIGHLIGHT_BLUE, FILL_HIGHLIGHT_GREEN, FILL_HIGHLIGHT_RED,
    THIN_BORDER, ALIGN_CENTER, ALIGN_LEFT, ALIGN_RIGHT, ALIGN_WRAP,
    FMT_CURRENCY, FMT_PERCENT,
)
from ..utils.currency_helpers import to_decimal


class Phase2ExcelReport:
    """Generates Phase 2 Excel workbook from trial balance / tax return analysis."""

    def __init__(self, analysis_results: dict, company_name: str = "", output_dir: str = "output"):
        self.results = analysis_results
        self.company_name = company_name
        self.output_dir = output_dir
        self.wb = Workbook()

    def generate(self) -> str:
        """Generate the full workbook and return the file path."""
        self._create_executive_summary()
        self._create_book_tax_differences()
        self._create_trial_balance_summary()
        self._create_deferred_revenue_rollforward()
        self._create_findings_detail()
        self._create_adjustment_opportunities()
        self._create_tax_return_analysis()
        self._create_phase3_roadmap()

        if "Sheet" in self.wb.sheetnames and len(self.wb.sheetnames) > 1:
            del self.wb["Sheet"]

        os.makedirs(self.output_dir, exist_ok=True)
        company = self.company_name.replace(" ", "_") or "Company"
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"RevRec_Phase2_{company}_{timestamp}.xlsx"
        filepath = os.path.join(self.output_dir, filename)
        self.wb.save(filepath)
        return filepath

    def _create_executive_summary(self):
        """Create executive summary for Phase 2."""
        ws = self.wb.active
        ws.title = "Executive Summary"
        metrics = self.results["metrics"]

        apply_title(ws, 1, 1, 6, f"Revenue Recognition Analysis — Phase 2")
        ws.cell(row=2, column=1, value=f"Trial Balance, Tax Return & Work Paper Analysis").font = FONT_BODY
        ws.cell(row=3, column=1, value=f"Generated: {datetime.now().strftime('%B %d, %Y')}").font = FONT_BODY

        row = 5
        apply_section_header(ws, row, 1, 6, "Phase 2 Summary")
        row += 1

        summary_items = [
            ("Total Findings", metrics.get("finding_count", 0)),
            ("Book-Tax Differences Identified", metrics.get("book_tax_diff_count", 0)),
            ("Adjustment Opportunities", metrics.get("adjustment_count", 0)),
            ("Total Temporary Differences", f"${float(metrics.get('total_temporary_differences', 0)):,.0f}"),
            ("Total Permanent Differences", f"${float(metrics.get('total_permanent_differences', 0)):,.0f}"),
            ("Est. Tax Impact (Temporary)", f"${float(metrics.get('estimated_tax_impact_temporary', 0)):,.0f}"),
            ("Total Adjustment Opportunity", f"${float(metrics.get('total_adjustment_opportunity', 0)):,.0f}"),
            ("Est. Tax Savings Opportunity", f"${float(metrics.get('total_tax_savings_opportunity', 0)):,.0f}"),
        ]

        for label, value in summary_items:
            ws.cell(row=row, column=1, value=label).font = FONT_BODY_BOLD
            ws.cell(row=row, column=2, value=value).font = FONT_BODY
            ws.cell(row=row, column=1).border = THIN_BORDER
            ws.cell(row=row, column=2).border = THIN_BORDER
            if "Savings" in label or "Adjustment Opportunity" in label:
                ws.cell(row=row, column=2).font = get_risk_font("high")
            row += 1

        auto_fit_columns(ws)
        ws.sheet_properties.tabColor = "2E75B6"

    def _create_book_tax_differences(self):
        """Create book-tax differences detail sheet."""
        ws = self.wb.create_sheet("Book-Tax Differences")
        apply_title(ws, 1, 1, 7, "Book vs. Tax Revenue Recognition Differences")

        row = 3
        headers = ["Item", "Book Amount", "Tax Amount", "Difference", "Type", "Explanation", "Action"]
        for i, h in enumerate(headers):
            ws.cell(row=row, column=i + 1, value=h)
        apply_header_row(ws, row, 1, len(headers))
        row += 1

        for idx, diff in enumerate(self.results["book_tax_differences"]):
            data = [
                diff["item"],
                float(to_decimal(diff["book_amount"])),
                float(to_decimal(diff["tax_amount"])),
                float(to_decimal(diff["difference"])),
                diff.get("type", ""),
                diff.get("explanation", ""),
                diff.get("action", ""),
            ]
            write_data_row(ws, row, data, alternate=(idx % 2 == 0))

            for col in [2, 3, 4]:
                format_currency_cell(ws.cell(row=row, column=col))

            ws.cell(row=row, column=6).alignment = ALIGN_WRAP
            ws.cell(row=row, column=7).alignment = ALIGN_WRAP
            ws.row_dimensions[row].height = 40
            row += 1

        # Totals
        row += 1
        ws.cell(row=row, column=1, value="TOTAL — Temporary").font = FONT_BODY_BOLD
        temp_total = sum(
            to_decimal(d["difference"]) for d in self.results["book_tax_differences"]
            if d.get("type", "").lower() == "temporary"
        )
        cell = ws.cell(row=row, column=4, value=float(temp_total))
        format_currency_cell(cell)
        cell.font = FONT_BODY_BOLD
        row += 1

        ws.cell(row=row, column=1, value="TOTAL — Permanent").font = FONT_BODY_BOLD
        perm_total = sum(
            to_decimal(d["difference"]) for d in self.results["book_tax_differences"]
            if d.get("type", "").lower() == "permanent"
        )
        cell = ws.cell(row=row, column=4, value=float(perm_total))
        format_currency_cell(cell)
        cell.font = FONT_BODY_BOLD

        auto_fit_columns(ws, max_width=50)
        ws.column_dimensions["F"].width = 40
        ws.column_dimensions["G"].width = 35

    def _create_trial_balance_summary(self):
        """Create trial balance summary sheet."""
        ws = self.wb.create_sheet("Trial Balance Summary")
        apply_title(ws, 1, 1, 6, "Trial Balance Summary by Account Type")

        row = 3
        headers = ["Account Type", "# Accounts", "Beginning Balance", "Ending Balance", "Total Debits", "Total Credits"]
        for i, h in enumerate(headers):
            ws.cell(row=row, column=i + 1, value=h)
        apply_header_row(ws, row, 1, len(headers))
        row += 1

        tb_summary = self.results.get("trial_balance_summary", {})
        for idx, (acct_type, totals) in enumerate(sorted(tb_summary.items())):
            data = [
                acct_type.replace("_", " ").title(),
                totals["count"],
                float(totals["total_beginning"]),
                float(totals["total_ending"]),
                float(totals["total_debits"]),
                float(totals["total_credits"]),
            ]
            write_data_row(ws, row, data, alternate=(idx % 2 == 0))
            for col in [3, 4, 5, 6]:
                format_currency_cell(ws.cell(row=row, column=col))
            row += 1

        auto_fit_columns(ws)

    def _create_deferred_revenue_rollforward(self):
        """Create deferred revenue rollforward sheet."""
        ws = self.wb.create_sheet("Deferred Rev Rollforward")
        apply_title(ws, 1, 1, 7, "Deferred Revenue Rollforward Analysis")

        row = 3
        headers = ["Category", "Opening Balance", "Additions", "Recognized", "Adjustments", "Closing Balance", "Notes"]
        for i, h in enumerate(headers):
            ws.cell(row=row, column=i + 1, value=h)
        apply_header_row(ws, row, 1, len(headers))
        row += 1

        total_opening = Decimal("0")
        total_additions = Decimal("0")
        total_recognized = Decimal("0")
        total_adjustments = Decimal("0")
        total_closing = Decimal("0")

        for idx, rf in enumerate(self.results.get("deferred_rollforwards", [])):
            data = [
                rf.category,
                float(rf.opening_balance),
                float(rf.additions),
                float(rf.recognized),
                float(rf.adjustments),
                float(rf.closing_balance),
                rf.notes,
            ]
            write_data_row(ws, row, data, alternate=(idx % 2 == 0))
            for col in [2, 3, 4, 5, 6]:
                format_currency_cell(ws.cell(row=row, column=col))

            total_opening += rf.opening_balance
            total_additions += rf.additions
            total_recognized += rf.recognized
            total_adjustments += rf.adjustments
            total_closing += rf.closing_balance
            row += 1

        # Totals row
        row += 1
        ws.cell(row=row, column=1, value="TOTAL").font = FONT_BODY_BOLD
        totals = [float(total_opening), float(total_additions), float(total_recognized),
                  float(total_adjustments), float(total_closing)]
        for i, val in enumerate(totals):
            cell = ws.cell(row=row, column=i + 2, value=val)
            format_currency_cell(cell)
            cell.font = FONT_BODY_BOLD

        auto_fit_columns(ws)

    def _create_findings_detail(self):
        """Create detailed findings sheet."""
        ws = self.wb.create_sheet("Findings")
        apply_title(ws, 1, 1, 6, "Phase 2 Detailed Findings")

        row = 3
        headers = ["#", "Category", "Account", "Finding", "Detail", "Risk Level", "Recommendation"]
        for i, h in enumerate(headers):
            ws.cell(row=row, column=i + 1, value=h)
        apply_header_row(ws, row, 1, len(headers))
        row += 1

        priority_order = {"High": 0, "Medium": 1, "Low": 2}
        sorted_findings = sorted(
            self.results["findings"],
            key=lambda x: priority_order.get(x.get("risk_level", "Low"), 3)
        )

        for idx, finding in enumerate(sorted_findings, 1):
            data = [
                idx,
                finding["category"],
                finding.get("account", ""),
                finding["finding"],
                finding.get("detail", ""),
                finding.get("risk_level", ""),
                finding.get("recommendation", ""),
            ]
            write_data_row(ws, row, data, alternate=(idx % 2 == 0))
            risk_cell = ws.cell(row=row, column=6)
            risk_cell.font = get_risk_font(finding.get("risk_level", ""))
            risk_cell.fill = get_risk_fill(finding.get("risk_level", ""))
            ws.cell(row=row, column=5).alignment = ALIGN_WRAP
            ws.cell(row=row, column=7).alignment = ALIGN_WRAP
            ws.row_dimensions[row].height = 45
            row += 1

        auto_fit_columns(ws, max_width=55)
        ws.column_dimensions["E"].width = 45
        ws.column_dimensions["G"].width = 40

    def _create_adjustment_opportunities(self):
        """Create adjustment opportunities sheet."""
        ws = self.wb.create_sheet("Adjustments")
        apply_title(ws, 1, 1, 7, "Revenue Recognition Adjustment Opportunities")

        row = 3
        headers = ["Type", "Description", "Est. Impact ($)", "Est. Tax Savings ($)",
                    "Timing", "Implementation", "Risk Level"]
        for i, h in enumerate(headers):
            ws.cell(row=row, column=i + 1, value=h)
        apply_header_row(ws, row, 1, len(headers))
        row += 1

        for idx, adj in enumerate(self.results["adjustments"]):
            data = [
                adj.get("type", ""),
                adj.get("description", ""),
                float(to_decimal(adj.get("estimated_impact", 0))),
                float(adj.get("tax_impact", 0)),
                adj.get("timing", ""),
                adj.get("implementation", ""),
                adj.get("risk", ""),
            ]
            write_data_row(ws, row, data, alternate=(idx % 2 == 0))
            for col in [3, 4]:
                format_currency_cell(ws.cell(row=row, column=col))
            ws.cell(row=row, column=2).alignment = ALIGN_WRAP
            ws.cell(row=row, column=6).alignment = ALIGN_WRAP
            ws.row_dimensions[row].height = 40
            row += 1

        # Total
        if self.results["adjustments"]:
            row += 1
            ws.cell(row=row, column=1, value="TOTAL").font = FONT_BODY_BOLD
            total_impact = sum(float(to_decimal(a.get("estimated_impact", 0))) for a in self.results["adjustments"])
            total_tax = sum(float(a.get("tax_impact", 0)) for a in self.results["adjustments"])
            cell = ws.cell(row=row, column=3, value=total_impact)
            format_currency_cell(cell)
            cell.font = FONT_BODY_BOLD
            cell = ws.cell(row=row, column=4, value=total_tax)
            format_currency_cell(cell)
            cell.font = FONT_BODY_BOLD

        auto_fit_columns(ws, max_width=50)
        ws.column_dimensions["B"].width = 45
        ws.column_dimensions["F"].width = 35

    def _create_tax_return_analysis(self):
        """Create tax return analysis sheet."""
        ws = self.wb.create_sheet("Tax Return Analysis")
        tr = self.results.get("tax_return")
        if not tr:
            apply_title(ws, 1, 1, 4, "Tax Return Analysis — No Data Provided")
            return

        apply_title(ws, 1, 1, 4, "Tax Return Revenue Analysis")

        row = 3
        apply_section_header(ws, row, 1, 4, "Return Information")
        row += 1
        info_fields = [
            ("Form Type", tr.form_type),
            ("Tax Year", tr.tax_year),
            ("Accounting Method", tr.accounting_method),
            ("§451(c) Election", "Yes" if tr.section_451c_election else "No"),
            ("§451(b) AFS", "Yes" if tr.section_451b_afs else "No"),
        ]
        for label, value in info_fields:
            ws.cell(row=row, column=1, value=label).font = FONT_BODY_BOLD
            ws.cell(row=row, column=2, value=value).font = FONT_BODY
            row += 1

        row += 1
        apply_section_header(ws, row, 1, 4, "Revenue / Gross Receipts")
        row += 1
        rev_fields = [
            ("Gross Receipts (Line 1a)", float(tr.gross_receipts_line)),
            ("Returns & Allowances", float(tr.returns_and_allowances)),
            ("Net Receipts", float(tr.net_receipts)),
            ("Book Income (M-1)", float(tr.book_income)),
            ("Tax Income", float(tr.tax_income)),
        ]
        for label, value in rev_fields:
            ws.cell(row=row, column=1, value=label).font = FONT_BODY_BOLD
            cell = ws.cell(row=row, column=2, value=value)
            format_currency_cell(cell)
            row += 1

        row += 1
        apply_section_header(ws, row, 1, 4, "Deferred Revenue — Tax Basis")
        row += 1
        ws.cell(row=row, column=1, value="Current Year").font = FONT_BODY_BOLD
        cell = ws.cell(row=row, column=2, value=float(tr.tax_deferred_revenue_current))
        format_currency_cell(cell)
        row += 1
        ws.cell(row=row, column=1, value="Prior Year").font = FONT_BODY_BOLD
        cell = ws.cell(row=row, column=2, value=float(tr.tax_deferred_revenue_prior))
        format_currency_cell(cell)

        # M-1 adjustments
        if tr.m1_revenue_adjustments:
            row += 2
            apply_section_header(ws, row, 1, 4, "Schedule M-1 Revenue Adjustments")
            row += 1
            m1_headers = ["Description", "Book Amount", "Tax Amount", "Difference"]
            for i, h in enumerate(m1_headers):
                ws.cell(row=row, column=i + 1, value=h)
            apply_header_row(ws, row, 1, 4)
            row += 1

            for adj in tr.m1_revenue_adjustments:
                book_amt = float(to_decimal(adj.get("book_amount", 0)))
                tax_amt = float(to_decimal(adj.get("tax_amount", 0)))
                data = [
                    adj.get("description", ""),
                    book_amt,
                    tax_amt,
                    book_amt - tax_amt,
                ]
                write_data_row(ws, row, data)
                for col in [2, 3, 4]:
                    format_currency_cell(ws.cell(row=row, column=col))
                row += 1

        auto_fit_columns(ws)

    def _create_phase3_roadmap(self):
        """Create Phase 3 investigation roadmap."""
        ws = self.wb.create_sheet("Phase 3 Roadmap")
        apply_title(ws, 1, 1, 5, "Phase 3: Contract-Level Analysis Roadmap")
        ws.cell(row=2, column=1, value="Full contract analysis with ASC 606 5-step model").font = FONT_BODY

        row = 4
        apply_section_header(ws, row, 1, 5, "Required Documents for Phase 3")
        row += 1
        documents = [
            ("Customer Contracts", "All material revenue contracts",
             "Apply 5-step ASC 606 model; identify performance obligations"),
            ("Contract Amendments", "Modifications and change orders",
             "Assess modification accounting impact on book and tax"),
            ("Revenue Detail by Contract", "Monthly recognition by contract",
             "Validate recognition timing; test cutoff procedures"),
            ("Variable Consideration Docs", "Rebate, discount, penalty calculations",
             "Quantify constraint impact; compare book vs. tax treatment"),
            ("Standalone Selling Prices", "SSP evidence and allocation",
             "Validate price allocation across bundled arrangements"),
        ]

        headers = ["Document", "Description", "Phase 3 Analysis"]
        for i, h in enumerate(headers):
            ws.cell(row=row, column=i + 1, value=h)
        apply_header_row(ws, row, 1, 3)
        row += 1

        for idx, (doc, desc, analysis) in enumerate(documents):
            write_data_row(ws, row, [doc, desc, analysis], alternate=(idx % 2 == 0))
            ws.cell(row=row, column=3).alignment = ALIGN_WRAP
            ws.row_dimensions[row].height = 35
            row += 1

        # Key items to carry forward from Phase 2
        row += 1
        apply_section_header(ws, row, 1, 5, "High-Priority Items to Investigate in Phase 3")
        row += 1

        headers = ["Priority", "Finding from Phase 2", "Contract-Level Action"]
        for i, h in enumerate(headers):
            ws.cell(row=row, column=i + 1, value=h)
        apply_header_row(ws, row, 1, 3)
        row += 1

        high_findings = [f for f in self.results["findings"] if f.get("risk_level") == "High"]
        for idx, finding in enumerate(high_findings):
            data = [
                finding["risk_level"],
                finding["finding"],
                f"Review underlying contracts for: {finding.get('recommendation', '')}",
            ]
            write_data_row(ws, row, data, alternate=(idx % 2 == 0))
            ws.cell(row=row, column=1).font = get_risk_font("High")
            ws.cell(row=row, column=1).fill = get_risk_fill("High")
            ws.cell(row=row, column=3).alignment = ALIGN_WRAP
            ws.row_dimensions[row].height = 35
            row += 1

        auto_fit_columns(ws, max_width=50)
        ws.column_dimensions["C"].width = 50
