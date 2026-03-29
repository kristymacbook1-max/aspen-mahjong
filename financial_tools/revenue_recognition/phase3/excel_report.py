"""Phase 3: Excel report generator for full contract-level revenue recognition analysis."""

import os
from datetime import datetime
from decimal import Decimal
from openpyxl import Workbook

from ..utils.excel_styles import (
    apply_title, apply_section_header, apply_header_row,
    write_data_row, auto_fit_columns, format_currency_cell, format_percent_cell,
    get_risk_font, get_risk_fill,
    FONT_BODY, FONT_BODY_BOLD, FONT_TITLE,
    FILL_HIGHLIGHT_BLUE, FILL_HIGHLIGHT_GREEN, FILL_HIGHLIGHT_RED, FILL_HIGHLIGHT_ORANGE,
    THIN_BORDER, BOTTOM_BORDER, ALIGN_CENTER, ALIGN_LEFT, ALIGN_RIGHT, ALIGN_WRAP,
    FMT_CURRENCY, FMT_PERCENT,
)
from ..utils.currency_helpers import to_decimal


class Phase3ExcelReport:
    """Generates Phase 3 Excel workbook from contract-level analysis."""

    def __init__(self, analysis_results: dict, company_name: str = "", output_dir: str = "output"):
        self.results = analysis_results
        self.company_name = company_name
        self.output_dir = output_dir
        self.wb = Workbook()

    def generate(self) -> str:
        """Generate the full workbook and return the file path."""
        self._create_executive_summary()
        self._create_contract_portfolio()
        self._create_book_tax_reconciliation()
        self._create_asc606_detail()
        self._create_recognition_schedule()
        self._create_variable_consideration()
        self._create_aggregate_findings()
        self._create_implementation_plan()

        if "Sheet" in self.wb.sheetnames and len(self.wb.sheetnames) > 1:
            del self.wb["Sheet"]

        os.makedirs(self.output_dir, exist_ok=True)
        company = self.company_name.replace(" ", "_") or "Company"
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"RevRec_Phase3_{company}_{timestamp}.xlsx"
        filepath = os.path.join(self.output_dir, filename)
        self.wb.save(filepath)
        return filepath

    def _create_executive_summary(self):
        """Create executive summary for Phase 3."""
        ws = self.wb.active
        ws.title = "Executive Summary"
        metrics = self.results["metrics"]

        apply_title(ws, 1, 1, 6, "Revenue Recognition — Phase 3 Contract Analysis")
        ws.cell(row=2, column=1, value="Full ASC 606 5-Step Model Application").font = FONT_BODY
        ws.cell(row=3, column=1, value=f"Generated: {datetime.now().strftime('%B %d, %Y')}").font = FONT_BODY

        row = 5
        apply_section_header(ws, row, 1, 6, "Portfolio Summary")
        row += 1

        summary_items = [
            ("Total Contracts Analyzed", metrics.get("total_contracts", 0)),
            ("Total Contract Value", f"${float(metrics.get('total_contract_value', 0)):,.0f}"),
            ("Total Book Revenue Recognized", f"${float(metrics.get('total_book_revenue', 0)):,.0f}"),
            ("Total Tax Revenue Recognized", f"${float(metrics.get('total_tax_revenue', 0)):,.0f}"),
            ("Net Book-Tax Difference", f"${float(metrics.get('total_book_tax_difference', 0)):,.0f}"),
            ("Tax Impact at Statutory Rate", f"${float(metrics.get('tax_impact_at_statutory', 0)):,.0f}"),
            ("Total Deferred Revenue", f"${float(metrics.get('total_deferred_revenue', 0)):,.0f}"),
            ("Overall Recognition Rate", f"{metrics.get('overall_recognition_rate', 0):.1%}"),
            ("Total Findings", metrics.get("total_findings", 0)),
            ("High Priority Findings", metrics.get("high_priority_findings", 0)),
        ]

        for label, value in summary_items:
            ws.cell(row=row, column=1, value=label).font = FONT_BODY_BOLD
            cell = ws.cell(row=row, column=2, value=value)
            cell.font = FONT_BODY
            ws.cell(row=row, column=1).border = THIN_BORDER
            cell.border = THIN_BORDER
            if "Tax Impact" in label or "Book-Tax Difference" in label:
                cell.font = get_risk_font("high")
            row += 1

        auto_fit_columns(ws)
        ws.sheet_properties.tabColor = "548235"

    def _create_contract_portfolio(self):
        """Create contract portfolio overview sheet."""
        ws = self.wb.create_sheet("Contract Portfolio")
        apply_title(ws, 1, 1, 9, "Contract Portfolio Summary")

        row = 3
        headers = [
            "Contract ID", "Customer", "Total Value",
            "Book Recognized", "Tax Recognized", "Book-Tax Diff",
            "Tax Impact", "# Obligations", "Tax Method"
        ]
        for i, h in enumerate(headers):
            ws.cell(row=row, column=i + 1, value=h)
        apply_header_row(ws, row, 1, len(headers))
        row += 1

        statutory_rate = 0.21
        for idx, ca in enumerate(self.results["contract_analyses"]):
            book_tax = float(to_decimal(ca["book_tax_difference"]))
            tax_impact = book_tax * statutory_rate
            data = [
                ca["contract_id"],
                ca["customer"],
                float(to_decimal(ca["total_price"])),
                float(to_decimal(ca["book_revenue_total"])),
                float(to_decimal(ca["tax_revenue_total"])),
                book_tax,
                tax_impact,
                ca["step2"]["obligation_count"],
                ca["tax_analysis"]["tax_method"],
            ]
            write_data_row(ws, row, data, alternate=(idx % 2 == 0))
            for col in [3, 4, 5, 6, 7]:
                format_currency_cell(ws.cell(row=row, column=col))
            row += 1

        # Totals
        row += 1
        ws.cell(row=row, column=1, value="TOTAL").font = FONT_BODY_BOLD
        metrics = self.results["metrics"]
        total_cells = [
            (3, float(metrics.get("total_contract_value", 0))),
            (4, float(metrics.get("total_book_revenue", 0))),
            (5, float(metrics.get("total_tax_revenue", 0))),
            (6, float(metrics.get("total_book_tax_difference", 0))),
            (7, float(metrics.get("tax_impact_at_statutory", 0))),
        ]
        for col, val in total_cells:
            cell = ws.cell(row=row, column=col, value=val)
            format_currency_cell(cell)
            cell.font = FONT_BODY_BOLD

        auto_fit_columns(ws)

    def _create_book_tax_reconciliation(self):
        """Create book-to-tax revenue reconciliation sheet."""
        ws = self.wb.create_sheet("Book-Tax Reconciliation")
        apply_title(ws, 1, 1, 3, "Book-to-Tax Revenue Reconciliation")

        row = 3
        headers = ["Line Item", "Amount"]
        for i, h in enumerate(headers):
            ws.cell(row=row, column=i + 1, value=h)
        apply_header_row(ws, row, 1, 2)
        row += 1

        for item in self.results["book_tax_reconciliation"]:
            ws.cell(row=row, column=1, value=item["line"]).font = FONT_BODY
            cell = ws.cell(row=row, column=2, value=float(to_decimal(item["amount"])))
            format_currency_cell(cell)
            ws.cell(row=row, column=1).border = THIN_BORDER
            cell.border = THIN_BORDER

            # Bold totals and key lines
            if item["line"].startswith("Total") or item["line"].startswith("Net") or item["line"].startswith("Tax Impact"):
                ws.cell(row=row, column=1).font = FONT_BODY_BOLD
                cell.font = FONT_BODY_BOLD
                if "Difference" in item["line"] or "Tax Impact" in item["line"]:
                    cell.font = get_risk_font("high")

            row += 1

        auto_fit_columns(ws)
        ws.column_dimensions["A"].width = 55

    def _create_asc606_detail(self):
        """Create per-contract ASC 606 5-step detail sheet."""
        ws = self.wb.create_sheet("ASC 606 Detail")
        apply_title(ws, 1, 1, 7, "ASC 606 Five-Step Analysis by Contract")

        row = 3
        for ca in self.results["contract_analyses"]:
            # Contract header
            apply_section_header(
                ws, row, 1, 7,
                f"{ca['contract_id']} — {ca['customer']} (${float(to_decimal(ca['total_price'])):,.0f})"
            )
            row += 1

            # Steps summary
            steps = [
                ("Step 1: Identify Contract", ca["step1"]["assessment"]),
                ("Step 2: Performance Obligations", f"{ca['step2']['obligation_count']} identified"),
                ("Step 3: Transaction Price", f"${float(to_decimal(ca['step3']['final_transaction_price'])):,.0f}"),
                ("Step 4: Allocation Method", ca["step4"]["allocation_method"]),
                ("Step 5: Recognition", f"{ca['step5']['completion_percentage']:.1%} complete"),
            ]
            for label, value in steps:
                ws.cell(row=row, column=1, value=label).font = FONT_BODY_BOLD
                ws.cell(row=row, column=2, value=value).font = FONT_BODY
                row += 1

            # Findings for this contract
            if ca["findings"]:
                row += 1
                finding_headers = ["Step", "Finding", "Risk", "Recommendation"]
                for i, h in enumerate(finding_headers):
                    ws.cell(row=row, column=i + 1, value=h)
                apply_header_row(ws, row, 1, 4)
                row += 1

                for finding in ca["findings"]:
                    data = [
                        finding.get("step", ""),
                        finding["finding"],
                        finding.get("risk_level", ""),
                        finding.get("recommendation", ""),
                    ]
                    write_data_row(ws, row, data)
                    ws.cell(row=row, column=3).font = get_risk_font(finding.get("risk_level", ""))
                    ws.cell(row=row, column=3).fill = get_risk_fill(finding.get("risk_level", ""))
                    ws.cell(row=row, column=4).alignment = ALIGN_WRAP
                    ws.row_dimensions[row].height = 35
                    row += 1

            row += 2  # Space between contracts

        auto_fit_columns(ws, max_width=55)
        ws.column_dimensions["B"].width = 45
        ws.column_dimensions["D"].width = 50

    def _create_recognition_schedule(self):
        """Create revenue recognition schedule sheet."""
        ws = self.wb.create_sheet("Recognition Schedule")
        apply_title(ws, 1, 1, 8, "Revenue Recognition Schedule by Obligation")

        row = 3
        headers = [
            "Contract", "Customer", "Obligation", "Pattern",
            "Allocated Price", "Recognized", "Deferred", "% Complete"
        ]
        for i, h in enumerate(headers):
            ws.cell(row=row, column=i + 1, value=h)
        apply_header_row(ws, row, 1, len(headers))
        row += 1

        for ca in self.results["contract_analyses"]:
            for idx, sched in enumerate(ca["recognition_schedule"]):
                data = [
                    ca["contract_id"],
                    ca["customer"],
                    sched["description"],
                    sched["pattern"],
                    float(to_decimal(sched["allocated_price"])),
                    float(to_decimal(sched["recognized"])),
                    float(to_decimal(sched["deferred"])),
                    sched.get("percent_complete", 0),
                ]
                write_data_row(ws, row, data, alternate=(row % 2 == 0))
                for col in [5, 6, 7]:
                    format_currency_cell(ws.cell(row=row, column=col))
                format_percent_cell(ws.cell(row=row, column=8))
                row += 1

        auto_fit_columns(ws)

    def _create_variable_consideration(self):
        """Create variable consideration analysis sheet."""
        ws = self.wb.create_sheet("Variable Consideration")
        apply_title(ws, 1, 1, 8, "Variable Consideration Analysis")

        row = 3
        headers = [
            "Contract", "Type", "Description", "Estimated Amount",
            "Constrained Amount", "Excluded", "Method", "Book-Tax Diff"
        ]
        for i, h in enumerate(headers):
            ws.cell(row=row, column=i + 1, value=h)
        apply_header_row(ws, row, 1, len(headers))
        row += 1

        has_vc = False
        for contract in self.results["contracts"]:
            for vc in contract.variable_consideration:
                has_vc = True
                excluded = float(vc.estimated_amount - vc.constrained_amount)
                data = [
                    contract.id,
                    vc.type,
                    vc.description,
                    float(vc.estimated_amount),
                    float(vc.constrained_amount),
                    excluded,
                    vc.estimation_method,
                    float(vc.book_tax_difference),
                ]
                write_data_row(ws, row, data, alternate=(row % 2 == 0))
                for col in [4, 5, 6, 8]:
                    format_currency_cell(ws.cell(row=row, column=col))
                row += 1

        if not has_vc:
            ws.cell(row=row, column=1, value="No variable consideration elements identified.").font = FONT_BODY

        auto_fit_columns(ws)

    def _create_aggregate_findings(self):
        """Create aggregate findings sheet."""
        ws = self.wb.create_sheet("Aggregate Findings")
        apply_title(ws, 1, 1, 5, "Portfolio-Level Findings & Opportunities")

        row = 3
        headers = ["Category", "Finding", "Tax Impact", "Recommendation", "Risk Level"]
        for i, h in enumerate(headers):
            ws.cell(row=row, column=i + 1, value=h)
        apply_header_row(ws, row, 1, len(headers))
        row += 1

        for idx, finding in enumerate(self.results["aggregate_findings"]):
            tax_impact = finding.get("tax_impact")
            tax_str = f"${tax_impact:,.0f}" if tax_impact else "To be determined"
            data = [
                finding["category"],
                finding["finding"],
                tax_str,
                finding.get("recommendation", ""),
                finding.get("risk_level", ""),
            ]
            write_data_row(ws, row, data, alternate=(idx % 2 == 0))
            ws.cell(row=row, column=5).font = get_risk_font(finding.get("risk_level", ""))
            ws.cell(row=row, column=5).fill = get_risk_fill(finding.get("risk_level", ""))
            ws.cell(row=row, column=4).alignment = ALIGN_WRAP
            ws.row_dimensions[row].height = 40
            row += 1

        auto_fit_columns(ws, max_width=50)
        ws.column_dimensions["B"].width = 50
        ws.column_dimensions["D"].width = 45

    def _create_implementation_plan(self):
        """Create implementation plan sheet with actionable next steps."""
        ws = self.wb.create_sheet("Implementation Plan")
        apply_title(ws, 1, 1, 6, "Revenue Recognition — Implementation Plan")

        row = 3
        apply_section_header(ws, row, 1, 6, "Recommended Actions")
        row += 1

        headers = ["Priority", "Action", "Description", "Est. Tax Savings", "Timeline", "Complexity"]
        for i, h in enumerate(headers):
            ws.cell(row=row, column=i + 1, value=h)
        apply_header_row(ws, row, 1, len(headers))
        row += 1

        # Build action items from aggregate findings + contract findings
        actions = []

        for finding in self.results["aggregate_findings"]:
            if finding.get("risk_level") == "High":
                actions.append({
                    "priority": "High",
                    "action": finding["category"],
                    "description": finding["recommendation"],
                    "tax_savings": finding.get("tax_impact"),
                    "timeline": "Immediate — current tax year",
                    "complexity": "Medium",
                })

        for ca in self.results["contract_analyses"]:
            high_findings = [f for f in ca["findings"] if f.get("risk_level") == "High"]
            for finding in high_findings:
                actions.append({
                    "priority": "High",
                    "action": f"{ca['customer']} — {finding.get('step', '')}",
                    "description": finding["recommendation"],
                    "tax_savings": None,
                    "timeline": "Current period",
                    "complexity": "Medium",
                })

        # Deduplicate by action
        seen = set()
        for action in actions:
            key = action["action"]
            if key in seen:
                continue
            seen.add(key)

            savings_str = f"${action['tax_savings']:,.0f}" if action["tax_savings"] else "TBD"
            data = [
                action["priority"],
                action["action"],
                action["description"],
                savings_str,
                action["timeline"],
                action["complexity"],
            ]
            write_data_row(ws, row, data, alternate=(len(seen) % 2 == 0))
            ws.cell(row=row, column=1).font = get_risk_font(action["priority"])
            ws.cell(row=row, column=1).fill = get_risk_fill(action["priority"])
            ws.cell(row=row, column=3).alignment = ALIGN_WRAP
            ws.row_dimensions[row].height = 40
            row += 1

        # Filing requirements
        row += 2
        apply_section_header(ws, row, 1, 6, "Filing Requirements")
        row += 1
        filings = [
            ("Form 3115", "Change in accounting method (§451(c) election, method changes)",
             "File with current year return or in advance"),
            ("Amended Returns", "If prior-year positions should be corrected",
             "Generally 3-year lookback from filing date"),
            ("Schedule M-1/M-3", "Update book-tax reconciliation for identified differences",
             "Include with current year return"),
            ("Documentation", "Support for all positions taken",
             "Maintain contemporaneous documentation"),
        ]
        headers = ["Form/Filing", "Purpose", "Timing"]
        for i, h in enumerate(headers):
            ws.cell(row=row, column=i + 1, value=h)
        apply_header_row(ws, row, 1, 3)
        row += 1

        for idx, (form, purpose, timing) in enumerate(filings):
            write_data_row(ws, row, [form, purpose, timing], alternate=(idx % 2 == 0))
            ws.cell(row=row, column=2).alignment = ALIGN_WRAP
            ws.row_dimensions[row].height = 35
            row += 1

        auto_fit_columns(ws, max_width=50)
        ws.column_dimensions["C"].width = 45
        ws.column_dimensions["B"].width = 45
