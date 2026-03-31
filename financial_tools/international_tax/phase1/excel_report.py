"""Phase 1 International Tax Excel Report Generator.

Produces a 7-tab workbook:
  1. Executive Summary
  2. GILTI Analysis
  3. FDII Analysis
  4. FTC Limitation
  5. Subpart F
  6. CFC Detail
  7. Phase 2 Roadmap
"""

import os, sys
from decimal import Decimal
from openpyxl import Workbook

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', '..'))
from revenue_recognition.utils.excel_styles import (
    FONT_TITLE, FONT_BODY, FONT_BODY_BOLD, THIN_BORDER, FMT_CURRENCY,
    apply_header_row, apply_section_header, apply_title,
    auto_fit_columns, write_data_row, format_currency_cell, get_risk_font,
)
from revenue_recognition.utils.currency_helpers import to_decimal, format_currency


class Phase1InternationalReport:

    def generate(self, results: dict, output_path: str) -> str:
        self._r = results
        self._wb = Workbook()
        self._wb.remove(self._wb.active)

        self._create_executive_summary()
        self._create_gilti()
        self._create_fdii()
        self._create_ftc()
        self._create_subpart_f()
        self._create_cfc_detail()
        self._create_roadmap()

        os.makedirs(os.path.dirname(output_path) or '.', exist_ok=True)
        self._wb.save(output_path)
        return output_path

    def _create_executive_summary(self):
        ws = self._wb.create_sheet("Executive Summary")
        apply_title(ws, 1, 1, 5, f"International Tax Analysis — {self._r['company'].name}")
        row = 3

        gilti = self._r["gilti_analysis"]
        fdii = self._r["fdii_analysis"]
        ftc = self._r["ftc_analysis"]
        info = [
            ("GILTI Inclusion", format_currency(gilti["gilti_inclusion"])),
            ("§250 GILTI Deduction", format_currency(gilti["gilti_deduction"])),
            ("GILTI Effective Rate", f"{gilti['effective_rate']*100:.1f}%"),
            ("FDII Amount", format_currency(fdii["fdii_amount"])),
            ("FDII Deduction", format_currency(fdii["fdii_deduction"])),
            ("FDII Effective Rate", f"{fdii['effective_rate']*100:.3f}%"),
            ("Total FTC Allowed", format_currency(ftc["total_credits_allowed"])),
            ("Excess FTCs", format_currency(ftc["total_excess_credits"])),
            ("CFC Count", str(gilti["cfc_count"])),
        ]
        for label, val in info:
            ws.cell(row=row, column=1, value=label).font = FONT_BODY_BOLD
            ws.cell(row=row, column=2, value=val).font = FONT_BODY
            row += 1

        row += 1
        apply_section_header(ws, row, 1, 5, "Key Findings")
        row += 1
        for f in self._r["findings"]:
            ws.cell(row=row, column=1, value=f["area"]).font = FONT_BODY_BOLD
            ws.cell(row=row, column=2, value=f["description"]).font = FONT_BODY
            ws.cell(row=row, column=3, value=f["risk_level"].upper()).font = get_risk_font(f["risk_level"])
            row += 1
        auto_fit_columns(ws)

    def _create_gilti(self):
        ws = self._wb.create_sheet("GILTI Analysis")
        apply_title(ws, 1, 1, 3, "§951A GILTI Computation")
        row = 3

        g = self._r["gilti_analysis"]
        lines = [
            ("Total Tested Income", g["total_tested_income"]),
            ("Total Tested Loss", g["total_tested_loss"]),
            ("Net CFC Tested Income", g["net_tested_income"]),
            ("", ""),
            ("Total QBAI", g["total_qbai"]),
            ("Deemed Tangible Return (10% × QBAI)", g["deemed_tangible_return"]),
            ("", ""),
            ("GILTI Inclusion", g["gilti_inclusion"]),
            ("§78 Gross-Up", g["section_78_gross_up"]),
            (f"§250 Deduction ({g['section_250_deduction_rate']*100:.0f}%)", g["gilti_deduction"]),
            ("Taxable GILTI", g["taxable_gilti"]),
            ("Effective Rate", f"{g['effective_rate']*100:.1f}%"),
        ]
        for label, val in lines:
            if not label:
                row += 1
                continue
            ws.cell(row=row, column=1, value=label).font = FONT_BODY_BOLD
            if isinstance(val, str):
                ws.cell(row=row, column=2, value=val).font = FONT_BODY
            else:
                c = ws.cell(row=row, column=2, value=float(to_decimal(val)))
                format_currency_cell(c)
            row += 1
        auto_fit_columns(ws)

    def _create_fdii(self):
        ws = self._wb.create_sheet("FDII Analysis")
        apply_title(ws, 1, 1, 3, "§250 Foreign-Derived Intangible Income")
        row = 3

        f = self._r["fdii_analysis"]
        lines = [
            ("Deduction Eligible Income (DEI)", f["dei"]),
            ("Foreign-Derived DEI (FDDEI)", f["fddei"]),
            ("Foreign-Derived Ratio", f"{f['foreign_ratio']*100:.1f}%"),
            ("", ""),
            ("Domestic QBAI", f["domestic_qbai"]),
            ("Deemed Tangible Return (10% × QBAI)", f["deemed_tangible_return"]),
            ("Deemed Intangible Income", f["deemed_intangible_income"]),
            ("", ""),
            ("FDII Amount", f["fdii_amount"]),
            (f"FDII Deduction ({f['deduction_rate']*100:.1f}%)", f["fdii_deduction"]),
            ("Effective Rate on FDII", f"{f['effective_rate']*100:.3f}%"),
            ("Tax Savings", f["tax_savings"]),
        ]
        for label, val in lines:
            if not label:
                row += 1
                continue
            ws.cell(row=row, column=1, value=label).font = FONT_BODY_BOLD
            if isinstance(val, str):
                ws.cell(row=row, column=2, value=val).font = FONT_BODY
            else:
                c = ws.cell(row=row, column=2, value=float(to_decimal(val)))
                format_currency_cell(c)
            row += 1
        auto_fit_columns(ws)

    def _create_ftc(self):
        ws = self._wb.create_sheet("FTC Limitation")
        apply_title(ws, 1, 1, 7, "§904 Foreign Tax Credit Limitation by Basket")
        row = 3

        headers = ["Category", "Foreign Source Income", "Foreign Taxes Paid",
                    "FTC Limitation", "Credits Allowed", "Excess Credits", "Eff. Foreign Rate"]
        for i, h in enumerate(headers, 1):
            ws.cell(row=row, column=i, value=h)
        apply_header_row(ws, row, 1, len(headers))
        row += 1

        for basket in self._r["ftc_analysis"]["baskets"]:
            write_data_row(ws, row, [
                basket["category"],
                float(to_decimal(basket["foreign_source_income"])),
                float(to_decimal(basket["foreign_taxes_paid"])),
                float(to_decimal(basket["limitation"])),
                float(to_decimal(basket["credits_allowed"])),
                float(to_decimal(basket["excess_credits"])),
                f"{basket['effective_foreign_rate']*100:.1f}%",
            ], alternate=(row % 2 == 0))
            for c in [2, 3, 4, 5, 6]:
                format_currency_cell(ws.cell(row=row, column=c))
            row += 1

        row += 1
        ftc = self._r["ftc_analysis"]
        ws.cell(row=row, column=1, value="US Tax Before FTC").font = FONT_BODY_BOLD
        c = ws.cell(row=row, column=2, value=float(to_decimal(ftc["us_tax_before_ftc"])))
        format_currency_cell(c)
        row += 1
        ws.cell(row=row, column=1, value="Total Credits Allowed").font = FONT_BODY_BOLD
        c = ws.cell(row=row, column=2, value=float(to_decimal(ftc["total_credits_allowed"])))
        format_currency_cell(c)
        row += 1
        ws.cell(row=row, column=1, value="Net US Tax").font = FONT_BODY_BOLD
        c = ws.cell(row=row, column=2, value=float(to_decimal(ftc["net_us_tax"])))
        format_currency_cell(c)
        c.font = FONT_BODY_BOLD
        auto_fit_columns(ws)

    def _create_subpart_f(self):
        ws = self._wb.create_sheet("Subpart F")
        apply_title(ws, 1, 1, 4, "Subpart F Income Inclusions")
        row = 3

        sf = self._r["subpart_f_analysis"]
        if not sf["items"]:
            ws.cell(row=row, column=1, value="No Subpart F income identified.").font = FONT_BODY
        else:
            headers = ["CFC", "Jurisdiction", "Subpart F Income", "Foreign Taxes"]
            for i, h in enumerate(headers, 1):
                ws.cell(row=row, column=i, value=h)
            apply_header_row(ws, row, 1, len(headers))
            row += 1

            for item in sf["items"]:
                write_data_row(ws, row, [
                    item["cfc_name"], item["jurisdiction"],
                    float(to_decimal(item["subpart_f_income"])),
                    float(to_decimal(item["foreign_taxes"])),
                ], alternate=(row % 2 == 0))
                format_currency_cell(ws.cell(row=row, column=3))
                format_currency_cell(ws.cell(row=row, column=4))
                row += 1

            row += 1
            ws.cell(row=row, column=1, value="Total Subpart F").font = FONT_BODY_BOLD
            c = ws.cell(row=row, column=3, value=float(to_decimal(sf["total_subpart_f"])))
            format_currency_cell(c)
            c.font = FONT_BODY_BOLD
        auto_fit_columns(ws)

    def _create_cfc_detail(self):
        ws = self._wb.create_sheet("CFC Detail")
        apply_title(ws, 1, 1, 6, "CFC-by-CFC GILTI Detail")
        row = 3

        cfc_details = self._r["gilti_analysis"]["cfc_details"]
        if not cfc_details:
            ws.cell(row=row, column=1, value="No CFC data provided.").font = FONT_BODY
        else:
            headers = ["CFC Name", "Jurisdiction", "Tested Income", "Tested Loss", "QBAI", "Foreign Taxes"]
            for i, h in enumerate(headers, 1):
                ws.cell(row=row, column=i, value=h)
            apply_header_row(ws, row, 1, len(headers))
            row += 1

            for cfc in cfc_details:
                write_data_row(ws, row, [
                    cfc["name"], cfc["jurisdiction"],
                    float(to_decimal(cfc["tested_income"])),
                    float(to_decimal(cfc["tested_loss"])),
                    float(to_decimal(cfc["qbai"])),
                    float(to_decimal(cfc["foreign_taxes"])),
                ], alternate=(row % 2 == 0))
                for c in [3, 4, 5, 6]:
                    format_currency_cell(ws.cell(row=row, column=c))
                row += 1
        auto_fit_columns(ws)

    def _create_roadmap(self):
        ws = self._wb.create_sheet("Phase 2 Roadmap")
        apply_title(ws, 1, 1, 4, "Phase 2 Recommendations")
        row = 3

        headers = ["Priority", "Area", "Description", "Data Needed"]
        for i, h in enumerate(headers, 1):
            ws.cell(row=row, column=i, value=h)
        apply_header_row(ws, row, 1, len(headers))
        row += 1

        for item in self._r["phase2_roadmap"]:
            write_data_row(ws, row, [
                item["priority"], item["area"], item["description"], item["data_needed"],
            ], alternate=(row % 2 == 0))
            row += 1
        auto_fit_columns(ws)
