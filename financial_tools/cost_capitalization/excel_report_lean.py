"""Streamlined (lean) Cost Capitalization Excel Report.

One self-contained working tab — "263A Cost Cap" — that does the whole calc:
parameters, the §263A absorption-ratio block, the auto-classified trial-balance
grid with one column per capitalization provision, and a Schedule M-1 strip.

The classifier's output (Provision / Cap % / Basis / Confidence) is written as
values; the capitalization split columns are LIVE Excel formulas off
Provision + Cap % + Amount, so an analyst override recomputes everything.
The §263A reg + IRS Practice Unit knowledge rides on each line as its Basis.
"""

import os
import sys

from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..'))
from revenue_recognition.utils.excel_styles import (
    FONT_BODY, FONT_BODY_BOLD,
    FILL_HIGHLIGHT_BLUE, FILL_HIGHLIGHT_GREEN, FILL_HIGHLIGHT_ORANGE, FILL_HIGHLIGHT_RED,
    THIN_BORDER, ALIGN_LEFT, ALIGN_RIGHT,
    FMT_CURRENCY, FMT_PERCENT,
    apply_title, apply_section_header, apply_header_row,
)

PROVISION_LIST = ('"book_capitalized,mixed,deductible,266,263(a),263A,'
                  'other:174A,other:197,other:195,other:248,other:709,'
                  'other:263(c),other:263(g),other:461(g)"')

GRID_HEADERS = [
    "Account #", "Account Name", "Cost Center", "Provision", "Cap %", "Amount",
    "§266", "§263(a)", "§263A", "Other Cap", "Book-Cap", "Deductible",
    "Basis (section / IRS Practice Unit)", "Conf",
]
BLANK_TEMPLATE_ROWS = 5


class CostCapitalizationLeanReport:
    """Generates the one-tab lean cost capitalization workbook."""

    def generate(self, results: dict, output_path: str) -> str:
        self._r = results
        self._inp = results["_input"]
        self._wb = Workbook()
        ws = self._wb.active
        ws.title = "263A Cost Cap"
        self._ws = ws

        apply_title(ws, 1, 1, 14,
                    f"§266 / §263(a) / §263A Cost Capitalization — {results['company_name']} "
                    f"(TY {results['tax_year']})")
        self._label(ws, 2, 1,
                    "Auto-classified from a department/cost-center trial balance. Provision, Cap %, "
                    "Basis & Conf are editable suggestions; split columns are live formulas.")
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=14)

        row = self._params_block(4)
        row = self._unicap_block(row + 1)
        self._grid_and_summary(row + 1)

        self._set_widths()
        os.makedirs(os.path.dirname(output_path) or '.', exist_ok=True)
        self._wb.save(output_path)
        return output_path

    # ---------------- helpers ----------------
    def _label(self, ws, r, c, text, bold=False):
        cell = ws.cell(row=r, column=c, value=text)
        cell.font = FONT_BODY_BOLD if bold else FONT_BODY
        cell.alignment = ALIGN_LEFT
        return cell

    def _money(self, ws, r, c, value, fill=None, bold=False):
        cell = ws.cell(row=r, column=c, value=value)
        cell.number_format = FMT_CURRENCY
        cell.font = FONT_BODY_BOLD if bold else FONT_BODY
        cell.alignment = ALIGN_RIGHT
        cell.border = THIN_BORDER
        if fill:
            cell.fill = fill
        return cell

    def _pct(self, ws, r, c, value, fill=None):
        cell = ws.cell(row=r, column=c, value=value)
        cell.number_format = FMT_PERCENT
        cell.font = FONT_BODY
        cell.alignment = ALIGN_RIGHT
        cell.border = THIN_BORDER
        if fill:
            cell.fill = fill
        return cell

    # ---------------- A. parameters ----------------
    def _params_block(self, start):
        ws = self._ws
        sb = self._inp.small_business
        s263a = self._inp.section_263a
        el = self._inp.section_263a_elections
        apply_section_header(ws, start, 1, 6, "Parameters")
        r = start + 1
        self._label(ws, r, 1, "§263A method", bold=True)
        self._label(ws, r, 2, s263a.method)
        r += 1
        self._label(ws, r, 1, "§448(c) 3-yr avg gross receipts", bold=True)
        self._money(ws, r, 2, float(sb.avg_annual_gross_receipts), fill=FILL_HIGHLIGHT_BLUE)
        self._gr_row = r
        r += 1
        self._label(ws, r, 1, "§448(c) threshold", bold=True)
        self._money(ws, r, 2, float(sb.threshold), fill=FILL_HIGHLIGHT_BLUE)
        self._thr_row = r
        r += 1
        self._label(ws, r, 1, "UNICAP required?", bold=True)
        cell = ws.cell(row=r, column=2,
                       value=(f'=IF(B{self._gr_row}<=B{self._thr_row},'
                              f'"No - small-business exempt (§263A(i)/§471(c))","Yes")'))
        cell.font = FONT_BODY_BOLD
        cell.fill = FILL_HIGHLIGHT_ORANGE if sb.exempt else FILL_HIGHLIGHT_GREEN
        r += 1
        self._label(ws, r, 1, "De minimis ceiling (§1.263(a)-1(f))", bold=True)
        self._money(ws, r, 2, float(el.de_minimis_ceiling), fill=FILL_HIGHLIGHT_BLUE)
        return r

    # ---------------- B. §263A absorption ----------------
    def _unicap_block(self, start):
        ws = self._ws
        s = self._inp.section_263a
        apply_section_header(ws, start, 1, 6, "§263A UNICAP — absorption ratio & §263A(f) interest")
        r = start + 1
        self._label(ws, r, 1, "§471 costs incurred", bold=True)
        self._money(ws, r, 2, float(s.section_471_costs), fill=FILL_HIGHLIGHT_BLUE)
        r471 = r
        r += 1
        self._label(ws, r, 1, "Additional §263A costs incurred", bold=True)
        self._money(ws, r, 2, float(s.additional_263a_costs), fill=FILL_HIGHLIGHT_BLUE)
        radd = r
        r += 1
        self._label(ws, r, 1, "Absorption ratio = additional / §471", bold=True)
        cell = ws.cell(row=r, column=2, value=f"=IFERROR(B{radd}/B{r471},0)")
        cell.number_format = FMT_PERCENT
        cell.alignment = ALIGN_RIGHT
        cell.border = THIN_BORDER
        cell.fill = FILL_HIGHLIGHT_GREEN
        rratio = r
        r += 1
        self._label(ws, r, 1, "§471 costs in ending inventory", bold=True)
        self._money(ws, r, 2, float(s.ending_inventory_471), fill=FILL_HIGHLIGHT_BLUE)
        rend = r
        r += 1
        self._label(ws, r, 1, "Additional §263A capitalized to ending inv", bold=True)
        self._money(ws, r, 2, None, fill=FILL_HIGHLIGHT_GREEN, bold=True).value = \
            f"=B{rend}*B{rratio}"
        r += 1
        # §263A(f) interest on one compact row: APE × rate = result
        self._label(ws, r, 1, "§263A(f) interest:  APE", bold=True)
        self._money(ws, r, 2, float(s.accumulated_production_expenditures), fill=FILL_HIGHLIGHT_BLUE)
        self._label(ws, r, 3, "× rate")
        self._pct(ws, r, 4, float(s.avoided_cost_rate), fill=FILL_HIGHLIGHT_BLUE)
        self._label(ws, r, 5, "=")
        designated = "1" if s.designated_property else "0"
        cell = ws.cell(row=r, column=6, value=f"=IF({designated}=1,B{r}*D{r},0)")
        cell.number_format = FMT_CURRENCY
        cell.font = FONT_BODY_BOLD
        cell.alignment = ALIGN_RIGHT
        cell.border = THIN_BORDER
        cell.fill = FILL_HIGHLIGHT_GREEN
        return r

    # ---------------- C. grid + D. summary ----------------
    def _grid_and_summary(self, start):
        ws = self._ws
        hdr = start
        for i, h in enumerate(GRID_HEADERS):
            ws.cell(row=hdr, column=1 + i, value=h)
        apply_header_row(ws, hdr, 1, len(GRID_HEADERS))

        dv = DataValidation(type="list", formula1=PROVISION_LIST, allow_blank=True)
        ws.add_data_validation(dv)

        per_line = self._r["per_line"]
        first = hdr + 1
        n = len(per_line) + BLANK_TEMPLATE_ROWS
        for i in range(n):
            r = first + i
            pl = per_line[i] if i < len(per_line) else None
            if pl:
                line = pl["line"]
                self._label(ws, r, 1, line.account_number)
                self._label(ws, r, 2, line.account_name)
                self._label(ws, r, 3, line.cost_center_code)
                self._label(ws, r, 4, line.provision)
                self._pct(ws, r, 5, float(line.cap_pct))
                self._money(ws, r, 6, float(line.amount), fill=FILL_HIGHLIGHT_BLUE)
                self._label(ws, r, 13, line.basis)
                self._label(ws, r, 14, line.confidence)
            else:
                for c in (1, 2, 3, 4, 13, 14):
                    ws.cell(row=r, column=c).border = THIN_BORDER
                self._pct(ws, r, 5, 1.0)
                self._money(ws, r, 6, None, fill=FILL_HIGHLIGHT_BLUE)
            dv.add(ws.cell(row=r, column=4))

            # Live split formulas (G..L)
            ws.cell(row=r, column=7, value=f'=IF($D{r}="266",$F{r}*$E{r},0)')
            ws.cell(row=r, column=8, value=f'=IF($D{r}="263(a)",$F{r}*$E{r},0)')
            ws.cell(row=r, column=9, value=f'=IF(OR($D{r}="263A",$D{r}="mixed"),$F{r}*$E{r},0)')
            ws.cell(row=r, column=10, value=f'=IF(LEFT($D{r},5)="other",$F{r}*$E{r},0)')
            ws.cell(row=r, column=11, value=f'=IF($D{r}="book_capitalized",$F{r},0)')
            ws.cell(row=r, column=12, value=f'=$F{r}-($G{r}+$H{r}+$I{r}+$J{r}+$K{r})')
            for c in range(7, 13):
                cell = ws.cell(row=r, column=c)
                cell.number_format = FMT_CURRENCY
                cell.alignment = ALIGN_RIGHT
                cell.border = THIN_BORDER
                cell.fill = FILL_HIGHLIGHT_GREEN
        last = first + n - 1

        # Totals row
        tr = last + 1
        self._label(ws, tr, 1, "TOTAL", bold=True)
        for c in range(6, 13):
            cl = ws.cell(row=tr, column=c).column_letter
            cell = ws.cell(row=tr, column=c, value=f"=SUM({cl}{first}:{cl}{last})")
            cell.number_format = FMT_CURRENCY
            cell.font = FONT_BODY_BOLD
            cell.alignment = ALIGN_RIGHT
            cell.border = THIN_BORDER

        # Conditional formatting: low-confidence + blank provision
        rng = f"A{first}:N{last}"
        ws.conditional_formatting.add(
            rng, FormulaRule(formula=[f'$N{first}="low"'], fill=FILL_HIGHLIGHT_ORANGE))
        ws.conditional_formatting.add(
            f"D{first}:D{last}",
            FormulaRule(formula=[f'AND($F{first}<>"",$D{first}="")'], fill=FILL_HIGHLIGHT_RED))

        # D. Schedule M-1 / summary strip
        sr = tr + 2
        apply_section_header(ws, sr, 1, 4, "Summary & Schedule M-1")
        sr += 1
        rows = [
            ("Total tax-capitalized (§266+§263(a)+§263A+Other)",
             f"=G{tr}+H{tr}+I{tr}+J{tr}", FILL_HIGHLIGHT_GREEN),
            ("Already capitalized for book", f"=K{tr}", None),
            ("Total deductible", f"=L{tr}", None),
            ("Trial balance total (check)", f"=F{tr}", None),
            ("Tie check (must be 0)",
             f"=(G{tr}+H{tr}+I{tr}+J{tr}+K{tr}+L{tr})-F{tr}", FILL_HIGHLIGHT_ORANGE),
            ("Schedule M-1 addback (book-deducted, now capitalized)",
             f"=G{tr}+H{tr}+I{tr}+J{tr}", FILL_HIGHLIGHT_GREEN),
        ]
        for label, formula, fill in rows:
            self._label(ws, sr, 1, label, bold=True)
            self._money(ws, sr, 3, None, fill=fill, bold=True).value = formula
            sr += 1
        self._label(ws, sr + 1, 1,
                    "Mixed lines route their capitalized % to §263A (mixed service cost, §1.263A-1(e)(4)); "
                    "the §263A absorption add-on above feeds inventory basis separately.")
        ws.merge_cells(start_row=sr + 1, start_column=1, end_row=sr + 1, end_column=14)

    def _set_widths(self):
        widths = {"A": 12, "B": 30, "C": 18, "D": 16, "E": 8, "F": 15,
                  "G": 13, "H": 13, "I": 13, "J": 13, "K": 13, "L": 14,
                  "M": 60, "N": 8}
        for col, w in widths.items():
            self._ws.column_dimensions[col].width = w
